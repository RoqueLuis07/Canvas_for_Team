"""Borradores editables tipo Excel ("staging"), vivos dentro de la web.

Reemplaza el flujo de "bajar el Excel de OneDrive, editarlo en Excel, volver
a subirlo" por una grilla editable (agregar/borrar filas, editar celdas,
pegar un rango copiado de Excel) persistida en la base, que se puede
sincronizar con el archivo real de OneDrive en cualquier momento — no un
reemplazo de OneDrive, un espejo editable de ida y vuelta.

El procesamiento real (crear cuentas en Canvas/Teams, matricular, dar de
baja) NO se reimplementa acá: se reutiliza tal cual el de app.routers.excel,
ya probado en producción. Este router solo se encarga de la parte nueva
(edición en vivo + sincronización con OneDrive) y, al procesar, arma un
`DiplomadosUrlRequest` apuntando al mismo archivo/pestaña de OneDrive —ya
actualizado con lo que esté en la grilla en ese momento— y llama a la
función de procesamiento existente, igual que si el operador hubiera usado
el flujo de siempre.
"""
import io
import logging

import openpyxl
from openpyxl.styles import Font
from fastapi import APIRouter, HTTPException, BackgroundTasks
from pydantic import BaseModel

from app.core import staging_db
from app.core import jobs
from app.services import teams_client as graph
from app.routers import excel as excel_router

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/staging", tags=["Staging (grilla editable)"])

# Tipos de flujo soportados hoy. "masivo" y "egreso" procesan sincrónico
# (devuelven el resultado directo); "diplomados" y "matriculaciones" corren
# en segundo plano como ya lo hacían, y se consulta el progreso con el
# mismo endpoint de siempre: GET /excel/jobs/{job_id}.
FLOW_TYPES = {"diplomados", "matriculaciones", "egreso", "masivo"}


class CreateSheetRequest(BaseModel):
    flow_type: str
    title: str | None = None
    onedrive_url: str | None = None
    sheet_name: str | None = None


class UpdateSheetRequest(BaseModel):
    title: str | None = None
    onedrive_url: str | None = None
    sheet_name: str | None = None


class UpsertRowsRequest(BaseModel):
    rows: list[dict]
    """Cada item: {"row_index": int, "data": {"<header>": "<valor>", ...}}"""


class AddRowRequest(BaseModel):
    data: dict | None = None


class ProcessRequest(BaseModel):
    delete_account: bool = False
    cc: list[str] = []
    report_url: str | None = None


def _require_flow_type(flow_type: str):
    if flow_type not in FLOW_TYPES:
        raise HTTPException(status_code=400, detail=f"flow_type inválido. Debe ser uno de: {', '.join(sorted(FLOW_TYPES))}")


async def _get_sheet_or_404(sheet_id: int) -> dict:
    sheet = await staging_db.get_sheet(sheet_id)
    if not sheet:
        raise HTTPException(status_code=404, detail="Borrador no encontrado.")
    return sheet


# ── CRUD del borrador ────────────────────────────────────────────────────

@router.post("/sheets", summary="Crear un borrador nuevo")
async def create_sheet(req: CreateSheetRequest) -> dict:
    _require_flow_type(req.flow_type)
    sheet_id = await staging_db.create_sheet(
        flow_type=req.flow_type,
        title=req.title or req.flow_type,
        onedrive_url=req.onedrive_url,
        sheet_name=req.sheet_name,
        headers=[],
        created_by="admin",
    )
    return await _get_sheet_or_404(sheet_id)


@router.get("/sheets", summary="Listar borradores")
async def list_sheets(flow_type: str | None = None) -> list[dict]:
    return await staging_db.list_sheets(flow_type)


@router.get("/sheets/{sheet_id}", summary="Obtener un borrador con sus filas")
async def get_sheet(sheet_id: int) -> dict:
    sheet = await _get_sheet_or_404(sheet_id)
    rows = await staging_db.get_rows(sheet_id)
    return {"sheet": sheet, "rows": rows}


@router.patch("/sheets/{sheet_id}", summary="Actualizar título/URL de OneDrive/pestaña")
async def update_sheet(sheet_id: int, req: UpdateSheetRequest) -> dict:
    await _get_sheet_or_404(sheet_id)
    import json
    from app.core.staging_db import _run
    def _upd(conn):
        cursor = conn.cursor()
        fields, params = [], []
        if req.title is not None:
            fields.append("title = %s"); params.append(req.title)
        if req.onedrive_url is not None:
            fields.append("onedrive_url = %s"); params.append(req.onedrive_url)
        if req.sheet_name is not None:
            fields.append("sheet_name = %s"); params.append(req.sheet_name)
        if not fields:
            return
        params.append(sheet_id)
        cursor.execute(f"UPDATE staging_sheets SET {', '.join(fields)}, updated_at = CURRENT_TIMESTAMP WHERE id = %s", params)
        conn.commit()
    await _run(_upd)
    return await _get_sheet_or_404(sheet_id)


@router.delete("/sheets/{sheet_id}", summary="Eliminar un borrador")
async def delete_sheet(sheet_id: int) -> dict:
    deleted = await staging_db.delete_sheet(sheet_id)
    if not deleted:
        raise HTTPException(status_code=404, detail="Borrador no encontrado.")
    return {"status": "deleted"}


# ── Edición de filas ─────────────────────────────────────────────────────

@router.put("/sheets/{sheet_id}/rows", summary="Editar/pegar filas (upsert masivo)")
async def upsert_rows(sheet_id: int, req: UpsertRowsRequest) -> dict:
    await _get_sheet_or_404(sheet_id)
    for r in req.rows:
        if "row_index" not in r or "data" not in r:
            raise HTTPException(status_code=400, detail='Cada fila necesita "row_index" y "data".')
    await staging_db.upsert_rows(sheet_id, req.rows)
    return {"status": "ok", "updated": len(req.rows)}


@router.post("/sheets/{sheet_id}/rows", summary="Agregar una fila vacía")
async def add_row(sheet_id: int, req: AddRowRequest) -> dict:
    await _get_sheet_or_404(sheet_id)
    row_index = await staging_db.add_row(sheet_id, req.data)
    return {"row_index": row_index}


@router.delete("/sheets/{sheet_id}/rows/{row_index}", summary="Borrar una fila")
async def delete_row(sheet_id: int, row_index: int) -> dict:
    await _get_sheet_or_404(sheet_id)
    deleted = await staging_db.delete_row(sheet_id, row_index)
    if not deleted:
        raise HTTPException(status_code=404, detail="Fila no encontrada.")
    return {"status": "deleted"}


# ── Sincronización con OneDrive ──────────────────────────────────────────

@router.post("/sheets/{sheet_id}/onedrive-sheets", summary="Listar las pestañas del archivo de OneDrive configurado")
async def list_onedrive_sheets(sheet_id: int) -> list[str]:
    sheet = await _get_sheet_or_404(sheet_id)
    if not sheet.get("onedrive_url"):
        raise HTTPException(status_code=400, detail="El borrador todavía no tiene una URL de OneDrive configurada.")
    encoded_url = excel_router._encode_share_url(sheet["onedrive_url"])
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
        sheets = wb.sheetnames
        wb.close()
        return sheets
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo leer el archivo de OneDrive. {e}")


@router.post("/sheets/{sheet_id}/pull", summary="Traer datos desde OneDrive (reemplaza la grilla actual)")
async def pull_from_onedrive(sheet_id: int) -> dict:
    sheet = await _get_sheet_or_404(sheet_id)
    if not sheet.get("onedrive_url") or not sheet.get("sheet_name"):
        raise HTTPException(status_code=400, detail="Configurá la URL de OneDrive y la pestaña antes de traer los datos.")

    encoded_url = excel_router._encode_share_url(sheet["onedrive_url"])
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo de OneDrive. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="El archivo no es un Excel válido.")

    if sheet["sheet_name"] not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"La pestaña '{sheet['sheet_name']}' no existe. Disponibles: {', '.join(wb.sheetnames)}")

    ws = wb[sheet["sheet_name"]]
    header_row_idx, headers_dict, headers_raw = excel_router._find_header_row_and_headers(ws)
    if not header_row_idx:
        wb.close()
        raise HTTPException(status_code=400, detail="No se encontró la fila de encabezados en esa pestaña.")

    headers = [h for h in headers_raw if h]
    rows: list[dict] = []
    for row_idx in range(header_row_idx + 1, ws.max_row + 1):
        row_vals = [str(ws.cell(row=row_idx, column=c).value or "").strip() for c in range(1, len(headers_raw) + 1)]
        if not any(row_vals):
            continue
        row_data = {h: v for h, v in zip(headers_raw, row_vals) if h}
        rows.append(row_data)
    wb.close()

    count = await staging_db.replace_rows(sheet_id, headers, rows)
    return {"status": "ok", "headers": headers, "rows_imported": count}


@router.post("/sheets/{sheet_id}/push", summary="Guardar la grilla actual en el archivo de OneDrive")
async def push_to_onedrive(sheet_id: int) -> dict:
    sheet = await _get_sheet_or_404(sheet_id)
    if not sheet.get("onedrive_url") or not sheet.get("sheet_name"):
        raise HTTPException(status_code=400, detail="Configurá la URL de OneDrive y la pestaña antes de guardar.")

    headers: list[str] = sheet.get("headers") or []
    if not headers:
        raise HTTPException(status_code=400, detail="El borrador no tiene columnas definidas todavía (traé datos de OneDrive primero, o cargalas a mano).")

    rows = await staging_db.get_rows(sheet_id)

    encoded_url = excel_router._encode_share_url(sheet["onedrive_url"])
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo de OneDrive. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents))
    except Exception:
        raise HTTPException(status_code=400, detail="El archivo no es un Excel válido.")

    if sheet["sheet_name"] in wb.sheetnames:
        ws = wb[sheet["sheet_name"]]
        header_row_idx, _, existing_headers_raw = excel_router._find_header_row_and_headers(ws)
        prev_max_row = ws.max_row
    else:
        ws = wb.create_sheet(sheet["sheet_name"])
        header_row_idx = None
        existing_headers_raw = []
        prev_max_row = 0

    if not header_row_idx:
        header_row_idx = 1

    # Encabezados (se preserva/crea el orden que ya tiene el borrador).
    for col_idx, h in enumerate(headers, 1):
        cell = excel_router._safe_set_cell(ws, header_row_idx, col_idx, h)
        cell.font = Font(bold=True)

    # Limpiar el área de datos previa (debajo del encabezado) antes de
    # reescribirla con las filas actuales del borrador — evita dejar filas
    # viejas "fantasma" si el borrador ahora tiene menos filas que antes.
    for row_idx in range(header_row_idx + 1, max(prev_max_row, header_row_idx) + 1):
        for col_idx in range(1, len(headers) + 1):
            excel_router._safe_set_cell(ws, row_idx, col_idx, None)

    for offset, row in enumerate(rows):
        row_idx = header_row_idx + 1 + offset
        data = row["data"]
        for col_idx, h in enumerate(headers, 1):
            excel_router._safe_set_cell(ws, row_idx, col_idx, data.get(h))

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()

    try:
        await graph.put_raw(f"/shares/{encoded_url}/driveItem/content", buf.getvalue())
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo guardar en OneDrive. {e}")

    return {"status": "ok", "rows_written": len(rows)}


# ── Procesar (reutiliza la lógica real ya probada de app.routers.excel) ──

@router.post("/sheets/{sheet_id}/process", summary="Guardar en OneDrive y procesar (crear cuentas/matricular/dar de baja)")
async def process_sheet(sheet_id: int, req: ProcessRequest, bg_tasks: BackgroundTasks) -> dict:
    sheet = await _get_sheet_or_404(sheet_id)
    flow_type = sheet["flow_type"]
    _require_flow_type(flow_type)

    # Siempre se guarda en OneDrive primero: el procesamiento real (igual
    # que hoy) vuelve a leer el archivo de OneDrive, así que si no
    # sincronizamos antes, procesaría datos viejos y la edición en la
    # grilla no tendría ningún efecto.
    await push_to_onedrive(sheet_id)

    onedrive_req = excel_router.DiplomadosUrlRequest(
        url=sheet["onedrive_url"],
        sheet_name=sheet["sheet_name"],
        delete_account=req.delete_account,
        cc=req.cc,
        report_url=req.report_url,
    )
    encoded_url = excel_router._encode_share_url(sheet["onedrive_url"])

    if flow_type in ("diplomados", "matriculaciones"):
        try:
            contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"No se pudo releer el archivo recién guardado en OneDrive. {e}")

        if flow_type == "diplomados":
            job_id = await jobs.create_job(job_type="diplomados_onedrive", operation="Alta de Diplomados (desde grilla)", username="admin")
            bg_tasks.add_task(excel_router._process_diplomados_bg, job_id, onedrive_req, contents, encoded_url)
        else:
            job_id = await jobs.create_job(job_type="matriculacion_masiva_onedrive", operation="Matriculación (desde grilla)", username="admin")
            bg_tasks.add_task(excel_router._process_matriculaciones_bg, job_id, onedrive_req, contents, encoded_url)

        return {"status": "processing", "job_id": job_id, "poll_url": f"/excel/jobs/{job_id}"}

    if flow_type == "masivo":
        result = await excel_router.import_masivo_onedrive(onedrive_req)
        return {"status": "completed", "result": result}

    if flow_type == "egreso":
        result = await excel_router._import_egreso_onedrive_inner(onedrive_req)
        return {"status": "completed", "result": result}

    raise HTTPException(status_code=400, detail=f"Procesamiento no implementado todavía para '{flow_type}'.")
