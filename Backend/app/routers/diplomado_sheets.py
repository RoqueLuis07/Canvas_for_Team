"""Planillas de trabajo de Diplomado — vista real (no maqueta) para alta de
alumnos y envío de credenciales, con un panel de candidatos de solo lectura
traído de la planilla comercial maestra de OneDrive ("Datos UBS").

Diseño (ver Frontend/templates/canvas/diplomados_sheet_preview.html para la
maqueta visual original aprobada por el usuario):
  - La tabla `diplomado_sheets`/`diplomado_rows` (app.core.diplomado_sheets_db)
    ES la fuente de verdad — no hay OneDrive de por medio para el trabajo
    diario, a diferencia del flujo "staging" de la rama hermana.
  - El panel de candidatos lee "Datos UBS" en SOLO LECTURA vía Graph — nunca
    se escribe ahí, ni siquiera para marcar quién ya fue procesado. El cruce
    "ya tiene credenciales" / "falta enviar" se calcula comparando por
    cédula contra las filas que YA están en `diplomado_rows`.

Puente "espejo a OneDrive" (POST .../process): la creación real de cuentas
en app.routers.excel (_process_diplomados_bg) depende intrínsecamente de
leer y volver a escribir un archivo real de OneDrive — así es como esa
lógica ya probada en producción sabe qué filas ya están creadas en corridas
futuras (lo persiste en las columnas Usuario/Contraseña/Estado del propio
archivo). En vez de reimplementar esa lógica de negocio en paralelo (con el
riesgo de que diverja de la ya probada), este router genera un archivo
Excel descartable a partir de esta tabla, lo sube al OneDrive del buzón
SMTP_FROM (permiso de aplicación 'Files.ReadWrite.All', ya concedido), y
dispara la misma `_process_diplomados_bg` de excel.py contra ese archivo.
La tabla sigue siendo la fuente de verdad: el archivo es un detalle de
implementación que el usuario nunca necesita abrir — después de procesar,
`/sync-from-mirror` trae de vuelta Usuario/Contraseña/Estado a esta tabla.
"""
import io
import logging
import re

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font
from fastapi import APIRouter, BackgroundTasks, HTTPException
from pydantic import BaseModel

from app.core import diplomado_sheets_db as db
from app.core import jobs
from app.core.config import settings
from app.services import teams_client as graph
from app.routers.excel import (
    _encode_share_url,
    _find_header_row_and_headers,
    _clean_cedula,
    _norm,
    _process_diplomados_bg,
    DiplomadosUrlRequest,
)


def _match_col_idx(headers: dict, *keys: str) -> int | None:
    """excel.py en esta rama no tiene un `_match_col_idx` compartido (cada
    función define su propio `get_col_idx` local con coincidencia por
    substring) — se replica acá la misma estrategia para leer de vuelta el
    archivo-espejo con el mismo criterio que usó `_process_diplomados_bg`
    para escribirlo."""
    for k in keys:
        nk = _norm(k)
        for h, idx in headers.items():
            if nk in h:
                return idx
    return None

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/diplomado-sheets", tags=["Planillas de Diplomado"])

# Nombre fijo de la única pestaña del archivo-espejo — se genera y se lee
# siempre con este mismo nombre, nunca lo ve el usuario.
_MIRROR_SHEET_NAME = "Planilla"
# Orden de columnas fijo del archivo-espejo (1-indexado para openpyxl).
_COL_FECHA, _COL_NOMBRE, _COL_CEDULA, _COL_CORREO, _COL_TELEFONO, _COL_USUARIO, _COL_CONTRA, _COL_ESTADO = range(1, 9)


def _build_mirror_workbook(sheet: dict, rows: list[dict]) -> bytes:
    """Arma el archivo-espejo desde cero a partir del estado ACTUAL de la
    tabla (incluyendo usuario/contraseña/estado de corridas anteriores, si
    los hay) — así _process_diplomados_bg reconoce correctamente qué filas
    ya están creadas y no las vuelve a procesar ni duplica cuentas."""
    wb = Workbook()
    ws = wb.active
    ws.title = _MIRROR_SHEET_NAME

    # Fila 1: título (nombre del diplomado, lo que _find_row1_title busca) +
    # el ID del equipo de Teams ya creado, si lo hay, en la columna que va a
    # ser "Usuario" — mismo lugar donde _process_diplomados_bg lo busca para
    # no crear un Team duplicado en cada corrida.
    ws.cell(row=1, column=1, value=sheet["diplomado_name"])
    if sheet.get("team_id"):
        ws.cell(row=1, column=_COL_USUARIO, value=sheet["team_id"])

    headers = ["Fecha", "Nombre", "Cedula", "Correo", "Telefono", "Usuario", "Contraseña", "Estado"]
    for col_idx, h in enumerate(headers, 1):
        ws.cell(row=2, column=col_idx, value=h).font = Font(bold=True)

    for r_offset, row in enumerate(rows, start=3):
        ws.cell(row=r_offset, column=_COL_FECHA, value=row.get("fecha") or "")
        ws.cell(row=r_offset, column=_COL_NOMBRE, value=row.get("nombre_completo") or "")
        ws.cell(row=r_offset, column=_COL_CEDULA, value=row.get("cedula") or "")
        ws.cell(row=r_offset, column=_COL_CORREO, value=row.get("correo") or "")
        ws.cell(row=r_offset, column=_COL_TELEFONO, value=row.get("telefono") or "")
        if row.get("usuario"):
            ws.cell(row=r_offset, column=_COL_USUARIO, value=row["usuario"])
        if row.get("contrasena"):
            ws.cell(row=r_offset, column=_COL_CONTRA, value=row["contrasena"])
        if row.get("estado"):
            ws.cell(row=r_offset, column=_COL_ESTADO, value=row["estado"])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _mirror_drive_path(sheet_id: int, diplomado_name: str) -> str:
    safe_name = re.sub(r"[^\w\- ]", "", diplomado_name)[:60].strip() or "Diplomado"
    return f"Planillas en vivo/Diplomados/{sheet_id} - {safe_name}.xlsx"


class CreateSheetRequest(BaseModel):
    diplomado_name: str
    programa_ubs: str | None = None
    team_id: str | None = None
    owners: list[str] = []


class UpdateSheetRequest(BaseModel):
    diplomado_name: str | None = None
    programa_ubs: str | None = None
    team_id: str | None = None
    owners: list[str] | None = None


class RowData(BaseModel):
    fecha: str | None = None
    nombre_completo: str | None = None
    cedula: str | None = None
    correo: str | None = None
    telefono: str | None = None
    vendedor: str | None = None
    rol: str | None = None
    estado: str | None = None


class TransferRequest(BaseModel):
    candidates: list[RowData]


async def _get_sheet_or_404(sheet_id: int) -> dict:
    sheet = await db.get_sheet(sheet_id)
    if not sheet:
        raise HTTPException(status_code=404, detail="Planilla de diplomado no encontrada.")
    return sheet


# ── CRUD de la planilla ──────────────────────────────────────────────────

@router.post("/sheets")
async def create_sheet(req: CreateSheetRequest) -> dict:
    if not req.diplomado_name.strip():
        raise HTTPException(status_code=400, detail="El nombre del diplomado es obligatorio.")
    sheet_id = await db.create_sheet(
        diplomado_name=req.diplomado_name.strip(),
        programa_ubs=(req.programa_ubs or "").strip() or None,
        team_id=(req.team_id or "").strip() or None,
        owners=req.owners,
        created_by="admin",
    )
    return await _get_sheet_or_404(sheet_id)


@router.get("/sheets")
async def list_sheets() -> list[dict]:
    return await db.list_sheets()


@router.get("/sheets/{sheet_id}")
async def get_sheet(sheet_id: int) -> dict:
    sheet = await _get_sheet_or_404(sheet_id)
    rows = await db.get_rows(sheet_id)
    return {"sheet": sheet, "rows": rows}


@router.patch("/sheets/{sheet_id}")
async def update_sheet(sheet_id: int, req: UpdateSheetRequest) -> dict:
    await _get_sheet_or_404(sheet_id)
    await db.update_sheet(sheet_id, req.diplomado_name, req.programa_ubs, req.team_id, req.owners)
    return await _get_sheet_or_404(sheet_id)


@router.delete("/sheets/{sheet_id}")
async def delete_sheet(sheet_id: int) -> dict:
    deleted = await db.delete_sheet(sheet_id)
    if not deleted:
        raise HTTPException(status_code=404, detail="Planilla de diplomado no encontrada.")
    return {"status": "deleted"}


# ── Filas de la planilla de trabajo ──────────────────────────────────────

@router.post("/sheets/{sheet_id}/rows")
async def add_row(sheet_id: int, req: RowData) -> dict:
    await _get_sheet_or_404(sheet_id)
    if not (req.nombre_completo or "").strip() and not (req.cedula or "").strip():
        # Fila vacía agregada desde el botón "Agregar alumno" — se permite,
        # se completa a mano y se guarda cuando el operador escribe algo.
        return await db.add_row(sheet_id, {})
    return await db.add_row(sheet_id, req.model_dump(exclude_none=True))


@router.patch("/sheets/{sheet_id}/rows/{row_id}")
async def update_row(sheet_id: int, row_id: int, req: RowData) -> dict:
    await _get_sheet_or_404(sheet_id)
    updated = await db.update_row(sheet_id, row_id, req.model_dump(exclude_none=True))
    if not updated:
        raise HTTPException(status_code=404, detail="Fila no encontrada.")
    return {"status": "ok"}


@router.delete("/sheets/{sheet_id}/rows/{row_id}")
async def delete_row(sheet_id: int, row_id: int) -> dict:
    await _get_sheet_or_404(sheet_id)
    deleted = await db.delete_row(sheet_id, row_id)
    if not deleted:
        raise HTTPException(status_code=404, detail="Fila no encontrada.")
    return {"status": "deleted"}


# ── Candidatos (solo lectura, desde la planilla comercial de OneDrive) ───

@router.get("/sheets/{sheet_id}/candidates")
async def get_candidates(sheet_id: int) -> dict:
    """Lee 'Datos UBS' en SOLO LECTURA (nunca escribe ahí), filtra por el
    programa de este diplomado, y cruza por cédula contra las filas que ya
    están en la planilla de trabajo real. No requiere que el archivo tenga
    ninguna marca manual — el cruce lo hace el sistema."""
    sheet = await _get_sheet_or_404(sheet_id)

    if not settings.comercial_onedrive_url:
        raise HTTPException(
            status_code=400,
            detail="No hay una planilla comercial configurada (COMERCIAL_ONEDRIVE_URL en el .env). "
                   "Pedile al administrador que la configure para poder ver candidatos.",
        )

    encoded_url = _encode_share_url(settings.comercial_onedrive_url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"No se pudo leer la planilla comercial desde OneDrive. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=502, detail="La planilla comercial descargada no es un Excel válido.")

    sheet_name = settings.comercial_sheet_name
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise HTTPException(status_code=502, detail=f"La pestaña '{sheet_name}' no existe en la planilla comercial.")

    ws = wb[sheet_name]
    header_row_idx, headers, _raw = _find_header_row_and_headers(ws)
    if not header_row_idx:
        wb.close()
        raise HTTPException(status_code=502, detail="No se encontró la fila de encabezados en 'Datos UBS'.")

    def col(*keys):
        for k in keys:
            for h, idx in headers.items():
                if _norm(k) in h:
                    return idx
        return None

    col_nombre = col("nombre completo", "nombre")
    col_cedula = col("cedula", "cédula")
    col_correo = col("email", "correo")
    col_telefono = col("telefono", "celular")
    col_vendedor = col("asesor", "vendedor")
    col_programa = col("programa ubs", "programa")
    col_situacion = col("situacion")

    if not col_nombre or not col_cedula or not col_programa:
        wb.close()
        raise HTTPException(status_code=502, detail="No se encontraron las columnas esperadas (Nombre/Cédula/Programa) en 'Datos UBS'.")

    programa_filtro = _norm(sheet.get("programa_ubs") or sheet["diplomado_name"])
    existing_cedulas = await db.get_existing_cedulas(sheet_id)

    candidates = []
    for row_idx in range(header_row_idx + 1, ws.max_row + 1):
        nombre = str(ws.cell(row=row_idx, column=col_nombre).value or "").strip()
        cedula = _clean_cedula(str(ws.cell(row=row_idx, column=col_cedula).value or "").strip())
        programa = str(ws.cell(row=row_idx, column=col_programa).value or "").strip()
        if not nombre or not cedula or not programa:
            continue
        # Coincidencia parcial normalizada: los nombres de programa en Datos
        # UBS suelen traer sufijos de fecha ("... - 12 Octubre") que no
        # necesariamente coinciden carácter a carácter con el nombre del
        # diplomado en el sistema.
        if programa_filtro not in _norm(programa) and _norm(programa) not in programa_filtro:
            continue
        candidates.append({
            "nombre_completo": nombre,
            "cedula": cedula,
            "correo": str(ws.cell(row=row_idx, column=col_correo).value or "").strip() if col_correo else "",
            "telefono": str(ws.cell(row=row_idx, column=col_telefono).value or "").strip() if col_telefono else "",
            "vendedor": str(ws.cell(row=row_idx, column=col_vendedor).value or "").strip() if col_vendedor else "",
            "programa_ubs": programa,
            "situacion": str(ws.cell(row=row_idx, column=col_situacion).value or "").strip() if col_situacion else "",
            "ya_tiene_credenciales": cedula in existing_cedulas,
        })
    wb.close()

    return {"candidates": candidates, "programa_filtro": sheet.get("programa_ubs") or sheet["diplomado_name"]}


@router.post("/sheets/{sheet_id}/transfer")
async def transfer_candidates(sheet_id: int, req: TransferRequest) -> dict:
    """Copia los candidatos elegidos a la planilla de trabajo real (en la
    base de datos). Nunca toca la planilla comercial de origen."""
    await _get_sheet_or_404(sheet_id)
    from datetime import date
    today = date.today().strftime("%d/%m/%Y")
    rows = []
    for c in req.candidates:
        data = c.model_dump(exclude_none=True)
        data.setdefault("fecha", today)
        data.setdefault("rol", "Alumno")
        data.setdefault("estado", "pendiente")
        rows.append(data)
    inserted = await db.bulk_add_rows(sheet_id, rows)
    return {"status": "ok", "inserted": inserted}


# ── Procesar / enviar credenciales ───────────────────────────────────────

@router.post("/sheets/{sheet_id}/process")
async def process_sheet(sheet_id: int, bg_tasks: BackgroundTasks) -> dict:
    """Genera (o actualiza) el archivo-espejo en OneDrive a partir del
    estado actual de la planilla y dispara la misma creación de cuentas ya
    probada en producción (excel.py:_process_diplomados_bg) contra ese
    archivo. Devuelve un job_id para seguir el progreso igual que el resto
    de los flujos de alta masiva (GET /excel/jobs/{job_id})."""
    sheet = await _get_sheet_or_404(sheet_id)
    rows = await db.get_rows(sheet_id)
    if not rows:
        raise HTTPException(status_code=400, detail="La planilla no tiene alumnos cargados todavía.")

    content = _build_mirror_workbook(sheet, rows)
    drive_path = _mirror_drive_path(sheet_id, sheet["diplomado_name"])

    try:
        await graph.create_or_update_drive_file(drive_path, content)
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"No se pudo subir el archivo-espejo a OneDrive. {e}")

    mirror_url = sheet.get("mirror_url")
    if not mirror_url:
        try:
            mirror_url = await graph.create_sharing_link(drive_path)
        except Exception as e:
            raise HTTPException(status_code=502, detail=f"No se pudo generar el link del archivo-espejo. {e}")
        if not mirror_url:
            raise HTTPException(status_code=502, detail="No se pudo generar el link del archivo-espejo en OneDrive.")
        await db.set_mirror_url(sheet_id, mirror_url)

    req = DiplomadosUrlRequest(url=mirror_url, sheet_name=_MIRROR_SHEET_NAME)
    encoded_url = _encode_share_url(mirror_url)

    job_id = await jobs.create_job(
        job_type="diplomados_onedrive",
        operation=f"Alta de Diplomados (planilla en vivo #{sheet_id})",
        username="admin",
    )
    bg_tasks.add_task(_process_diplomados_bg, job_id, req, content, encoded_url)

    return {"status": "success", "job_id": job_id, "message": "Alta de Diplomados iniciada en segundo plano."}


@router.post("/sheets/{sheet_id}/sync-from-mirror")
async def sync_from_mirror(sheet_id: int) -> dict:
    """Después de que el job de /process termina, trae de vuelta
    Usuario/Contraseña/Estado desde el archivo-espejo hacia esta tabla, para
    que la grilla en pantalla refleje el resultado real. Se llama manual
    (botón "Refrescar estado") o automáticamente al terminar el polling del
    job, según lo resuelva el frontend."""
    sheet = await _get_sheet_or_404(sheet_id)
    mirror_url = sheet.get("mirror_url")
    if not mirror_url:
        raise HTTPException(status_code=400, detail="Esta planilla todavía no se procesó ninguna vez.")

    encoded_url = _encode_share_url(mirror_url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"No se pudo leer el archivo-espejo. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        ws = wb[_MIRROR_SHEET_NAME] if _MIRROR_SHEET_NAME in wb.sheetnames else wb.worksheets[0]
    except Exception:
        raise HTTPException(status_code=502, detail="El archivo-espejo descargado no es un Excel válido.")

    header_row_idx, headers, _raw = _find_header_row_and_headers(ws)
    if not header_row_idx:
        raise HTTPException(status_code=502, detail="No se encontró la fila de encabezados en el archivo-espejo.")

    col_cedula = _match_col_idx(headers, "cedula", "cédula")
    col_usuario = _match_col_idx(headers, "usuario")
    col_contra = _match_col_idx(headers, "contrasena", "contraseña")
    col_estado = _match_col_idx(headers, "estado", "enviado")

    if not col_cedula:
        raise HTTPException(status_code=502, detail="El archivo-espejo no tiene columna de Cédula — no se pudo sincronizar.")

    updated = 0
    for r_idx in range(header_row_idx + 1, ws.max_row + 1):
        cedula = _clean_cedula(str(ws.cell(row=r_idx, column=col_cedula).value or "").strip())
        if not cedula:
            continue
        data: dict = {}
        if col_usuario:
            v = str(ws.cell(row=r_idx, column=col_usuario).value or "").strip()
            if v:
                data["usuario"] = v
        if col_contra:
            v = str(ws.cell(row=r_idx, column=col_contra).value or "").strip()
            if v:
                data["contrasena"] = v
        if col_estado:
            v = str(ws.cell(row=r_idx, column=col_estado).value or "").strip()
            if v:
                data["estado"] = v
        if data and await db.update_row_by_cedula(sheet_id, cedula, data):
            updated += 1

    return {"status": "ok", "updated": updated}
