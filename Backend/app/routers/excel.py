"""Excel import endpoints + template downloads.

All templates use Spanish headers; upload endpoints accept both Spanish
and English column names via the _get() helper that tries multiple keys.
"""
import asyncio
import io
import logging
import re
import unicodedata
from typing import Any

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
from fastapi import APIRouter, HTTPException, UploadFile, File, Depends, BackgroundTasks
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from starlette.exceptions import HTTPException as StarletteHTTPException

from app.core.config import settings
from app.models.canvas import BulkResult
from app.services import canvas_client as canvas
from app.services import teams_client as graph
from app.services import user_service
from app.services.credential_generator import generate_password
from app.services import email_service
from app.services.teams_client import create_team_via_group
from app.core import jobs
from app.core import database

def _err(exc: Exception) -> str:
    return getattr(exc, "detail", str(exc))

router = APIRouter(tags=["Excel Import/Export"])
_ACCOUNT = settings.canvas_account_id
logger = logging.getLogger(__name__)

# Propietarios que deben agregarse a todo Team de Diplomados nuevo, sin excepción.
_DIPLOMADO_DEFAULT_OWNERS = ["ldure@usil.edu.py", "jfleitas@usil.edu.py", "ralvarez@usil.edu.py"]

# Propietario adicional exclusivo de los Teams de MBA (se suma a los de
# arriba, no los reemplaza).
_MBA_EXTRA_OWNER = "amichelena@usil.edu.py"

# Límites de seguridad
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB
MAX_ROWS = 10000
ALLOWED_EXCEL_TYPES = {
    "application/vnd.ms-excel",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/x-xlsx"
}


# ── Excel read/normalize helpers ─────────────────────────────────────────────

def _norm(s: str) -> str:
    """Normalize a header string: lowercase, strip accents, spaces→underscore."""
    s = s.strip().lower()
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_")


def _match_col_idx(headers: dict, *keys: str) -> int | None:
    """Encuentra la columna cuyo encabezado corresponde a alguno de los
    `keys` (ej. "cedula", "ci"), probando primero por COINCIDENCIA EXACTA
    del encabezado normalizado, y solo si no hay ninguna cayendo a
    coincidencia parcial (substring) — en ese orden de prioridad.

    Claves cortas tipo "ci" son abreviaturas que también aparecen como
    substring de un montón de palabras comunes en español ("inscripción",
    "comercial", "servicio"...). Sin priorizar la coincidencia exacta,
    una columna de fecha o de cualquier otra cosa puede "ganarle" a la
    columna real de Cédula/CI y corromper en silencio la contraseña
    generada (cedula-Iniciales termina siendo fecha-Iniciales) — bug real
    visto en una planilla de MBA con una columna no relacionada que
    contenía "ci" en su encabezado antes de llegar a la columna "CI".
    """
    norm_keys = [_norm(k) for k in keys]
    for nk in norm_keys:
        for h, idx in headers.items():
            if h == nk:
                return idx
    for nk in norm_keys:
        for h, idx in headers.items():
            if nk in h:
                return idx
    return None


def _collision_note(creds: dict) -> str:
    """Arma el aviso de colisión de nombre para la columna de Estado, según
    si se pudo cruzar por cédula (SIS ID / postalCode) o no:
    - "different": la cédula también difiere → confirmado que es otra persona.
    - "unverified": el campo de cédula está vacío en Microsoft/Canvas → no se
      pudo verificar, hay que revisarlo a mano.
    - sin cédula para comparar (colisión solo por nombre, caso general)."""
    collision_name = creds.get("name_collision_with")
    if not collision_name:
        return ""
    cedula_status = creds.get("name_collision_cedula_status")
    if cedula_status == "different":
        return f" ⚠️ Verificar antes de compartir: nombre parecido a '{collision_name}' (otra persona, cédula distinta confirmada)"
    if cedula_status == "unverified":
        return f" ⚠️ Verificar antes de compartir: nombre parecido a '{collision_name}' (no se pudo verificar por cédula, campo vacío en Microsoft/Canvas — revisar manualmente)"
    return f" ⚠️ Verificar antes de compartir: nombre parecido a '{collision_name}' (otra persona)"


def _safe_mail_nickname(name: str, suffix: str = "") -> str:
    """Genera un mailNickname válido para Microsoft Graph a partir de un nombre libre.

    Ver app.services.teams_client.safe_mail_nickname (misma lógica, centralizada
    ahí porque también la necesita el router de Teams)."""
    return graph.safe_mail_nickname(name, suffix=suffix)


def _clean_cedula(v: str) -> str:
    """Quita puntos/guiones/espacios de una cédula leída de Excel (ej. '3.517.784'
    → '3517784'), para que el formato de credenciales (cedula-Iniciales) sea
    siempre consistente sin importar cómo esté formateada la celda de origen."""
    return re.sub(r"[.\-\s]", "", v or "")


_coordinador_role_id_cache: str | None = None


async def _get_coordinador_role_id() -> str | None:
    """Busca el rol personalizado 'Coordinador' (basado en Designer) en la cuenta
    de Canvas. Se cachea en memoria porque no cambia entre requests. Devuelve
    None si el rol todavía no existe (hay que crearlo en Canvas Admin > Permisos)."""
    global _coordinador_role_id_cache
    if _coordinador_role_id_cache:
        return _coordinador_role_id_cache
    try:
        roles = await canvas.paginate(
            f"/accounts/{settings.canvas_account_id}/roles", params={"show_inactive": "false"}
        )
        for r in roles:
            if (r.get("label") or "").strip().lower() == "coordinador":
                _coordinador_role_id_cache = str(r.get("id"))
                return _coordinador_role_id_cache
    except Exception:
        pass
    return None


def _validate_file(file: UploadFile) -> None:
    """Validar tipo y tamaño del archivo."""
    # Validar tipo MIME
    if file.content_type not in ALLOWED_EXCEL_TYPES:
        logger.warning(f"Tipo MIME no permitido: {file.content_type}")
        raise HTTPException(
            status_code=400,
            detail=f"Solo archivos Excel permitidos. Recibido: {file.content_type}"
        )

    # Validar extensión
    filename = file.filename or ""
    if not filename.lower().endswith(('.xlsx', '.xls')):
        logger.warning(f"Extensión no permitida: {filename}")
        raise HTTPException(
            status_code=400,
            detail="Solo archivos .xlsx o .xls permitidos"
        )


def _read_excel(file: UploadFile) -> list[dict]:
    """Leer y normalizar archivo Excel con validaciones de seguridad."""
    _validate_file(file)

    contents = file.file.read()

    # Validar tamaño
    if len(contents) > MAX_FILE_SIZE:
        logger.warning(f"Archivo excede límite de {MAX_FILE_SIZE} bytes")
        raise HTTPException(
            status_code=413,
            detail=f"Archivo demasiado grande. Máximo: {MAX_FILE_SIZE / 1024 / 1024:.0f} MB"
        )

    try:
        df = pd.read_excel(io.BytesIO(contents), dtype=str)
    except Exception as e:
        logger.error(f"Error leyendo Excel: {e}")
        raise HTTPException(status_code=400, detail="Archivo Excel inválido")

    # Validar número de filas
    if len(df) > MAX_ROWS:
        logger.warning(f"Archivo excede {MAX_ROWS} filas")
        raise HTTPException(
            status_code=413,
            detail=f"Máximo {MAX_ROWS} filas permitidas"
        )

    df = df.where(pd.notnull(df), None)
    # Normalize column headers
    df.columns = [_norm(str(c)) for c in df.columns]
    return df.to_dict(orient="records")


def _get(row: dict, *keys) -> Any:
    """Try multiple key variants (Spanish + English) and return the first non-empty value."""
    for k in keys:
        v = row.get(k) or row.get(_norm(k))
        if v is not None and str(v).strip():
            return str(v).strip()
    return None

def _find_header_row_and_headers(ws, max_scan_rows=15):
    """
    Escanea las primeras `max_scan_rows` buscando la fila de encabezados.
    Retorna (header_row_idx, dict_de_headers_normalizados, array_de_headers) o (None, {}, []).
    """
    keywords = ["nombre", "curso", "usuario", "correo", "cedula", "cédula", "documento", "id canvas", "id teams"]
    for row_idx in range(1, min(max_scan_rows, ws.max_row + 1)):
        row_vals = [str(ws.cell(row=row_idx, column=c).value or "").strip() for c in range(1, min(50, ws.max_column + 1))]
        valid_cols = [v for v in row_vals if v]
        if len(valid_cols) >= 2:
            row_str = " ".join(valid_cols).lower()
            matches = sum(1 for keyword in keywords if keyword in row_str)
            # Una fila de título (ej. "Nombre del Diplomado" en la fila 1, junto
            # al ID del equipo) puede coincidir con una sola palabra clave por
            # casualidad. Una fila de encabezados real siempre tiene varias
            # columnas reconocibles, así que exigimos 2+ coincidencias o
            # suficientes celdas no vacías para no confundirlas.
            if matches >= 2 or (matches >= 1 and len(valid_cols) >= 4):
                headers_dict = {}
                for col_idx, val in enumerate(row_vals, 1):
                    if val:
                        headers_dict[_norm(val)] = col_idx
                return row_idx, headers_dict, row_vals
                
    return None, {}, []


def _safe_set_cell(ws, row: int, column: int, value):
    """Set a cell's value, unmerging its range first if it's part of one.

    Plantillas de OneDrive suelen tener un título en la fila 1 combinado
    (merge) a lo largo de varias columnas. Si ese merge llega a cubrir la
    columna donde el sistema necesita escribir (ej. el ID del equipo de
    Teams), openpyxl lanza AttributeError al intentar asignar `.value` en
    una MergedCell. Deshacemos el merge puntual para evitar el crash.
    """
    cell = ws.cell(row=row, column=column)
    if isinstance(cell, MergedCell):
        for merged_range in list(ws.merged_cells.ranges):
            if cell.coordinate in merged_range:
                ws.unmerge_cells(str(merged_range))
                break
        cell = ws.cell(row=row, column=column)
    cell.value = value
    return cell


_UUID_RE = re.compile(r'^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}$')


def _find_row1_title(ws, max_scan_cols: int = 30) -> str:
    """Return the diplomado/curso title stored in row 1, wherever it is.

    Distintas plantillas insertan columnas nuevas (Fecha, Telefono,
    Vendedor, etc.) antes del título, corriéndolo de la columna A hacia
    la derecha. En vez de asumir que siempre está en la columna 1,
    recorremos la fila y devolvemos la primera celda con texto que no
    sea, a su vez, el ID (GUID) del equipo de Teams.
    """
    for c in range(1, min(max_scan_cols, ws.max_column) + 1):
        v = str(ws.cell(row=1, column=c).value or "").strip()
        if v and not _UUID_RE.match(v):
            return v
    return ""


# ── Template generation ───────────────────────────────────────────────────────

_HEADER_FILL  = PatternFill("solid", fgColor="1A2035")
_EXAMPLE_FILL = PatternFill("solid", fgColor="EEF2FF")
_HEADER_FONT  = Font(bold=True, color="FFFFFF", size=10)
_EXAMPLE_FONT = Font(italic=True, color="555555", size=9)
_BORDER_SIDE  = Side(style="thin", color="C5C5C5")
_CELL_BORDER  = Border(left=_BORDER_SIDE, right=_BORDER_SIDE,
                       bottom=_BORDER_SIDE, top=_BORDER_SIDE)


def _build_template(
    headers: list[str],
    examples: list[list],
    col_widths: list[int] | None = None,
) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active

    # Header row
    ws.append(headers)
    for i, cell in enumerate(ws[1], 1):
        cell.fill   = _HEADER_FILL
        cell.font   = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = _CELL_BORDER
        ws.column_dimensions[get_column_letter(i)].width = (
            col_widths[i - 1] if col_widths and i <= len(col_widths) else 22
        )
    ws.row_dimensions[1].height = 30

    # Example rows
    for row in examples:
        ws.append(row)
        for cell in ws[ws.max_row]:
            cell.fill   = _EXAMPLE_FILL
            cell.font   = _EXAMPLE_FONT
            cell.border = _CELL_BORDER

    ws.freeze_panes = "A2"
    return wb


def _wb_response(wb: openpyxl.Workbook, filename: str) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _excel_response(rows: list[dict], filename: str) -> StreamingResponse:
    """Simple response for exported data (no special formatting)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    if rows:
        headers = list(rows[0].keys())
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in rows:
            ws.append(list(r.values()))
    return _wb_response(wb, filename)


# ═══════════════════════════════════════════════════════════════════════════════
# Canvas – Usuarios
# ═══════════════════════════════════════════════════════════════════════════════

@router.get("/excel/template/canvas-users", summary="Descargar plantilla usuarios Canvas")
async def template_canvas_users():
    wb = _build_template(
        headers=["Nombre Completo *", "Email *", "Login / Usuario *", "Contraseña",
                 "ID SIS (cédula)"],
        examples=[
            ["Karen Gonzalez", "karen.gonzalez@usil.edu.py", "karen.gonzalez",
             "6868066-Kg", "6868066"],
            ["Juan Perez", "juan.perez@usil.edu.py", "juan.perez",
             "1234567-Jp", "1234567"],
        ],
        col_widths=[28, 32, 22, 18, 18],
    )
    return _wb_response(wb, "plantilla_canvas_usuarios.xlsx")


@router.post("/excel/canvas/users", summary="Importar usuarios Canvas desde Excel")
async def import_canvas_users(file: UploadFile = File(...)) -> BulkResult:
    rows = _read_excel(file)
    result = BulkResult()

    async def _create(row: dict):
        try:
            name     = _get(row, "nombre_completo", "nombre", "name")
            email    = _get(row, "email", "correo", "email_institucional")
            login_id = _get(row, "login_usuario", "login", "login_id")
            password = _get(row, "contrasena", "password")
            sis_id   = _get(row, "id_sis_cedula", "id_sis", "sis_user_id")

            if not name or not email or not login_id:
                raise ValueError("Nombre Completo, Email y Login son obligatorios")

            payload = {
                "user": {"name": name, "short_name": name},
                "pseudonym": {"unique_id": login_id, "send_confirmation": False},
                "communication_channel": {
                    "type": "email", "address": email, "skip_confirmation": True
                },
            }
            if password:
                payload["pseudonym"]["password"] = password
            if sis_id:
                payload["pseudonym"]["sis_user_id"] = sis_id

            data = await canvas.post(f"/accounts/{_ACCOUNT}/users", payload)
            result.succeeded.append(data)
        except Exception as exc:
            result.failed.append({"input": row, "error": _err(exc)})

    await asyncio.gather(*[_create(r) for r in rows])
    return result


# ═══════════════════════════════════════════════════════════════════════════════
# Canvas – Cursos
# ═══════════════════════════════════════════════════════════════════════════════

@router.get("/excel/template/canvas-courses", summary="Descargar plantilla cursos Canvas")
async def template_canvas_courses():
    wb = _build_template(
        headers=["Nombre del Curso *", "Código del Curso *", "ID SIS Curso",
                 "Fecha Inicio (YYYY-MM-DD)", "Fecha Fin (YYYY-MM-DD)"],
        examples=[
            ["Matemáticas I – 2025", "MAT-I-2025", "MAT-I-2025",
             "2025-03-01", "2025-07-31"],
            ["Comunicación Oral – 2025", "COM-ORAL-2025", "COM-ORAL-2025",
             "2025-03-01", "2025-07-31"],
        ],
        col_widths=[32, 22, 18, 26, 26],
    )
    return _wb_response(wb, "plantilla_canvas_cursos.xlsx")


@router.post("/excel/canvas/courses", summary="Importar cursos Canvas desde Excel")
async def import_canvas_courses(file: UploadFile = File(...)) -> BulkResult:
    rows = _read_excel(file)
    result = BulkResult()

    async def _create(row: dict):
        try:
            name    = _get(row, "nombre_del_curso", "nombre", "name")
            code    = _get(row, "codigo_del_curso", "codigo", "course_code")
            sis_id  = _get(row, "id_sis_curso", "sis_course_id")
            start   = _get(row, "fecha_inicio_yyyy_mm_dd", "fecha_inicio", "start_at")
            end     = _get(row, "fecha_fin_yyyy_mm_dd", "fecha_fin", "end_at")

            if not name or not code:
                raise ValueError("Nombre del Curso y Código son obligatorios")

            payload = {"course": {
                "name": name, "course_code": code,
                "sis_course_id": sis_id or None,
                "start_at": start or None,
                "end_at": end or None,
            }}
            data = await canvas.post(f"/accounts/{_ACCOUNT}/courses", payload)
            result.succeeded.append(data)
        except Exception as exc:
            result.failed.append({"input": row, "error": _err(exc)})

    await asyncio.gather(*[_create(r) for r in rows])
    return result


# ═══════════════════════════════════════════════════════════════════════════════
# Canvas – Matrículas
# ═══════════════════════════════════════════════════════════════════════════════

@router.get("/excel/template/canvas-enrollments",
            summary="Descargar plantilla matrículas Canvas")
async def template_canvas_enrollments():
    wb = _build_template(
        headers=["ID Curso *", "ID Usuario *",
                 "Rol * (StudentEnrollment / TeacherEnrollment / TaEnrollment)",
                 "Estado (active / invited)", "Notificar (true / false)"],
        examples=[
            ["103", "958", "StudentEnrollment", "active", "false"],
            ["103", "1355", "TeacherEnrollment", "active", "false"],
        ],
        col_widths=[14, 14, 44, 24, 24],
    )
    return _wb_response(wb, "plantilla_canvas_matriculas.xlsx")


@router.post("/excel/canvas/enrollments",
             summary="Matricular usuarios Canvas desde Excel")
async def import_canvas_enrollments(file: UploadFile = File(...)) -> BulkResult:
    rows = _read_excel(file)
    result = BulkResult()

    async def _enroll(row: dict):
        try:
            course_id = _get(row, "id_curso", "course_id")
            user_id   = _get(row, "id_usuario", "user_id")
            rol       = _get(row, "rol", "type") or "StudentEnrollment"
            estado    = _get(row, "estado", "enrollment_state") or "invited"
            notif_raw = _get(row, "notificar", "notify") or ""

            if not course_id or not user_id:
                raise ValueError("ID Curso e ID Usuario son obligatorios")

            payload = {"enrollment": {
                "user_id": str(user_id),
                "type": rol,
                "enrollment_state": estado,
                "notify": notif_raw.lower() in ("1", "true", "si", "yes"),
            }}
            data = await canvas.post(f"/courses/{course_id}/enrollments", payload)
            result.succeeded.append(data)
        except Exception as exc:
            result.failed.append({"input": row, "error": _err(exc)})

    await asyncio.gather(*[_enroll(r) for r in rows])
    return result


@router.post("/excel/canvas/unenrollments",
             summary="Desmatricular usuarios Canvas desde Excel")
async def import_canvas_unenrollments(file: UploadFile = File(...)) -> BulkResult:
    rows = _read_excel(file)
    result = BulkResult()

    async def _unenroll(row: dict):
        try:
            course_id     = _get(row, "id_curso", "course_id")
            enrollment_id = _get(row, "id_matricula", "enrollment_id")
            task          = _get(row, "accion", "task") or "conclude"

            if not course_id or not enrollment_id:
                raise ValueError("ID Curso e ID Matrícula son obligatorios")

            data = await canvas.delete(
                f"/courses/{course_id}/enrollments/{enrollment_id}",
                {"task": task},
            )
            result.succeeded.append({"enrollment_id": enrollment_id, **data})
        except Exception as exc:
            result.failed.append({"input": row, "error": _err(exc)})

    await asyncio.gather(*[_unenroll(r) for r in rows])
    return result


# ═══════════════════════════════════════════════════════════════════════════════
# Teams – Usuarios Azure AD
# ═══════════════════════════════════════════════════════════════════════════════

@router.get("/excel/template/teams-users",
            summary="Descargar plantilla usuarios Azure AD")
async def template_teams_users():
    wb = _build_template(
        headers=["Nombre Completo *", "Usuario Principal (UPN) *",
                 "Alias de Correo *", "Contraseña *",
                 "País * (PY / PE / US)", "Rol * (student / teacher)",
                 "Nombre", "Apellido", "Departamento", "Cargo"],
        examples=[
            ["Karen Gonzalez", "karen.gonzalez@usil.edu.py",
             "karen_gonzalez", "6868066-Kg", "PY", "student",
             "Karen", "Gonzalez", "Sistemas", "Alumno"],
            ["Prof. Juan Perez", "juan.perez@usil.edu.py",
             "juan_perez", "1234567-Jp", "PY", "teacher",
             "Juan", "Perez", "Docentes", "Profesor"],
        ],
        col_widths=[28, 32, 22, 18, 16, 24, 18, 18, 20, 20],
    )
    return _wb_response(wb, "plantilla_teams_usuarios.xlsx")


@router.post("/excel/teams/users",
             summary="Importar usuarios Azure AD desde Excel")
async def import_teams_users(file: UploadFile = File(...)) -> BulkResult:
    rows = _read_excel(file)
    result = BulkResult()

    async def _create(row: dict):
        try:
            display = _get(row, "nombre_completo", "nombre", "display_name")
            upn     = _get(row, "usuario_principal_upn", "usuario_principal",
                           "user_principal_name")
            alias   = _get(row, "alias_de_correo", "alias", "mail_nickname")
            pwd     = _get(row, "contrasena", "password")
            country = _get(row, "pais_py_pe_us", "pais", "usage_location") or settings.usage_location
            role    = (_get(row, "rol_student_teacher", "rol", "role") or "student").lower()
            fname   = _get(row, "nombre", "given_name")
            lname   = _get(row, "apellido", "surname")
            dept    = _get(row, "departamento", "department")
            title   = _get(row, "cargo", "job_title")

            missing = [f for f, v in [("Nombre Completo", display),
                                       ("Usuario Principal", upn),
                                       ("Alias de Correo", alias),
                                       ("Contraseña", pwd)] if not v]
            if missing:
                raise ValueError(f"Campos obligatorios faltantes: {', '.join(missing)}")

            if role not in ("student", "teacher"):
                role = "student"

            payload: dict[str, Any] = {
                "displayName": display,
                "userPrincipalName": upn,
                "mailNickname": alias,
                "usageLocation": country,
                "accountEnabled": True,
                "passwordProfile": {
                    "forceChangePasswordNextSignIn": False,
                    "password": pwd,
                },
            }
            for src, dst in [(fname, "givenName"), (lname, "surname"),
                             (dept, "department"), (title, "jobTitle")]:
                if src:
                    payload[dst] = src

            data = await graph.post("/users", payload)
            sku = settings.azure_sku_teachers if role == "teacher" else settings.azure_sku_students
            await graph.assign_license(data["id"], sku)
            result.succeeded.append({**data, "role": role})
        except Exception as exc:
            result.failed.append({
                "input": {k: v for k, v in row.items() if "contrasena" not in k and "password" not in k},
                "error": str(exc),
            })

    await asyncio.gather(*[_create(r) for r in rows])
    return result


# ═══════════════════════════════════════════════════════════════════════════════
# Teams – Miembros
# ═══════════════════════════════════════════════════════════════════════════════

@router.get("/excel/template/teams-members",
            summary="Descargar plantilla miembros Teams")
async def template_teams_members():
    wb = _build_template(
        headers=["ID Equipo *", "ID Usuario (Azure AD) *",
                 "Rol * (member / owner)"],
        examples=[
            ["d5e9a352-4d7e-4e61-b965-2637856568e0",
             "46ad7584-b4f4-44ea-aa8e-2a77127fdb82", "member"],
            ["d5e9a352-4d7e-4e61-b965-2637856568e0",
             "7f3c9b21-a1d2-4e33-b421-998765432100", "owner"],
        ],
        col_widths=[40, 40, 20],
    )
    return _wb_response(wb, "plantilla_teams_miembros.xlsx")


@router.post("/excel/teams/members",
             summary="Añadir miembros a Teams desde Excel")
async def import_teams_members(file: UploadFile = File(...)) -> BulkResult:
    rows = _read_excel(file)
    result = BulkResult()

    async def _add(row: dict):
        try:
            team_id = _get(row, "id_equipo", "team_id")
            user_id = _get(row, "id_usuario_azure_ad", "id_usuario", "user_id")
            role    = _get(row, "rol_member_owner", "rol", "role") or "member"

            if not team_id or not user_id:
                raise ValueError("ID Equipo e ID Usuario son obligatorios")

            payload = {
                "@odata.type": "#microsoft.graph.aadUserConversationMember",
                "roles": ["owner"] if role.lower() == "owner" else [],
                "user@odata.bind": f"https://graph.microsoft.com/v1.0/users('{user_id}')",
            }
            data = await graph.post(f"/teams/{team_id}/members", payload)
            result.succeeded.append(data)
        except Exception as exc:
            result.failed.append({"input": row, "error": _err(exc)})

    await asyncio.gather(*[_add(r) for r in rows])
    return result


# ═══════════════════════════════════════════════════════════════════════════════
# Nuevo Ingreso
# ═══════════════════════════════════════════════════════════════════════════════

@router.get("/excel/template/ingreso",
            summary="Descargar plantilla Nuevo Ingreso")
async def template_ingreso():
    wb = _build_template(
        headers=["Nombre Completo *", "Cédula *", "Correo Personal *",
                 "Rol * (student / teacher)",
                 "Tipo Programa * (grado / mba / diplomado)",
                 "Nombre del Programa",
                 "Plataforma (both / canvas / teams)"],
        examples=[
            ["Karen Gonzalez", "6868066", "karen@gmail.com",
             "student", "grado", "", "both"],
            ["Juan Perez", "1234567", "juan@gmail.com",
             "student", "mba", "Master in Business Administration", "both"],
            ["Prof. Maria Lopez", "9999999", "maria@gmail.com",
             "teacher", "diplomado", "Diplomado en Gestión Empresarial",
             "teams"],
        ],
        col_widths=[28, 14, 28, 26, 36, 36, 30],
    )
    return _wb_response(wb, "plantilla_nuevo_ingreso.xlsx")


@router.post("/excel/ingreso",
             summary="Crear cuentas conjuntas desde lista de alumnos")
async def import_ingreso(file: UploadFile = File(...)) -> BulkResult:
    rows = _read_excel(file)
    result = BulkResult()
    _ACCOUNT_LOCAL = settings.canvas_account_id

    async def _process(row: dict):
        try:
            full_name = _get(row, "nombre_completo", "nombre", "full_name")
            cedula    = _clean_cedula(_get(row, "cedula", "ci") or "")
            p_email   = _get(row, "correo_personal", "personal_email")

            if not full_name or not cedula:
                raise ValueError("Nombre Completo y Cédula son obligatorios")

            role         = (_get(row, "rol_student_teacher", "rol", "role") or "student").lower()
            platform     = (_get(row, "plataforma_both_canvas_teams", "plataforma", "platform") or "both").lower()
            program_type = (_get(row, "tipo_programa_grado_mba_diplomado", "tipo_programa", "program_type") or "grado").lower()
            creds, status = await user_service.generate_unique_credentials(full_name, cedula, platform)
            program_name = _get(row, "nombre_del_programa", "nombre_programa", "program_name") or ""
            # login = email institucional para ambas plataformas (alumnos y docentes)
            # SIS   = cédula siempre (identificador institucional)
            login_id = creds["email"]

            entry: dict[str, Any] = {
                "student": full_name, "role": role,
                "personal_email": p_email,
                "credentials": {**creds, "login_id": login_id},
            }

            if platform in ("canvas", "both"):
                if status in ("existing_cedula", "existing_name"):
                    entry["canvas"] = {"status": "ok", "msg": "Ya existía en Canvas"}
                else:
                    try:
                        cu = await canvas.post(f"/accounts/{_ACCOUNT_LOCAL}/users", {
                            "user": {"name": creds["full_name"]},
                            "pseudonym": {
                                "unique_id": login_id, "sis_user_id": cedula,
                                "password": creds["password"], "send_confirmation": False,
                            },
                            "communication_channel": {
                                "type": "email", "address": creds["email"],
                                "skip_confirmation": True,
                            },
                        })
                        entry["canvas"] = {"status": "ok", "id": cu.get("id")}
                    except Exception as exc:
                        if "sis_id_in_use" in str(exc).lower() or "unique_id_in_use" in str(exc).lower():
                            entry["canvas"] = {"status": "ok", "msg": "Ya existía en Canvas"}
                        else:
                            entry["canvas"] = {"status": "error", "error": exc.detail if isinstance(exc, HTTPException) else str(exc)}

            if platform in ("teams", "both"):
                parts = full_name.strip().split()
                sku = settings.azure_sku_teachers if role == "teacher" else settings.azure_sku_students
                try:
                    au_payload = {
                        "displayName": creds["full_name"],
                        "givenName": parts[0],
                        "surname": " ".join(parts[1:]) if len(parts) > 1 else "",
                        "userPrincipalName": creds["email"],
                        "mailNickname": creds["login_id"].replace(".", "_"),
                        "usageLocation": settings.usage_location,
                        "accountEnabled": True,
                        "jobTitle": "Docente" if role == "teacher" else "Alumno",
                        "country": "Paraguay",
                        "passwordProfile": {
                            "forceChangePasswordNextSignIn": False,
                            "password": creds["password"],
                        },
                    }
                    if cedula:
                        au_payload["postalCode"] = cedula
                    au = await graph.post("/users", au_payload)
                    await graph.assign_license(au["id"], sku)
                    entry["teams"] = {"status": "ok", "id": au.get("id")}
                except Exception as exc:
                    err_str = str(exc).lower()
                    if "already exists" in err_str or "request_badrequest" in err_str:
                        entry["teams"] = {"status": "ok", "msg": "Ya existía en Azure AD"}
                    else:
                        entry["teams"] = {"status": "error", "error": exc.detail if isinstance(exc, HTTPException) else str(exc)}

            # El correo de credenciales no se envía aquí: es una acción aparte,
            # disparada manualmente (POST /ingreso/resend-credentials o
            # /ingreso/bulk-resend) con los datos de `entry["credentials"]`.
            result.succeeded.append(entry)
        except Exception as exc:
            result.failed.append({
                "input": {k: v for k, v in row.items() if "cedula" not in k},
                "error": str(exc),
            })

    await asyncio.gather(*[_process(r) for r in rows])
    return result


# ═══════════════════════════════════════════════════════════════════════════════
# Sync Canvas → Teams
# ═══════════════════════════════════════════════════════════════════════════════

@router.get("/excel/template/sync",
            summary="Descargar plantilla sincronización Canvas→Teams")
async def template_sync():
    wb = _build_template(
        headers=["ID Curso Canvas *", "ID Owner (Azure AD) *",
                 "Visibilidad (Private / Public)",
                 "Alias de Correo del Equipo *"],
        examples=[
            ["103", "46ad7584-b4f4-44ea-aa8e-2a77127fdb82",
             "Private", "comunicacion-oral-2025"],
            ["104", "46ad7584-b4f4-44ea-aa8e-2a77127fdb82",
             "Private", "matematicas-i-2025"],
        ],
        col_widths=[22, 40, 28, 32],
    )
    return _wb_response(wb, "plantilla_sync_canvas_teams.xlsx")


@router.post("/excel/sync/canvas-teams",
             summary="Sincronización conjunta Canvas→Teams desde Excel")
async def bulk_sync_canvas_teams(file: UploadFile = File(...)) -> BulkResult:
    rows = _read_excel(file)
    result = BulkResult()

    for row in rows:
        try:
            course_id   = _get(row, "id_curso_canvas", "canvas_course_id")
            owner_id    = _get(row, "id_owner_azure_ad", "owner_id")
            visibility  = _get(row, "visibilidad_private_public", "visibilidad", "visibility") or "Private"
            mail_nick   = _get(row, "alias_de_correo_del_equipo", "alias", "mail_nickname")

            if not course_id or not owner_id:
                raise ValueError("ID Curso Canvas e ID Owner son obligatorios")
            if not mail_nick:
                raise ValueError("Alias de Correo del Equipo es obligatorio")

            course = await canvas.get(f"/courses/{course_id}")
            team = await create_team_via_group(
                display_name=course.get("name", f"Curso {course_id}"),
                mail_nickname=mail_nick,
                description=course.get("public_description") or course.get("name", ""),
                visibility=visibility,
                owner_ids=[owner_id],
            )
            result.succeeded.append({
                "canvas_course_id": course_id,
                "team_id": team.get("id", ""),
                "team_name": team.get("displayName", ""),
            })
        except Exception as exc:
            result.failed.append({"input": row, "error": _err(exc)})

    return result


# ═══════════════════════════════════════════════════════════════════════════════
# Sync – members endpoint (used by sync.html)
# ═══════════════════════════════════════════════════════════════════════════════

class SyncMembersRequest(BaseModel):
    canvas_course_id: str
    teams_team_id: str


@router.post("/sync/canvas-to-teams",
             summary="Añadir miembros del curso Canvas al Team")
async def sync_canvas_to_teams(body: SyncMembersRequest) -> BulkResult:
    result = BulkResult()
    try:
        enrollments = await canvas.paginate(
            f"/courses/{body.canvas_course_id}/enrollments",
            {"state[]": ["active"], "per_page": 100},
        )
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc))

    async def _add(enrollment: dict):
        email = (enrollment.get("user") or {}).get("email") or \
                enrollment.get("user", {}).get("login_id")
        if not email:
            result.failed.append({
                "enrollment_id": enrollment.get("id"), "error": "Sin email"
            })
            return
        try:
            data = await graph.post(
                f"/teams/{body.teams_team_id}/members",
                {
                    "@odata.type": "#microsoft.graph.aadUserConversationMember",
                    "roles": [],
                    "user@odata.bind": f"https://graph.microsoft.com/v1.0/users('{email}')",
                },
            )
            result.succeeded.append(data)
        except Exception as exc:
            result.failed.append({"enrollment_id": enrollment.get("id"), "error": _err(exc)})

    await asyncio.gather(*[_add(e) for e in enrollments])
    return result
@router.get("/template/unified-creation", summary="Descargar plantilla Excel para creación conjunta")
async def template_unified_creation():
    import io
    import openpyxl
    from fastapi.responses import StreamingResponse

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cursos y Equipos"

    headers = [
        "Nombre del Curso", "Código del Curso", "Identificador del Propietario"
    ]
    ws.append(headers)
    for col in ws.iter_cols(min_row=1, max_row=1):
        for cell in col:
            cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
            cell.fill = openpyxl.styles.PatternFill(start_color="5A67D8", end_color="5A67D8", fill_type="solid")
            cell.alignment = openpyxl.styles.Alignment(horizontal="center")

    ws.append(["Curso de Prueba", "PRB-101", "profesor@usil.edu.py"])
    for col_letter in ["A", "B", "C"]:
        ws.column_dimensions[col_letter].width = 25

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=plantilla_creacion_unificada.xlsx"},
    )

@router.get("/template/unified-enrollment", summary="Descargar plantilla Excel para matriculación conjunta")
async def template_unified_enrollment():
    import io
    import openpyxl
    from fastapi.responses import StreamingResponse

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Matriculaciones"

    headers = [
        "Identificador de Usuario", "ID Curso Canvas", "ID Equipo Teams", "Rol"
    ]
    ws.append(headers)
    for col in ws.iter_cols(min_row=1, max_row=1):
        for cell in col:
            cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
            cell.fill = openpyxl.styles.PatternFill(start_color="5A67D8", end_color="5A67D8", fill_type="solid")
            cell.alignment = openpyxl.styles.Alignment(horizontal="center")

    ws.append(["1234567", "10203", "1a2b3c-4d5e", "student"])
    for col_letter in ["A", "B", "C", "D"]:
        ws.column_dimensions[col_letter].width = 25

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=plantilla_matriculacion_unificada.xlsx"},
    )

# ═══════════════════════════════════════════════════════════════════════════════
# Diplomados (Lectura y Escritura de Planilla Original)
# ═══════════════════════════════════════════════════════════════════════════════

class DiplomadosUrlRequest(BaseModel):
    url: str
    sheet_name: str
    delete_account: bool = False
    cc: list[str] = []
    report_url: str | None = None



class MatriculacionesPreviewResponse(BaseModel):
    headers: list[str]
    sample_rows: list[dict]

class UrlOnlyRequest(BaseModel):
    url: str

class JsonDataRequest(BaseModel):
    data: list[dict]
class DocentesPreviewResponse(BaseModel):
    total_rows: int
    valid_rows: int
    headers: list[str] = []
    sample_rows: list[dict] = []
    already_processed_details: list[dict] = []

class PreviewResponse(BaseModel):
    sheet_name: str
    students_to_process: int
    students_already_processed: int
    headers: list[str] = []
    sample_rows: list[dict] = []
    student_details: list[dict] = []
    already_processed_details: list[dict] = []


class SendCredentialsPreviewResponse(BaseModel):
    """Vista previa de a quién le falta enviar credenciales — no envía nada,
    solo lee la planilla y aplica la misma lógica de filtrado que el envío
    real (Correo Enviado / Enviado legacy / cuenta creada / correo válido)."""
    pending: list[dict] = []
    already_sent_count: int = 0
    no_account_count: int = 0
    no_email_count: int = 0
    program_type: str = "diplomado"

def _encode_share_url(url: str) -> str:
    import base64
    b64 = base64.b64encode(url.encode()).decode()
    b64 = b64.replace("/", "_").replace("+", "-").rstrip("=")
    return f"u!{b64}"


@router.post("/excel/diplomados/sheets", response_model=list[str])
async def get_diplomados_sheets(req: UrlOnlyRequest) -> list[str]:
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL invalida.")
    
    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo. {e}")

    try:
        import io
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
        sheets = wb.sheetnames
        wb.close()
        return sheets
    except Exception as e:
        raise HTTPException(status_code=400, detail="El archivo no es un Excel válido.")

@router.post("/excel/egreso/sheets", response_model=list[str])
async def get_egreso_sheets(req: UrlOnlyRequest) -> list[str]:
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL invalida.")
    
    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo. {e}")

    try:
        import io
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
        sheets = wb.sheetnames
        wb.close()
        return sheets
    except Exception as e:
        raise HTTPException(status_code=400, detail="El archivo no es un Excel válido.")
@router.post("/excel/diplomados/preview", summary="Pre-visualizar planilla de Diplomados")
async def preview_diplomados_onedrive(req: DiplomadosUrlRequest) -> PreviewResponse:
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")
    
    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail="El archivo no es un Excel válido.")

    if req.sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"La pestaña '{req.sheet_name}' no existe. Disponibles: {', '.join(wb.sheetnames)}")

    ws = wb[req.sheet_name]

    sample_rows = []
    header_row_idx, headers_dict, headers_raw = _find_header_row_and_headers(ws)
            
    if not header_row_idx:
        raise HTTPException(status_code=400, detail="No se encontró la fila de encabezados.")

    headers = [h for h in headers_raw if h] # Just for returning to frontend

    col_estado = -1
    col_nombre = -1
    col_cedula = -1
    col_usuario = -1
    for i, h in enumerate(headers_raw):
        if not h: continue
        h_lower = h.lower()
        if "estado" in h_lower or ("canvas" in h_lower and "id" in h_lower):
            col_estado = i
        if "nombre" in h_lower:
            if col_nombre == -1: col_nombre = i
        if "cedula" in h_lower or "cédula" in h_lower or "ci" in h_lower:
            if col_cedula == -1: col_cedula = i
        if "usuario" in h_lower:
            if col_usuario == -1: col_usuario = i
    # Si no hay columna "Estado" propiamente dicha, "Enviado" cumple el mismo rol.
    if col_estado == -1:
        for i, h in enumerate(headers_raw):
            if h and "enviado" in h.lower():
                col_estado = i
                break

    students_to_process = 0
    students_already_processed = 0
    student_details = []
    already_processed_details = []

    for row_idx in range(header_row_idx + 1, ws.max_row + 1):
        row_vals = [str(ws.cell(row=row_idx, column=c).value or "").strip() for c in range(1, len(headers_raw) + 1)]

        # Check if the row has at least a name
        if not any(row_vals):
            continue

        nombre_val = row_vals[col_nombre] if col_nombre >= 0 else ""
        cedula_val = row_vals[col_cedula] if col_cedula >= 0 else ""
        if not nombre_val or not cedula_val:
            continue

        estado_val = row_vals[col_estado] if col_estado >= 0 else ""
        estado_lower = estado_val.lower()
        usuario_val = row_vals[col_usuario] if col_usuario >= 0 else ""
        # Misma condición que usa el procesamiento real (import_diplomados_onedrive):
        # respeta también el caso de "Estado vacío pero Usuario ya tiene @".
        if ("✅" in estado_val or estado_lower in ["creado", "ok", "si", "yes", "true", "enviado"]
                or "creado ok" in estado_lower or "ya exist" in estado_lower
                or (usuario_val and "@" in usuario_val)):
            students_already_processed += 1
            already_processed_details.append({
                "nombre": nombre_val,
                "cedula": cedula_val,
                "usuario": usuario_val,
                "estado": estado_val,
            })
        else:
            students_to_process += 1
            student_details.append({
                "nombre": nombre_val,
                "cedula": cedula_val,
            })

        if len(sample_rows) < 10:
            sample_rows.append({h: v for h, v in zip(headers_raw, row_vals) if h})

    wb.close()
    return PreviewResponse(
        sheet_name=req.sheet_name,
        students_to_process=students_to_process,
        students_already_processed=students_already_processed,
        headers=headers,
        sample_rows=sample_rows,
        student_details=student_details,
        already_processed_details=already_processed_details,
    )


@router.post("/excel/diplomados", summary="Procesar planilla de Diplomados directo en OneDrive")
async def import_diplomados_onedrive(req: DiplomadosUrlRequest, bg_tasks: BackgroundTasks) -> dict:
    """Dispara la creación de Diplomados en segundo plano.

    Antes era una request HTTP síncrona con un límite artificial de 50
    cuentas nuevas por corrida — no por una regla de negocio, sino porque
    crear usuarios en Azure AD + matricularlos en Teams uno por uno es
    demasiado lento para completarse dentro del timeout de una request
    síncrona a partir de cierto volumen (mismo motivo por el que Creación
    de Cursos y Matriculaciones ya corren en segundo plano). Al mover esto
    a un job, el límite deja de ser necesario por timeout y queda solo
    como resguardo de seguridad contra cargas accidentalmente enormes."""
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")

    encoded_url = _encode_share_url(req.url)

    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo de OneDrive. Verifica la URL y los permisos. Detalle: {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
        if req.sheet_name not in wb.sheetnames:
            raise HTTPException(status_code=400, detail=f"La pestaña '{req.sheet_name}' no existe en el archivo. Las disponibles son: {', '.join(wb.sheetnames)}")
        wb.close()
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=400, detail="El archivo descargado no es un Excel válido.")

    job_id = await jobs.create_job(job_type="diplomados_onedrive", operation="Alta de Diplomados", username="admin")
    bg_tasks.add_task(_process_diplomados_bg, job_id, req, contents, encoded_url)
    return {"status": "success", "job_id": job_id, "message": "Alta de Diplomados iniciada en segundo plano."}


async def _process_diplomados_bg(job_id: int, req: DiplomadosUrlRequest, contents: bytes, encoded_url: str):
    import json
    await jobs.start_job(job_id)

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents))
    except Exception as e:
        await jobs.fail_job(job_id, f"El archivo no es un Excel válido: {e}")
        return

    _ACCOUNT_LOCAL = settings.canvas_account_id
    result = BulkResult()

    if req.sheet_name not in wb.sheetnames:
        await jobs.fail_job(job_id, f"La pestaña '{req.sheet_name}' no existe en el archivo. Las disponibles son: {', '.join(wb.sheetnames)}")
        return

    # Buscar en la hoja especificada
    for sheet_name in [req.sheet_name]:
        ws = wb[sheet_name]
        
        header_row_idx, headers, _ = _find_header_row_and_headers(ws)

        
        if not header_row_idx:
            continue

        def get_col_idx(*keys):
            return _match_col_idx(headers, *keys)

        col_nombre = get_col_idx("nombre", "alumno", "estudiante")
        col_cedula = get_col_idx("cedula", "cédula", "ci")
        col_correo = get_col_idx("correo")
        col_telefono = get_col_idx("telefono", "teléfono", "celular")
        col_curso = get_col_idx("curso", "id curso", "canvas")
        col_equipo = get_col_idx("equipo", "id equipo", "teams")
        col_curso_nombre = get_col_idx("nombre del curso", "curso", "diplomado")
        
        col_usuario = get_col_idx("usuario")
        col_contra = get_col_idx("contrasena", "contraseña", "clave")
        col_enviado = get_col_idx("estado", "enviado")

        title_val_early = _find_row1_title(ws)
        is_mba_early = "mba" in req.sheet_name.lower() or "mba" in (title_val_early or "").lower()

        col_cc = get_col_idx("cc", "copia")
        sheet_cc_list = []
        if col_cc:
            for r_idx in range(header_row_idx + 1, ws.max_row + 1):
                cc_val = str(ws.cell(row=r_idx, column=col_cc).value or "").strip()
                if cc_val:
                    for email in cc_val.replace(";", ",").replace("\n", ",").split(","):
                        email = email.strip()
                        if "@" in email and email not in sheet_cc_list:
                            sheet_cc_list.append(email)

        if not col_nombre or not col_cedula:
            detected = list(headers.keys()) if headers else ['Ninguna']
            await jobs.fail_job(job_id, f"No se encontraron las columnas requeridas ('Nombre' y 'Cédula') en la pestaña seleccionada. Columnas detectadas: {detected}")
            return

        next_col = ws.max_column + 1

        # MBA además crea la cuenta en Canvas (a diferencia de Diplomados,
        # que hoy es solo Teams) — si la planilla no trae una columna de
        # curso de Canvas, se crea una para dejar el ID asentado.
        if is_mba_early and not col_curso:
            col_curso = next_col
            _safe_set_cell(ws, header_row_idx, col_curso, "ID Curso Canvas").font = Font(bold=True)
            next_col += 1

        if not col_usuario:
            col_usuario = next_col
            _safe_set_cell(ws, header_row_idx, col_usuario, "Usuario").font = Font(bold=True)
            next_col += 1
        if not col_contra:
            col_contra = next_col
            _safe_set_cell(ws, header_row_idx, col_contra, "Contraseña").font = Font(bold=True)
            next_col += 1
        if not col_enviado:
            col_enviado = next_col
            _safe_set_cell(ws, header_row_idx, col_enviado, "Enviado").font = Font(bold=True)

        # Si la planilla tiene una columna "Enviado" propia del usuario,
        # separada de "Estado" (que es la que el sistema prioriza para
        # decidir si ya se procesó una fila), la mantenemos sincronizada
        # con el mismo estado en vez de dejarla vacía/manual.
        col_enviado_mirror = None
        for h, idx in headers.items():
            if idx != col_enviado and "enviado" in h:
                col_enviado_mirror = idx
                break

        def _set_estado(r_idx, value, color, mirror=False):
            """Escribe el estado de creación de cuenta en la columna Estado.

            `mirror=True` además replica el valor en la columna "Enviado"
            preexistente — SOLO corresponde para sincronizar una fila que ya
            estaba creada de antes (Estado ya OK pero Enviado en blanco), no
            para cuentas recién creadas en esta misma corrida: "Enviado"
            representa si el correo de credenciales fue enviado, algo que
            hoy es una acción manual y separada (botón "Enviar
            Credenciales"), no un efecto automático de la creación de la
            cuenta.
            """
            cell = ws.cell(row=r_idx, column=col_enviado, value=value)
            cell.font = Font(color=color, bold=True)
            if mirror and col_enviado_mirror:
                mirror_cell = ws.cell(row=r_idx, column=col_enviado_mirror, value=value)
                mirror_cell.font = Font(color=color, bold=True)

        global_team_id = ""
        title_val = _find_row1_title(ws)
        global_team_name = title_val

        # Mismo criterio que en los endpoints de envío de credenciales: la
        # pestaña/hora de MBA se identifica por su nombre o por el título de
        # la fila 1 — se usa para sumar al propietario adicional de MBA en
        # cualquier Team que se cree desde esta pestaña.
        is_mba = "mba" in req.sheet_name.lower() or "mba" in (title_val or "").lower()
        team_owner_ids = list(_DIPLOMADO_DEFAULT_OWNERS) + ([_MBA_EXTRA_OWNER] if is_mba else [])
        if global_team_name and col_usuario:
            team_id_from_header = str(ws.cell(row=1, column=col_usuario).value or "").strip()
            if _UUID_RE.match(team_id_from_header):
                global_team_id = team_id_from_header
            else:
                try:
                    existing_tid = await graph.search_group_by_name(global_team_name)
                    if existing_tid:
                        global_team_id = existing_tid
                    else:
                        nickname = _safe_mail_nickname(global_team_name)

                        owner_ids = list(team_owner_ids)
                        try:
                            admin_user = await graph.get(f"/users/resteche@usil.edu.py", params={"$select": "id"})
                            if admin_user and admin_user.get("id"):
                                owner_ids.append(admin_user["id"])
                        except: pass

                        new_team = await graph.create_team_via_group(
                            display_name=global_team_name,
                            mail_nickname=nickname,
                            description=f"Grupo para {global_team_name}",
                            visibility="Private",
                            owner_ids=owner_ids
                        )
                        global_team_id = new_team.get("id", "")
                    if global_team_id:
                        _safe_set_cell(ws, 1, col_usuario, global_team_id).font = Font(bold=True)
                except Exception as e:
                    print(f"Error pre-creando equipo: {e}")

        # Curso de Canvas global para MBA: mismo criterio que el Team de
        # arriba — se crea (o reutiliza) una sola vez por pestaña, a partir
        # del título de la fila 1, y el ID queda asentado en la fila 1 de la
        # columna "ID Curso Canvas" para no recrearlo en corridas futuras.
        global_course_id = ""
        if is_mba and global_team_name and col_curso:
            course_id_from_header = str(ws.cell(row=1, column=col_curso).value or "").strip()
            if course_id_from_header.isdigit():
                global_course_id = course_id_from_header
            else:
                try:
                    existing_cid = await canvas.search_course_by_name(_ACCOUNT_LOCAL, global_team_name)
                    if existing_cid:
                        global_course_id = existing_cid
                    else:
                        global_course_id = await canvas.create_course(_ACCOUNT_LOCAL, global_team_name)
                    if global_course_id:
                        _safe_set_cell(ws, 1, col_curso, global_course_id).font = Font(bold=True)
                except Exception as e:
                    print(f"Error pre-creando curso Canvas (MBA): {e}")

        async def process_row(r_idx):
            nombre = str(ws.cell(row=r_idx, column=col_nombre).value or "").strip()
            cedula = _clean_cedula(str(ws.cell(row=r_idx, column=col_cedula).value or "").strip())
            correo = str(ws.cell(row=r_idx, column=col_correo).value or "").strip() if col_correo else ""
            telefono = str(ws.cell(row=r_idx, column=col_telefono).value or "").strip() if col_telefono else ""
            id_curso = str(ws.cell(row=r_idx, column=col_curso).value or "").strip() if col_curso else ""
            id_equipo = str(ws.cell(row=r_idx, column=col_equipo).value or "").strip() if col_equipo else ""
            curso_nombre = str(ws.cell(row=r_idx, column=col_curso_nombre).value or "").strip() if col_curso_nombre else ""

            enviado = str(ws.cell(row=r_idx, column=col_enviado).value or "").strip()
            
            if not nombre or not cedula or cedula == "None":
                return
            usuario_val = str(ws.cell(row=r_idx, column=col_usuario).value or "").strip() if col_usuario else ""
            enviado_lower = enviado.lower()
            if "✅" in enviado or enviado_lower in ["si", "yes", "true", "enviado", "ok"] or "creado ok" in enviado_lower or "ya exist" in enviado_lower or (usuario_val and "@" in usuario_val):
                if not enviado and col_enviado:
                    _set_estado(r_idx, "✅ Existente", "00B050", mirror=True)
                elif col_enviado_mirror:
                    mirror_val = str(ws.cell(row=r_idx, column=col_enviado_mirror).value or "").strip()
                    if not mirror_val:
                        _set_estado(r_idx, enviado, "00B050", mirror=True)
                # Alumno ya creado en una corrida anterior: igual verificamos que
                # esté agregado al grupo de Teams correspondiente (como miembro),
                # y lo agregamos si falta — un run previo pudo haber creado la
                # cuenta pero fallado al matricularla en el grupo.
                if usuario_val and "@" in usuario_val:
                    # Reparo de contraseña: la cuenta ya existe (hay Usuario)
                    # pero la celda de Contraseña está vacía o quedó mal
                    # generada (ej. por el bug de detección de columna ya
                    # corregido) — se recalcula con la cédula de ESTA fila,
                    # sin tocar la cuenta en Canvas/Teams (no hace falta,
                    # solo había que dejar asentada la contraseña real).
                    if col_contra and cedula:
                        contra_val = str(ws.cell(row=r_idx, column=col_contra).value or "").strip()
                        if not contra_val:
                            ws.cell(row=r_idx, column=col_contra, value=generate_password(cedula, nombre))
                    try:
                        target_equipo = id_equipo if (id_equipo and id_equipo != "None") else global_team_id
                        if not target_equipo and curso_nombre:
                            target_equipo = await graph.search_group_by_name(curso_nombre)
                        if target_equipo:
                            user_data = await graph.get(f"/users/{usuario_val}", params={"$select": "id"})
                            uid = user_data.get("id") if user_data else None
                            if uid:
                                is_member = True
                                try:
                                    await graph.get(f"/groups/{target_equipo}/members/{uid}", params={"$select": "id"})
                                except Exception:
                                    is_member = False
                                if not is_member:
                                    await graph.add_member_to_group(target_equipo, uid)
                    except Exception:
                        pass
                    # Mismo auto-reparo para Canvas: un run previo pudo haber
                    # creado la cuenta de Teams pero fallado al crear la de
                    # Canvas o al matricularla en el curso.
                    if is_mba and global_course_id:
                        try:
                            canvas_match = await canvas.find_user_exact(_ACCOUNT_LOCAL, usuario_val)
                            if canvas_match:
                                try:
                                    await canvas.post(f"/courses/{global_course_id}/enrollments", {
                                        "enrollment": {
                                            "user_id": canvas_match["id"],
                                            "type": "StudentEnrollment",
                                            "enrollment_state": "active",
                                            "notify": False,
                                        },
                                    })
                                except Exception:
                                    pass
                        except Exception:
                            pass
                return

            # platform="both" (no solo "teams"): si la persona ya tiene cuenta
            # en Canvas (de este mismo flujo de MBA, de Alta Docentes, de
            # Ingreso, etc.) se detecta por CÉDULA -el chequeo más confiable,
            # porque es el SIS ID exacto- y se reutiliza su correo real, en
            # vez de intentar adivinarlo de nuevo y terminar creando una
            # cuenta duplicada cuando el correo real tiene otro formato
            # (ej. por una colisión de nombre resuelta distinto la vez
            # anterior).
            creds, status = await user_service.generate_unique_credentials(nombre, cedula, "both")
            login_id = creds["email"]
            pwd = creds["password"]
            error = None
            
            au_id = None
            
            if not error:
                parts = creds["full_name"].strip().split()
                try:
                    au_payload = {
                        "displayName": creds["full_name"],
                        "givenName": parts[0],
                        "surname": " ".join(parts[1:]) if len(parts) > 1 else "",
                        "userPrincipalName": login_id,
                        "mailNickname": login_id.replace(".", "_").replace("@", "_"),
                        "usageLocation": settings.usage_location,
                        "accountEnabled": True,
                        "jobTitle": "Alumno",
                        "department": "UBS",
                        "country": "Paraguay",
                        "passwordProfile": {
                            "forceChangePasswordNextSignIn": False,
                            "password": pwd,
                        },
                    }
                    if cedula:
                        au_payload["postalCode"] = cedula
                    if telefono:
                        au_payload["mobilePhone"] = telefono
                    au = await graph.post("/users", au_payload)
                    au_id = au.get("id")
                    await graph.assign_license(au_id, settings.azure_sku_students)
                except Exception as e:
                    if "already exists" not in str(e).lower() and "Request_BadRequest" not in str(e):
                        error = str(e)
                    else:
                        error = "Ya existía en Azure AD" 
            
            if not error or "Ya existía" in str(error):
                try:
                    target_equipo = global_team_id
                    if id_equipo and id_equipo != "None":
                        target_equipo = id_equipo
                    elif not target_equipo and curso_nombre:
                        existing_tid = await graph.search_group_by_name(curso_nombre)
                        if existing_tid:
                            target_equipo = existing_tid
                        else:
                            nickname = _safe_mail_nickname(curso_nombre)
                            new_team = await graph.create_team_via_group(
                                display_name=curso_nombre,
                                mail_nickname=nickname,
                                description=f"Grupo para {curso_nombre}",
                                visibility="Private",
                                owner_ids=list(team_owner_ids)
                            )
                            target_equipo = new_team.get("id")

                    if target_equipo:
                        if col_equipo and target_equipo != id_equipo:
                            ws.cell(row=r_idx, column=col_equipo, value=target_equipo)
                        
                        uid = au_id
                        if not uid:
                            try:
                                user_data = await graph.get(f"/users/{login_id}", params={"$select": "id"})
                                if user_data and user_data.get("id"):
                                    uid = user_data["id"]
                            except Exception as get_err:
                                if "404" in str(get_err) and not error:
                                    pass
                                elif "404" in str(get_err) and "Ya existía" in str(error):
                                    error = f"{error} pero se encuentra eliminado (Papelera de Azure AD)"
                                else:
                                    raise get_err
                                    
                        if uid:
                            await graph.post(f"/groups/{target_equipo}/members/$ref", {"@odata.id": f"https://graph.microsoft.com/v1.0/directoryObjects/{uid}"})
                        else:
                            if not error:
                                error = "TeamsEnroll: No se pudo obtener el ID del usuario."
                except Exception as e:
                    if "already exist" not in str(e).lower():
                        error = str(error) + f" | TeamsEnroll: {e}" if error else f"TeamsEnroll: {e}"

            # MBA: a diferencia de Diplomados (Teams-only), acá además se da
            # de alta al alumno en Canvas y se lo matricula como estudiante
            # en el curso global de la pestaña.
            if is_mba and global_course_id:
                canvas_id = None
                try:
                    cu = await canvas.post(f"/accounts/{_ACCOUNT_LOCAL}/users", {
                        "user": {
                            "name": creds["full_name"],
                            "sortable_name": creds["full_name"],
                            "short_name": parts[0] + " " + parts[-1] if len(parts) > 1 else creds["full_name"],
                        },
                        "pseudonym": {
                            "unique_id": login_id,
                            "sis_user_id": cedula,
                            "password": pwd,
                            "send_confirmation": False,
                        },
                        "communication_channel": {
                            "type": "email", "address": login_id,
                            "skip_confirmation": True,
                        },
                    })
                    canvas_id = cu["id"]
                except Exception as e:
                    if "already exists" not in str(e).lower() and "sis_user_id" not in str(e).lower():
                        error = str(error) + f" | Canvas: {e}" if error else f"Canvas: {e}"
                    else:
                        try:
                            match = await canvas.find_user_exact(_ACCOUNT_LOCAL, login_id)
                            if match:
                                canvas_id = match["id"]
                        except Exception:
                            pass
                        if not canvas_id:
                            error = str(error) + f" | Canvas: {e}" if error else f"Canvas: {e}"

                if canvas_id:
                    try:
                        await canvas.post(f"/courses/{global_course_id}/enrollments", {
                            "enrollment": {
                                "user_id": canvas_id,
                                "type": "StudentEnrollment",
                                "enrollment_state": "active",
                                "notify": False,
                            },
                        })
                    except Exception as e:
                        if "already" not in str(e).lower():
                            error = str(error) + f" | CanvasEnroll: {e}" if error else f"CanvasEnroll: {e}"

            if not error or "Creado OK" in str(error) or "Ya existía" in str(error):
                ws.cell(row=r_idx, column=col_usuario, value=login_id)
                ws.cell(row=r_idx, column=col_contra, value=pwd)
                result.succeeded.append({"cedula": cedula, "nombre": creds["full_name"], "login_id": login_id})

                collision_note = _collision_note(creds)

                if not error:
                    _set_estado(r_idx, f"✅ OK{collision_note}", "00B050" if not collision_note else "D97706")
                else:
                    _set_estado(r_idx, f"⚠️ {error}{collision_note}", "D97706")
            else:
                _set_estado(r_idx, f"❌ Error: {error}", "FF0000")
                result.failed.append({"input": {"cedula": cedula}, "error": error})

        tasks = []
        pending_count = 0
        empty_count = 0
        for r_idx in range(header_row_idx + 1, ws.max_row + 1):
            nombre_val = str(ws.cell(row=r_idx, column=col_nombre).value or "").strip()
            cedula_val = str(ws.cell(row=r_idx, column=col_cedula).value or "").strip()

            if not nombre_val and not cedula_val:
                empty_count += 1
                if empty_count > 10:
                    break
                continue

            empty_count = 0

            # El límite de seguridad es sobre cuentas NUEVAS a crear, no sobre
            # el total de filas — si no, una planilla con la mayoría de filas
            # ya procesadas se bloquea aunque falten pocas por crear.
            enviado_check = str(ws.cell(row=r_idx, column=col_enviado).value or "").strip()
            enviado_check_lower = enviado_check.lower()
            usuario_check = str(ws.cell(row=r_idx, column=col_usuario).value or "").strip() if col_usuario else ""
            already_done = (
                "✅" in enviado_check or enviado_check_lower in ["si", "yes", "true", "enviado", "ok"]
                or "creado ok" in enviado_check_lower or "ya exist" in enviado_check_lower
                or (usuario_check and "@" in usuario_check)
            )
            if not already_done:
                pending_count += 1

            tasks.append(process_row(r_idx))

        if pending_count > 300:
            await jobs.fail_job(job_id, f"Límite de seguridad excedido: intentás crear {pending_count} cuentas nuevas a la vez (máximo 300 permitidas). Revisa el archivo para evitar accidentes.")
            return

        if len(tasks) > 0:
            try:
                # Verificar si el archivo está bloqueado antes de empezar a procesar alumnos
                await graph.put_raw(f"/shares/{encoded_url}/driveItem/content", contents)
            except Exception as e:
                await jobs.fail_job(job_id, f"El archivo Excel está abierto o el enlace es de Solo Lectura. Detalle real: {e}")
                return

        batch_size = 5
        for i in range(0, len(tasks), batch_size):
            await asyncio.gather(*tasks[i:i+batch_size])
            await jobs.update_job_progress(
                job_id, len(result.succeeded), len(result.failed),
                data_json=json.dumps({
                    "total_to_process": len(tasks),
                    "processed": len(result.succeeded) + len(result.failed),
                    "results": result.succeeded + result.failed,
                }),
            )

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    try:
        await graph.put_raw(f"/shares/{encoded_url}/driveItem/content", output.getvalue())
    except Exception as e:
        await jobs.complete_job(job_id, len(result.succeeded), len(result.failed), f"Procesado, pero no se pudo guardar en OneDrive: {e}")
        return

    if not result.succeeded and not result.failed:
        await jobs.complete_job(job_id, 0, 0, "No se procesó ninguna fila. Todas las filas ya estaban marcadas como procesadas, tenían correo asignado, o la planilla estaba vacía.")
        return

    await jobs.update_job_progress(
        job_id, len(result.succeeded), len(result.failed),
        data_json=json.dumps({
            "total_to_process": len(result.succeeded) + len(result.failed),
            "processed": len(result.succeeded) + len(result.failed),
            "results": result.succeeded + result.failed,
        }),
    )
    await jobs.complete_job(job_id, len(result.succeeded), len(result.failed), "Guardado en OneDrive correctamente.")


@router.post("/excel/diplomados/send-credentials/preview", summary="Vista previa: quién falta enviar credenciales (Diplomados)")
async def preview_send_diplomados_credentials(req: DiplomadosUrlRequest) -> SendCredentialsPreviewResponse:
    """Muestra quién está pendiente de recibir el correo de credenciales, sin
    enviar nada — misma lógica de filtrado exacta que /send-credentials."""
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")

    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo de OneDrive. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail="El archivo descargado no es un Excel válido.")

    if req.sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"La pestaña '{req.sheet_name}' no existe.")

    ws = wb[req.sheet_name]
    header_row_idx, headers, _ = _find_header_row_and_headers(ws)
    if not header_row_idx:
        raise HTTPException(status_code=400, detail="No se encontró la fila de encabezados.")

    def get_col_idx(*keys):
        return _match_col_idx(headers, *keys)

    col_nombre = get_col_idx("nombre", "alumno", "estudiante")
    col_correo = get_col_idx("correo")
    col_usuario = get_col_idx("usuario")
    col_contra = get_col_idx("contrasena", "contraseña", "clave")

    if not col_usuario or not col_contra:
        raise HTTPException(
            status_code=400,
            detail="No se encontraron las columnas de Usuario/Contraseña. Primero procesa la planilla con 'Carga Diplomados'.",
        )
    if not col_correo:
        raise HTTPException(status_code=400, detail="No se encontró una columna de Correo para enviar las credenciales.")

    col_correo_enviado = get_col_idx("correo enviado", "credenciales enviadas")
    col_enviado_legacy = headers.get("enviado")

    title_val = _find_row1_title(ws)
    is_mba = "mba" in req.sheet_name.lower() or "mba" in (title_val or "").lower()

    def _looks_sent(raw: str) -> bool:
        v = raw.strip().lower()
        return bool(v) and ("✅" in raw or v in ("si", "yes", "true", "enviado"))

    pending, already_sent, no_account, no_email = [], 0, 0, 0
    for r_idx in range(header_row_idx + 1, ws.max_row + 1):
        nombre_val = str(ws.cell(row=r_idx, column=col_nombre).value or "").strip() if col_nombre else ""
        usuario_val = str(ws.cell(row=r_idx, column=col_usuario).value or "").strip()
        contra_val = str(ws.cell(row=r_idx, column=col_contra).value or "").strip()
        correo_val = str(ws.cell(row=r_idx, column=col_correo).value or "").strip()
        if not nombre_val and not usuario_val and not correo_val:
            continue

        ya_enviado = str(ws.cell(row=r_idx, column=col_correo_enviado).value or "") if col_correo_enviado else ""
        ya_enviado_legacy = str(ws.cell(row=r_idx, column=col_enviado_legacy).value or "") if col_enviado_legacy else ""

        if not usuario_val or usuario_val == "None" or not contra_val or contra_val == "None":
            no_account += 1
            continue
        if not correo_val or "@" not in correo_val:
            no_email += 1
            continue
        if _looks_sent(ya_enviado) or _looks_sent(ya_enviado_legacy):
            already_sent += 1
            continue

        pending.append({"nombre": nombre_val, "correo": correo_val, "usuario": usuario_val})

    wb.close()
    return SendCredentialsPreviewResponse(
        pending=pending, already_sent_count=already_sent,
        no_account_count=no_account, no_email_count=no_email,
        program_type="mba" if is_mba else "diplomado",
    )


@router.post("/excel/diplomados/send-credentials", summary="Enviar correo de credenciales a alumnos ya creados (Diplomados)")
async def send_diplomados_credentials(req: DiplomadosUrlRequest) -> BulkResult:
    """Envía el correo de credenciales a los alumnos que ya tienen cuenta creada
    (Usuario + Contraseña ya generados) y que todavía no la recibieron.

    Acción separada de la creación de cuentas: permite reintentar el envío sin
    volver a crear nada, y no reenvía a quien ya fue marcado como enviado.
    """
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")

    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo de OneDrive. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents))
    except Exception as e:
        raise HTTPException(status_code=400, detail="El archivo descargado no es un Excel válido.")

    result = BulkResult()
    if req.sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"La pestaña '{req.sheet_name}' no existe.")

    ws = wb[req.sheet_name]
    header_row_idx, headers, _ = _find_header_row_and_headers(ws)
    if not header_row_idx:
        raise HTTPException(status_code=400, detail="No se encontró la fila de encabezados.")

    def get_col_idx(*keys):
        return _match_col_idx(headers, *keys)

    col_nombre = get_col_idx("nombre", "alumno", "estudiante")
    col_correo = get_col_idx("correo")
    col_usuario = get_col_idx("usuario")
    col_contra = get_col_idx("contrasena", "contraseña", "clave")
    col_curso_nombre = get_col_idx("nombre del curso", "curso", "diplomado")

    if not col_usuario or not col_contra:
        raise HTTPException(
            status_code=400,
            detail="No se encontraron las columnas de Usuario/Contraseña. Primero procesa la planilla con 'Carga Diplomados'.",
        )
    if not col_correo:
        raise HTTPException(status_code=400, detail="No se encontró una columna de Correo para enviar las credenciales.")

    # Columna dedicada para marcar el envío, separada de "Estado" (que ya usa
    # el procesamiento de creación de cuentas).
    col_correo_enviado = get_col_idx("correo enviado", "credenciales enviadas")
    if not col_correo_enviado:
        col_correo_enviado = ws.max_column + 1
        _safe_set_cell(ws, header_row_idx, col_correo_enviado, "Correo Enviado").font = Font(bold=True)

    # Muchas planillas ya traen su propia columna "Enviado" (marcada a mano o
    # por un proceso externo, antes de que existiera este sistema). Si existe
    # y ya dice que se envió, la respetamos como señal de "ya enviado" además
    # de nuestra propia columna "Correo Enviado" — para no reenviar algo que
    # ya salió por otra vía.
    col_enviado_legacy = headers.get("enviado")

    title_val = _find_row1_title(ws)

    # El mismo enlace de "Diplomados" se usa también para MBA (pestaña/hora
    # distinta del mismo Excel) — se distingue por el nombre de la pestaña o
    # el título de la fila 1 ("MBA DUAL 2026", "MBA 2026-01", etc.) para
    # elegir la plantilla de correo correcta (ver email_service._build_mba_message).
    is_mba = "mba" in req.sheet_name.lower() or "mba" in (title_val or "").lower()
    program_type = "mba" if is_mba else "diplomado"

    col_cc = get_col_idx("cc", "copia")
    sheet_cc_list: list[str] = []
    if col_cc:
        for r_idx in range(header_row_idx + 1, ws.max_row + 1):
            cc_val = str(ws.cell(row=r_idx, column=col_cc).value or "").strip()
            if cc_val:
                for email in cc_val.replace(";", ",").replace("\n", ",").split(","):
                    email = email.strip()
                    if "@" in email and email not in sheet_cc_list:
                        sheet_cc_list.append(email)

    def _looks_sent(raw: str) -> bool:
        v = raw.strip().lower()
        return bool(v) and ("✅" in raw or v in ("si", "yes", "true", "enviado"))

    rows_to_send = []
    for r_idx in range(header_row_idx + 1, ws.max_row + 1):
        usuario_val = str(ws.cell(row=r_idx, column=col_usuario).value or "").strip()
        contra_val = str(ws.cell(row=r_idx, column=col_contra).value or "").strip()
        correo_val = str(ws.cell(row=r_idx, column=col_correo).value or "").strip()
        ya_enviado = str(ws.cell(row=r_idx, column=col_correo_enviado).value or "")
        ya_enviado_legacy = str(ws.cell(row=r_idx, column=col_enviado_legacy).value or "") if col_enviado_legacy else ""

        if not usuario_val or usuario_val == "None" or not contra_val or contra_val == "None":
            continue  # cuenta no creada todavía
        if not correo_val or "@" not in correo_val:
            continue  # sin correo personal para enviar
        if _looks_sent(ya_enviado) or _looks_sent(ya_enviado_legacy):
            continue  # ya se le envió (por este sistema o por el proceso anterior)

        rows_to_send.append(r_idx)

    if len(rows_to_send) > 300:
        raise HTTPException(
            status_code=400,
            detail=f"Límite de seguridad excedido: intentas enviar {len(rows_to_send)} correos a la vez (máximo 300 permitidos).",
        )

    async def send_row(r_idx):
        nombre = str(ws.cell(row=r_idx, column=col_nombre).value or "").strip() if col_nombre else ""
        usuario_val = str(ws.cell(row=r_idx, column=col_usuario).value or "").strip()
        contra_val = str(ws.cell(row=r_idx, column=col_contra).value or "").strip()
        correo_val = str(ws.cell(row=r_idx, column=col_correo).value or "").strip()
        curso_nombre = str(ws.cell(row=r_idx, column=col_curso_nombre).value or "").strip() if col_curso_nombre else ""

        try:
            await email_service.send_credentials_email(
                to_email=correo_val,
                full_name=nombre or usuario_val,
                login_id=usuario_val,
                password=contra_val,
                program_type=program_type,
                program_name=curso_nombre or title_val,
                extra_cc=sheet_cc_list,
            )
            ws.cell(row=r_idx, column=col_correo_enviado, value="✅ Enviado")
            ws.cell(row=r_idx, column=col_correo_enviado).font = Font(color="00B050", bold=True)
            if col_enviado_legacy:
                ws.cell(row=r_idx, column=col_enviado_legacy, value="✅ Enviado")
                ws.cell(row=r_idx, column=col_enviado_legacy).font = Font(color="00B050", bold=True)
            result.succeeded.append({"correo": correo_val, "usuario": usuario_val})
        except Exception as exc:
            ws.cell(row=r_idx, column=col_correo_enviado, value=f"❌ Error: {exc}")
            ws.cell(row=r_idx, column=col_correo_enviado).font = Font(color="FF0000")
            result.failed.append({"correo": correo_val, "error": str(exc)})

    batch_size = 5
    for i in range(0, len(rows_to_send), batch_size):
        await asyncio.gather(*(send_row(r) for r in rows_to_send[i:i + batch_size]))

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    try:
        await graph.put_raw(f"/shares/{encoded_url}/driveItem/content", output.getvalue())
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"No se pudo guardar el archivo actualizado en OneDrive. {e}")

    if not result.succeeded and not result.failed:
        raise HTTPException(
            status_code=400,
            detail="No hay alumnos pendientes de envío: o no tienen cuenta creada todavía, o ya se les envió el correo, o no tienen correo personal cargado.",
        )

    return result


@router.post("/excel/docentes/send-credentials/preview", summary="Vista previa: quién falta enviar credenciales (Docentes)")
async def preview_send_docentes_credentials(req: DiplomadosUrlRequest) -> SendCredentialsPreviewResponse:
    """Muestra quién está pendiente de recibir el correo de credenciales, sin
    enviar nada — misma lógica de filtrado exacta que /send-credentials."""
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")

    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo de OneDrive. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail="El archivo descargado no es un Excel válido.")

    if req.sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"La pestaña '{req.sheet_name}' no existe.")

    ws = wb[req.sheet_name]
    header_row_idx, headers, _ = _find_header_row_and_headers(ws)
    if not header_row_idx:
        raise HTTPException(status_code=400, detail="No se encontró la fila de encabezados.")

    def get_col_idx(*keys):
        return _match_col_idx(headers, *keys)

    col_nombre = get_col_idx("nombre", "alumno", "estudiante")
    col_correo = get_col_idx("correo", "email")
    col_usuario = get_col_idx("usuario")
    col_contra = get_col_idx("contrasena", "contraseña", "clave")

    if not col_usuario or not col_contra:
        raise HTTPException(
            status_code=400,
            detail="No se encontraron las columnas de Usuario/Contraseña. Primero procesa la planilla con 'Alta Docentes'.",
        )
    if not col_correo:
        raise HTTPException(status_code=400, detail="No se encontró una columna de Correo para enviar las credenciales.")

    col_correo_enviado = get_col_idx("correo enviado", "credenciales enviadas")
    col_enviado_legacy = headers.get("enviado")

    def _looks_sent(raw: str) -> bool:
        v = raw.strip().lower()
        return bool(v) and ("✅" in raw or v in ("si", "yes", "true", "enviado"))

    pending, already_sent, no_account, no_email = [], 0, 0, 0
    for r_idx in range(header_row_idx + 1, ws.max_row + 1):
        nombre_val = str(ws.cell(row=r_idx, column=col_nombre).value or "").strip() if col_nombre else ""
        usuario_val = str(ws.cell(row=r_idx, column=col_usuario).value or "").strip()
        contra_val = str(ws.cell(row=r_idx, column=col_contra).value or "").strip()
        correo_val = str(ws.cell(row=r_idx, column=col_correo).value or "").strip()
        if not nombre_val and not usuario_val and not correo_val:
            continue

        ya_enviado = str(ws.cell(row=r_idx, column=col_correo_enviado).value or "") if col_correo_enviado else ""
        ya_enviado_legacy = str(ws.cell(row=r_idx, column=col_enviado_legacy).value or "") if col_enviado_legacy else ""

        if not usuario_val or usuario_val == "None" or not contra_val or contra_val == "None":
            no_account += 1
            continue
        if not correo_val or "@" not in correo_val:
            no_email += 1
            continue
        if _looks_sent(ya_enviado) or _looks_sent(ya_enviado_legacy):
            already_sent += 1
            continue

        pending.append({"nombre": nombre_val, "correo": correo_val, "usuario": usuario_val})

    wb.close()
    return SendCredentialsPreviewResponse(
        pending=pending, already_sent_count=already_sent,
        no_account_count=no_account, no_email_count=no_email,
    )


@router.post("/excel/docentes/send-credentials", summary="Enviar correo de credenciales a docentes ya creados")
async def send_docentes_credentials(req: DiplomadosUrlRequest) -> BulkResult:
    """Envía el correo de credenciales a los docentes que ya tienen cuenta creada
    (Usuario + Contraseña ya generados) y que todavía no la recibieron.

    Acción separada de la creación de cuentas (misma lógica que Diplomados):
    permite reintentar el envío sin volver a crear nada, y no reenvía a quien
    ya fue marcado como enviado.
    """
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")

    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo de OneDrive. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents))
    except Exception as e:
        raise HTTPException(status_code=400, detail="El archivo descargado no es un Excel válido.")

    result = BulkResult()
    if req.sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"La pestaña '{req.sheet_name}' no existe.")

    ws = wb[req.sheet_name]
    header_row_idx, headers, _ = _find_header_row_and_headers(ws)
    if not header_row_idx:
        raise HTTPException(status_code=400, detail="No se encontró la fila de encabezados.")

    def get_col_idx(*keys):
        return _match_col_idx(headers, *keys)

    col_nombre = get_col_idx("nombre", "alumno", "estudiante")
    col_correo = get_col_idx("correo", "email")
    col_usuario = get_col_idx("usuario")
    col_contra = get_col_idx("contrasena", "contraseña", "clave")
    col_curso_nombre = get_col_idx("nombre del curso", "curso", "diplomado")

    if not col_usuario or not col_contra:
        raise HTTPException(
            status_code=400,
            detail="No se encontraron las columnas de Usuario/Contraseña. Primero procesa la planilla con 'Alta Docentes'.",
        )
    if not col_correo:
        raise HTTPException(status_code=400, detail="No se encontró una columna de Correo para enviar las credenciales.")

    col_correo_enviado = get_col_idx("correo enviado", "credenciales enviadas")
    if not col_correo_enviado:
        col_correo_enviado = ws.max_column + 1
        _safe_set_cell(ws, header_row_idx, col_correo_enviado, "Correo Enviado").font = Font(bold=True)

    # Respeta una columna "Enviado" preexistente (marcada a mano o por un
    # proceso externo, antes de este sistema) además de nuestra "Correo
    # Enviado", para no reenviar algo que ya salió por otra vía.
    col_enviado_legacy = headers.get("enviado")

    col_cc = get_col_idx("cc", "copia")
    sheet_cc_list: list[str] = []
    if col_cc:
        for r_idx in range(header_row_idx + 1, ws.max_row + 1):
            cc_val = str(ws.cell(row=r_idx, column=col_cc).value or "").strip()
            if cc_val:
                for email in cc_val.replace(";", ",").replace("\n", ",").split(","):
                    email = email.strip()
                    if "@" in email and email not in sheet_cc_list:
                        sheet_cc_list.append(email)

    def _looks_sent(raw: str) -> bool:
        v = raw.strip().lower()
        return bool(v) and ("✅" in raw or v in ("si", "yes", "true", "enviado"))

    rows_to_send = []
    for r_idx in range(header_row_idx + 1, ws.max_row + 1):
        usuario_val = str(ws.cell(row=r_idx, column=col_usuario).value or "").strip()
        contra_val = str(ws.cell(row=r_idx, column=col_contra).value or "").strip()
        correo_val = str(ws.cell(row=r_idx, column=col_correo).value or "").strip()
        ya_enviado = str(ws.cell(row=r_idx, column=col_correo_enviado).value or "")
        ya_enviado_legacy = str(ws.cell(row=r_idx, column=col_enviado_legacy).value or "") if col_enviado_legacy else ""

        if not usuario_val or usuario_val == "None" or not contra_val or contra_val == "None":
            continue  # cuenta no creada todavía
        if not correo_val or "@" not in correo_val:
            continue  # sin correo personal para enviar
        if _looks_sent(ya_enviado) or _looks_sent(ya_enviado_legacy):
            continue  # ya se le envió (por este sistema o por el proceso anterior)

        rows_to_send.append(r_idx)

    if len(rows_to_send) > 300:
        raise HTTPException(
            status_code=400,
            detail=f"Límite de seguridad excedido: intentas enviar {len(rows_to_send)} correos a la vez (máximo 300 permitidos).",
        )

    async def send_row(r_idx):
        nombre = str(ws.cell(row=r_idx, column=col_nombre).value or "").strip() if col_nombre else ""
        usuario_val = str(ws.cell(row=r_idx, column=col_usuario).value or "").strip()
        contra_val = str(ws.cell(row=r_idx, column=col_contra).value or "").strip()
        correo_val = str(ws.cell(row=r_idx, column=col_correo).value or "").strip()
        curso_nombre = str(ws.cell(row=r_idx, column=col_curso_nombre).value or "").strip() if col_curso_nombre else ""

        try:
            await email_service.send_credentials_email(
                to_email=correo_val,
                full_name=nombre or usuario_val,
                login_id=usuario_val,
                password=contra_val,
                program_name=curso_nombre,
                extra_cc=sheet_cc_list,
            )
            ws.cell(row=r_idx, column=col_correo_enviado, value="✅ Enviado")
            ws.cell(row=r_idx, column=col_correo_enviado).font = Font(color="00B050", bold=True)
            if col_enviado_legacy:
                ws.cell(row=r_idx, column=col_enviado_legacy, value="✅ Enviado")
                ws.cell(row=r_idx, column=col_enviado_legacy).font = Font(color="00B050", bold=True)
            result.succeeded.append({"correo": correo_val, "usuario": usuario_val})
        except Exception as exc:
            ws.cell(row=r_idx, column=col_correo_enviado, value=f"❌ Error: {exc}")
            ws.cell(row=r_idx, column=col_correo_enviado).font = Font(color="FF0000")
            result.failed.append({"correo": correo_val, "error": str(exc)})

    batch_size = 5
    for i in range(0, len(rows_to_send), batch_size):
        await asyncio.gather(*(send_row(r) for r in rows_to_send[i:i + batch_size]))

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    try:
        await graph.put_raw(f"/shares/{encoded_url}/driveItem/content", output.getvalue())
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"No se pudo guardar el archivo actualizado en OneDrive. {e}")

    if not result.succeeded and not result.failed:
        raise HTTPException(
            status_code=400,
            detail="No hay docentes pendientes de envío: o no tienen cuenta creada todavía, o ya se les envió el correo, o no tienen correo personal cargado.",
        )

    return result


# ═══════════════════════════════════════════════════════════════════════════════
# Cursos Canvas (Lectura y Escritura de Planilla OneDrive)
# ═══════════════════════════════════════════════════════════════════════════════

class CoursesPreviewResponse(BaseModel):
    sheet_name: str
    courses_to_create: int
    courses_already_created: int
    headers: list[str]
    sample_rows: list[dict]
    headers: list[str] = []
    sample_rows: list[dict] = []
    course_details: list[dict] = []
    platform: str = "both"


@router.post("/excel/courses/sheets", response_model=list[str])
async def get_courses_sheets(req: UrlOnlyRequest) -> list[str]:
    """Lista las pestañas de un archivo Excel de cursos en OneDrive."""
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")

    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
        sheets = wb.sheetnames
        wb.close()
        return sheets
    except Exception as e:
        raise HTTPException(status_code=400, detail="El archivo no es un Excel válido.")


@router.post("/excel/courses/preview", summary="Pre-visualizar planilla de Cursos")
async def preview_courses_onedrive(req: DiplomadosUrlRequest) -> CoursesPreviewResponse:
    """Lee la planilla de cursos y devuelve un resumen de los que se van a crear."""
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")

    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail="El archivo no es un Excel válido.")

    if req.sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"La pestaña '{req.sheet_name}' no existe.")

    ws = wb[req.sheet_name]

    sample_rows = []


    header_row_idx, headers_dict, headers_raw = _find_header_row_and_headers(ws)


    if not header_row_idx:
        raise HTTPException(status_code=400, detail="No se encontró la fila de encabezados.")

    headers = [h for h in headers_raw if h]


    col_estado = -1
    col_nombre = -1
    col_sis = -1
    col_docente = -1
    col_coordinadores = -1
    for i, h in enumerate(headers_raw):
        h_lower = h.lower()
        if "estado" in h_lower or ("canvas" in h_lower and "id" in h_lower):
            col_estado = i
        if "nombre" in h_lower or "curso" in h_lower:
            if col_nombre == -1: col_nombre = i
        if "identificaci" in h_lower or "sis" in h_lower:
            if col_sis == -1: col_sis = i
        if "docente" in h_lower or "usuario institucional" in h_lower or "instructor" in h_lower or "profesor" in h_lower:
            if col_docente == -1: col_docente = i
        if "coordinador" in h_lower or "diseñador" in h_lower or "disenador" in h_lower or "designer" in h_lower:
            if col_coordinadores == -1: col_coordinadores = i

    courses_to_create = 0
    courses_already_created = 0
    course_details = []

    for row_idx in range(header_row_idx + 1, ws.max_row + 1):
        row_vals = [str(ws.cell(row=row_idx, column=c).value or "").strip() for c in range(1, len(headers_raw) + 1)]
        
        # Check if the row has at least a course name (assuming column 0 or 1 is name)
        if not any(row_vals):
            continue
            
        estado_val = row_vals[col_estado] if col_estado >= 0 else ""
        if "✅" in estado_val or estado_val.lower() in ["creado", "ok", "si"]:
            courses_already_created += 1
        else:
            courses_to_create += 1
            course_details.append({
                "nombre": row_vals[col_nombre] if col_nombre >= 0 else "",
                "sis_id": row_vals[col_sis] if col_sis >= 0 else "",
                "docente": row_vals[col_docente] if col_docente >= 0 else "",
                "coordinadores": row_vals[col_coordinadores] if col_coordinadores >= 0 else "",
                "row": {h: v for h, v in zip(headers_raw, row_vals) if h}
            })

        if len(sample_rows) < 10:
            sample_rows.append({h: v for h, v in zip(headers_raw, row_vals) if h})
            
    wb.close()

    # Misma detección de plataforma por nombre de pestaña que usa
    # _process_courses_bg — se informa acá para que el preview no prometa
    # matriculación en Teams cuando la planilla es de Canvas solamente (o
    # viceversa).
    sheet_lower = req.sheet_name.lower()
    platform = "both"
    if "canvas" in sheet_lower and "teams" not in sheet_lower:
        platform = "canvas"
    elif "teams" in sheet_lower and "canvas" not in sheet_lower:
        platform = "teams"

    return CoursesPreviewResponse(
        sheet_name=req.sheet_name,
        courses_to_create=courses_to_create,
        courses_already_created=courses_already_created,
        headers=headers,
        sample_rows=sample_rows,
        course_details=course_details,
        platform=platform,
    )


# ── Corregir rol de Coordinadores (Diseñador → Coordinador) ────────────────────
#
# Antes de que existiera el rol personalizado "Coordinador" en Canvas, los
# coordinadores/diseñadores se inscribían con el rol base "Designer". Esto
# detecta, para los cursos/coordinadores de una planilla ya cargada, cuáles
# quedaron con el rol base y los migra al rol personalizado "Coordinador"
# (se borra la inscripción vieja y se crea una nueva con el role_id correcto,
# porque Canvas no permite cambiar el rol de una inscripción existente).

class FixCoordinadorRoleItem(BaseModel):
    course_id: str
    course_name: str = ""
    sis_id: str = ""
    email: str = ""
    user_id: str
    enrollment_id: str
    current_role: str = ""


@router.post("/excel/courses/fix-coordinador-role/preview", summary="Detectar coordinadores inscritos como Diseñador para migrarlos al rol Coordinador")
async def preview_fix_coordinador_role(req: DiplomadosUrlRequest):
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")

    coordinador_role_id = await _get_coordinador_role_id()
    if not coordinador_role_id:
        raise HTTPException(
            status_code=400,
            detail="No se encontró el rol personalizado 'Coordinador' en la cuenta de Canvas (Admin > Permisos).",
        )

    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="El archivo no es un Excel válido.")

    if req.sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"La pestaña '{req.sheet_name}' no existe.")
    ws = wb[req.sheet_name]

    header_row_idx, _, headers_raw = _find_header_row_and_headers(ws)
    if not header_row_idx:
        raise HTTPException(status_code=400, detail="No se encontró la fila de encabezados.")

    col_nombre = -1
    col_sis = -1
    col_coordinadores = -1
    for i, h in enumerate(headers_raw):
        h_lower = h.lower()
        if ("nombre" in h_lower or "curso" in h_lower) and col_nombre == -1:
            col_nombre = i
        if ("identificaci" in h_lower or "sis" in h_lower) and col_sis == -1:
            col_sis = i
        if ("coordinador" in h_lower or "diseñador" in h_lower or "disenador" in h_lower or "designer" in h_lower) and col_coordinadores == -1:
            col_coordinadores = i

    rows = []
    for row_idx in range(header_row_idx + 1, ws.max_row + 1):
        row_vals = [str(ws.cell(row=row_idx, column=c).value or "").strip() for c in range(1, len(headers_raw) + 1)]
        if not any(row_vals):
            continue
        sis_id = row_vals[col_sis] if 0 <= col_sis < len(row_vals) else ""
        coordinadores_raw = row_vals[col_coordinadores] if 0 <= col_coordinadores < len(row_vals) else ""
        nombre = row_vals[col_nombre] if 0 <= col_nombre < len(row_vals) else ""
        emails = [e.strip() for e in re.split(r'[,;/\n]+', coordinadores_raw) if e.strip() and "@" in e.strip()]
        if sis_id and emails:
            rows.append({"nombre": nombre, "sis_id": sis_id, "emails": emails})
    wb.close()

    items: list[FixCoordinadorRoleItem] = []
    not_found: list[str] = []
    already_correct = 0

    async def process_row(row):
        nonlocal already_correct
        sis_id = row["sis_id"]
        try:
            course = await canvas.get(f"/courses/sis_course_id:{sis_id}")
        except Exception:
            course = None
        if not course:
            not_found.append(f"{row['nombre']} (SIS {sis_id}): curso no encontrado")
            return
        course_id = str(course.get("id"))
        course_name = course.get("name") or row["nombre"]

        async def process_email(email: str):
            nonlocal already_correct
            try:
                match = await canvas.find_user_exact(_ACCOUNT, email, fields=("email", "login_id"))
                user_id = str(match["id"]) if match else None
            except Exception:
                user_id = None
            if not user_id:
                not_found.append(f"{course_name}: {email} (no encontrado en Canvas)")
                return
            try:
                enrollments = await canvas.get(f"/courses/{course_id}/enrollments", params={"user_id": user_id})
            except Exception:
                enrollments = []

            target = None
            has_coordinador = False
            for en in enrollments or []:
                if en.get("enrollment_state") == "deleted":
                    continue
                if en.get("role") == "Coordinador":
                    has_coordinador = True
                elif en.get("type") == "DesignerEnrollment" and target is None:
                    target = en

            if target:
                items.append(FixCoordinadorRoleItem(
                    course_id=course_id,
                    course_name=course_name,
                    sis_id=sis_id,
                    email=email,
                    user_id=user_id,
                    enrollment_id=str(target.get("id")),
                    current_role=target.get("role") or target.get("type") or "",
                ))
            elif has_coordinador:
                already_correct += 1
            else:
                not_found.append(f"{course_name}: {email} (sin inscripción de Diseñador para migrar)")

        await asyncio.gather(*(process_email(e) for e in row["emails"]))

    await asyncio.gather(*(process_row(r) for r in rows))

    return {"items": [i.model_dump() for i in items], "not_found": not_found, "already_correct": already_correct}


class FixCoordinadorRoleByCourseRequest(BaseModel):
    entries: list[str]


@router.post("/excel/courses/fix-coordinador-role/resolve-by-course", summary="Detectar inscripciones de Diseñador a migrar a partir de una lista de IDs/SIS IDs de curso")
async def resolve_fix_coordinador_role_by_course(req: FixCoordinadorRoleByCourseRequest):
    coordinador_role_id = await _get_coordinador_role_id()
    if not coordinador_role_id:
        raise HTTPException(
            status_code=400,
            detail="No se encontró el rol personalizado 'Coordinador' en la cuenta de Canvas (Admin > Permisos).",
        )

    items: list[FixCoordinadorRoleItem] = []
    not_found: list[str] = []
    already_correct = 0
    seen: set[str] = set()

    async def process_entry(raw: str):
        nonlocal already_correct
        entry = raw.strip()
        if not entry or entry in seen:
            return
        seen.add(entry)

        course = None
        try:
            course = await canvas.get(f"/courses/sis_course_id:{entry}")
        except Exception:
            pass
        if not course:
            try:
                course = await canvas.get(f"/courses/{entry}")
            except Exception:
                pass
        if not course:
            not_found.append(f"{entry}: curso no encontrado")
            return

        course_id = str(course.get("id"))
        course_name = course.get("name") or entry

        try:
            enrollments = await canvas.paginate(f"/courses/{course_id}/enrollments", params={"type[]": "DesignerEnrollment"})
        except Exception as e:
            not_found.append(f"{course_name}: no se pudieron leer las inscripciones ({e})")
            return

        found_any = False
        for en in enrollments or []:
            if en.get("enrollment_state") == "deleted":
                continue
            if en.get("role") == "Coordinador":
                already_correct += 1
                continue
            found_any = True
            user = en.get("user") or {}
            items.append(FixCoordinadorRoleItem(
                course_id=course_id,
                course_name=course_name,
                sis_id=entry,
                email=user.get("login_id") or user.get("email") or user.get("name") or f"user {en.get('user_id')}",
                user_id=str(en.get("user_id")),
                enrollment_id=str(en.get("id")),
                current_role=en.get("role") or en.get("type") or "",
            ))
        if not found_any and not any(en.get("role") == "Coordinador" for en in (enrollments or [])):
            not_found.append(f"{course_name}: sin inscripciones de Diseñador")

    await asyncio.gather(*(process_entry(e) for e in req.entries))

    return {"items": [i.model_dump() for i in items], "not_found": not_found, "already_correct": already_correct}


class FixCoordinadorRoleRequest(BaseModel):
    items: list[FixCoordinadorRoleItem]


@router.post("/excel/courses/fix-coordinador-role/execute", summary="Migrar inscripciones seleccionadas de Diseñador a Coordinador")
async def execute_fix_coordinador_role(req: FixCoordinadorRoleRequest):
    coordinador_role_id = await _get_coordinador_role_id()
    if not coordinador_role_id:
        raise HTTPException(
            status_code=400,
            detail="No se encontró el rol personalizado 'Coordinador' en la cuenta de Canvas (Admin > Permisos).",
        )

    result = {"succeeded": [], "failed": []}

    async def fix_one(item: FixCoordinadorRoleItem):
        label = f"{item.course_name or item.course_id}: {item.email or item.user_id}"
        try:
            await canvas.delete(f"/courses/{item.course_id}/enrollments/{item.enrollment_id}", params={"task": "delete"})
        except Exception as e:
            result["failed"].append({"item": label, "error": f"No se pudo remover la inscripción de Diseñador: {e}"})
            return
        try:
            await canvas.post(f"/courses/{item.course_id}/enrollments", {
                "enrollment": {
                    "user_id": item.user_id,
                    "role_id": coordinador_role_id,
                    "enrollment_state": "active",
                    "notify": False,
                }
            })
            result["succeeded"].append(label)
        except Exception as e:
            result["failed"].append({"item": label, "error": f"Se removió Diseñador pero falló la inscripción como Coordinador: {e}"})

    await asyncio.gather(*(fix_one(i) for i in req.items))
    return result


# ── Corregir SIS ID en masa ─────────────────────────────────────────────────
#
# Para cursos ya creados en Canvas con un SIS ID mal cargado (ej. copiado de
# la fila equivocada). Reutiliza la misma planilla de creación de cursos:
# lee "ID CANVAS (Autogenerado)" para identificar el curso ya existente y
# "Identificación del SIS (Canvas)" como el valor CORRECTO al que debe
# quedar. Al hacer PUT /courses/{id} con el sis_course_id nuevo, Canvas
# libera automáticamente el valor viejo de ese curso (un curso solo puede
# tener un SIS ID a la vez) — no hace falta un paso aparte para "liberarlo".

class FixSisIdPreviewResponse(BaseModel):
    sheet_name: str
    to_fix: int
    already_correct: int
    no_course_found: int
    sample_rows: list[dict] = []


@router.post("/excel/courses/fix-sis-id/preview", summary="Detectar cursos ya creados con SIS ID incorrecto para corregirlo masivamente")
async def preview_fix_sis_id(req: DiplomadosUrlRequest) -> FixSisIdPreviewResponse:
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")

    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="El archivo no es un Excel válido.")

    if req.sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"La pestaña '{req.sheet_name}' no existe.")
    ws = wb[req.sheet_name]

    header_row_idx, headers_dict, headers_raw = _find_header_row_and_headers(ws)
    if not header_row_idx:
        raise HTTPException(status_code=400, detail="No se encontró la fila de encabezados.")

    def get_col(*keys):
        for k in keys:
            for h, idx in headers_dict.items():
                if _norm(k) in h:
                    return idx
        return None

    col_nombre = get_col("nombre", "curso")
    col_sis = get_col("sis", "sys")
    col_canvas_id = get_col("canvas id", "id canvas")
    if not col_canvas_id or not col_sis:
        raise HTTPException(status_code=400, detail="Faltan columnas de ID Canvas (Autogenerado) o SIS ID en la planilla.")

    rows = []
    for r_idx in range(header_row_idx + 1, ws.max_row + 1):
        canvas_id = str(ws.cell(row=r_idx, column=col_canvas_id).value or "").strip()
        new_sis = str(ws.cell(row=r_idx, column=col_sis).value or "").strip()
        nombre = str(ws.cell(row=r_idx, column=col_nombre).value or "").strip() if col_nombre else ""
        if canvas_id and canvas_id.isdigit() and new_sis and new_sis != "None":
            rows.append({"canvas_id": canvas_id, "new_sis": new_sis, "nombre": nombre})
    wb.close()

    to_fix = 0
    already_correct = 0
    no_course_found = 0
    sample_rows: list[dict] = []

    async def check(row):
        nonlocal to_fix, already_correct, no_course_found
        try:
            course = await canvas.get(f"/courses/{row['canvas_id']}")
            current_sis = str(course.get("sis_course_id") or "").strip()
        except Exception:
            no_course_found += 1
            return
        if current_sis == row["new_sis"]:
            already_correct += 1
        else:
            to_fix += 1
            if len(sample_rows) < 15:
                sample_rows.append({
                    "nombre": row["nombre"],
                    "canvas_id": row["canvas_id"],
                    "sis_actual": current_sis or "(vacío)",
                    "sis_nuevo": row["new_sis"],
                })

    await asyncio.gather(*(check(r) for r in rows))

    return FixSisIdPreviewResponse(
        sheet_name=req.sheet_name,
        to_fix=to_fix,
        already_correct=already_correct,
        no_course_found=no_course_found,
        sample_rows=sample_rows,
    )


@router.post("/excel/courses/fix-sis-id", summary="Corregir masivamente el SIS ID de cursos ya creados en Canvas")
async def fix_sis_id_onedrive(req: DiplomadosUrlRequest, bg_tasks: BackgroundTasks) -> dict:
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")

    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo de OneDrive. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
        if req.sheet_name not in wb.sheetnames:
            raise HTTPException(status_code=400, detail=f"La pestaña '{req.sheet_name}' no existe.")
        wb.close()
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=400, detail="El archivo no es un Excel válido.")

    job_id = await jobs.create_job(
        job_type="fix_sis_id_onedrive",
        operation="Corrección masiva de SIS ID",
        username="admin",
    )
    bg_tasks.add_task(_process_fix_sis_id_bg, job_id, req, contents, encoded_url)

    return {"status": "success", "job_id": job_id, "message": "Corrección de SIS ID iniciada en segundo plano."}


async def _process_fix_sis_id_bg(job_id: int, req: DiplomadosUrlRequest, contents: bytes, encoded_url: str):
    await jobs.start_job(job_id)

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents))
    except Exception as e:
        await jobs.fail_job(job_id, f"El archivo no es un Excel válido: {e}")
        return

    if req.sheet_name not in wb.sheetnames:
        await jobs.fail_job(job_id, f"La pestaña '{req.sheet_name}' no existe.")
        return

    ws = wb[req.sheet_name]
    header_row_idx, headers_dict, headers_raw = _find_header_row_and_headers(ws)
    if not header_row_idx:
        await jobs.fail_job(job_id, "No se encontró la fila de encabezados.")
        return

    def get_col(*keys):
        for k in keys:
            for h, idx in headers_dict.items():
                if _norm(k) in h:
                    return idx
        return None

    col_canvas_id = get_col("canvas id", "id canvas")
    col_sis = get_col("sis", "sys")
    col_estado = get_col("estado")
    if not col_canvas_id or not col_sis:
        await jobs.fail_job(job_id, "Faltan columnas de ID Canvas (Autogenerado) o SIS ID en la planilla.")
        return

    if not col_estado:
        col_estado = ws.max_column + 1
        ws.cell(row=header_row_idx, column=col_estado, value="Estado (Autogenerado)").font = Font(bold=True)

    rows = []
    for r_idx in range(header_row_idx + 1, ws.max_row + 1):
        canvas_id = str(ws.cell(row=r_idx, column=col_canvas_id).value or "").strip()
        new_sis = str(ws.cell(row=r_idx, column=col_sis).value or "").strip()
        if canvas_id and canvas_id.isdigit() and new_sis and new_sis != "None":
            rows.append((r_idx, canvas_id, new_sis))

    success_count = 0
    error_count = 0

    async def fix_one(r_idx: int, canvas_id: str, new_sis: str):
        nonlocal success_count, error_count
        try:
            course = await canvas.get(f"/courses/{canvas_id}")
            current_sis = str(course.get("sis_course_id") or "").strip()
        except Exception as e:
            ws.cell(row=r_idx, column=col_estado, value=f"❌ Error: curso {canvas_id} no encontrado ({e})").font = Font(color="FF0000", bold=True)
            error_count += 1
            return

        if current_sis == new_sis:
            ws.cell(row=r_idx, column=col_estado, value="✅ Ya estaba correcto").font = Font(color="00B050", bold=True)
            success_count += 1
            return

        try:
            await canvas.put(f"/courses/{canvas_id}", {"course": {"sis_course_id": new_sis}})
            ws.cell(row=r_idx, column=col_estado, value="✅ SIS corregido").font = Font(color="00B050", bold=True)
            success_count += 1
        except Exception as e:
            err = getattr(e, "detail", None) or str(e)
            if "sis" in str(err).lower() and "already" in str(err).lower():
                err = f"{err} — el SIS ID '{new_sis}' ya está en uso por otro curso; corregí primero ese curso para liberarlo."
            ws.cell(row=r_idx, column=col_estado, value=f"❌ Error: {err}").font = Font(color="FF0000", bold=True)
            error_count += 1

    batch_size = 5
    for i in range(0, len(rows), batch_size):
        chunk = rows[i:i + batch_size]
        await asyncio.gather(*(fix_one(*r) for r in chunk))
        await jobs.update_job_progress(job_id, success_count, error_count)

    out_io = io.BytesIO()
    wb.save(out_io)
    out_io.seek(0)

    try:
        await graph.put_raw(f"/shares/{encoded_url}/driveItem/content", out_io.read())
        await jobs.complete_job(job_id, success_count, error_count, "Guardado en OneDrive correctamente.")
    except Exception as e:
        await jobs.complete_job(job_id, success_count, error_count, f"Procesado, pero no se pudo guardar en OneDrive: {e}")


async def append_report_onedrive(report_url: str, succeeded: list, failed: list):
    try:
        import datetime, io, base64, openpyxl
        encoded = base64.urlsafe_b64encode(report_url.encode("utf-8")).decode("utf-8").rstrip("=")
        encoded_url = "u!" + encoded.replace("-", "+").replace("_", "/")
        
        r = await graph._client().get(f"{graph._GRAPH}/shares/{encoded_url}/driveItem/content", headers=graph._headers())
        if r.status_code != 200:
            print(f"No se pudo descargar el reporte maestro: {r.status_code}")
            return
            
        file_content = r.content
        wb = openpyxl.load_workbook(io.BytesIO(file_content))
        
        sheet_name = datetime.datetime.now().strftime("%d-%m-%Y")
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.create_sheet(title=sheet_name)
            
        if ws.max_row == 1 and not ws.cell(row=1, column=1).value:
            ws.append(['Fecha Hora', 'Nombre del Curso', 'Canvas ID', 'SIS Course ID', 'Estado', 'Teams ID'])
            
        now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        for c in succeeded:
            nombre = c.get("nombre") or c.get("input", {}).get("nombre", "")
            ws.append([now_str, nombre, c.get("canvas_id", ""), c.get("sis_course_id", ""), "Creado", c.get("teams_id", "")])
            
        for c in failed:
            nombre = c.get("nombre") or c.get("input", {}).get("nombre", "")
            ws.append([now_str, nombre, c.get("canvas_id", ""), c.get("sis_course_id", ""), f"Fallo: {c.get('error', '')}", ""])
            
        out_io = io.BytesIO()
        wb.save(out_io)
        out_io.seek(0)
        
        await graph.put_raw(f"/shares/{encoded_url}/driveItem/content", out_io.read())
        print("Reporte maestro actualizado exitosamente.")
    except Exception as e:
        print(f"Error actualizando reporte maestro: {str(e)}")

@router.post("/excel/courses", summary="Crear cursos en Canvas desde planilla OneDrive")
async def import_courses_onedrive(req: DiplomadosUrlRequest, bg_tasks: BackgroundTasks) -> dict:
    """Descarga y valida la planilla, y dispara la creación de cursos en segundo plano.

    Crear cursos + equipos de Teams (con espera de replicación de Azure AD
    de ~20-60s por equipo) es demasiado lento para una request HTTP síncrona
    con lotes de más de un puñado de filas — el proxy/navegador corta la
    conexión antes de que termine y en el frontend "no pasa nada". Por eso
    esto ahora sigue el mismo patrón de job en segundo plano que
    Matriculaciones Masivas (ver _process_matriculaciones_bg)."""
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")

    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo de OneDrive. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
        if req.sheet_name not in wb.sheetnames:
            raise HTTPException(status_code=400, detail=f"La pestaña '{req.sheet_name}' no existe.")
        wb.close()
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=400, detail="El archivo no es un Excel válido.")

    job_id = await jobs.create_job(
        job_type="creacion_cursos_onedrive",
        operation="Creación masiva de cursos",
        username="admin",
    )
    bg_tasks.add_task(_process_courses_bg, job_id, req, contents, encoded_url)

    return {"status": "success", "job_id": job_id, "message": "Creación de cursos iniciada en segundo plano."}


async def _process_courses_bg(job_id: int, req: DiplomadosUrlRequest, contents: bytes, encoded_url: str):
    """Crea cursos simultáneamente en Canvas y equipos en Teams leyendo de OneDrive."""
    await jobs.start_job(job_id)

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents))
    except Exception as e:
        await jobs.fail_job(job_id, f"El archivo no es un Excel válido: {e}")
        return

    if req.sheet_name not in wb.sheetnames:
        await jobs.fail_job(job_id, f"La pestaña '{req.sheet_name}' no existe.")
        return

    _ACCOUNT_LOCAL = settings.canvas_account_id
    result = BulkResult()
    ws = wb[req.sheet_name]

    header_row_idx, headers, _ = _find_header_row_and_headers(ws)


    if not header_row_idx:
        await jobs.fail_job(job_id, "No se encontró la fila de encabezados.")
        return

    def get_col(*keys):
        for k in keys:
            for h, idx in headers.items():
                if _norm(k) in h:
                    return idx
        return None

    col_fecha = get_col("fecha")
    col_nombre = get_col("nombre", "curso", "canvas")
    col_sis = get_col("sis", "sys")
    col_periodo = get_col("periodo")
    col_subcuenta = get_col("sub-cuenta", "subcuenta", "sub_cuenta")
    col_canvas_id = get_col("canvas id", "id canvas")
    col_teams_id = get_col("teams id", "id teams")
    col_estado = get_col("estado")
    col_docente = get_col("docente", "usuario institucional", "instructor", "profesor")
    col_coordinadores = get_col("coordinador", "coordinadores", "diseñador", "disenador", "designer")

    if not col_nombre:
        await jobs.fail_job(job_id, "No se encontró la columna de nombre del curso.")
        return

    next_col = ws.max_column + 1
    if not col_fecha:
        col_fecha = next_col
        ws.cell(row=header_row_idx, column=col_fecha, value="Fecha de Creación").font = Font(bold=True)
        next_col += 1
    if not col_canvas_id:
        col_canvas_id = next_col
        ws.cell(row=header_row_idx, column=col_canvas_id, value="CANVAS ID").font = Font(bold=True)
        next_col += 1
    if not col_teams_id:
        col_teams_id = next_col
        ws.cell(row=header_row_idx, column=col_teams_id, value="TEAMS ID").font = Font(bold=True)
        next_col += 1
    if not col_estado:
        col_estado = next_col
        ws.cell(row=header_row_idx, column=col_estado, value="ESTADO").font = Font(bold=True)
        next_col += 1

    from datetime import datetime
    import time
    from app.services.teams_client import create_team_via_group

    terms_data = []
    try:
        terms_res = await canvas.get(f"/accounts/{_ACCOUNT_LOCAL}/terms", params={"per_page": 100})
        terms_data = terms_res.get("enrollment_terms", [])
    except: pass

    # Mapa nombre→id de TODAS las subcuentas bajo la cuenta raíz (recursivo,
    # sin importar cuántos niveles de anidamiento tengan — ej. USIL > 2026-2
    # > CPEL > Administración). Se arma una sola vez para todo el job en vez
    # de una consulta por fila.
    # Se indexa con _norm() (mismo helper que usa el resto de la detección de
    # columnas/encabezados): saca tildes, mayúsculas y separadores, para que
    # "Administración" en Canvas matchee con "administracion", "ADMINISTRACIÓN"
    # o "Administración " (con espacio de más) en la planilla.
    subaccount_map: dict[str, int] = {}
    try:
        subaccounts_data = await canvas.paginate(
            f"/accounts/{_ACCOUNT_LOCAL}/sub_accounts", params={"recursive": "true", "per_page": 100}
        )
        for sa in subaccounts_data:
            sa_name = _norm(str(sa.get("name") or ""))
            if sa_name:
                subaccount_map[sa_name] = sa.get("id")
    except Exception:
        pass

    coordinador_role_id = await _get_coordinador_role_id()

    sheet_lower = req.sheet_name.lower()
    default_plat = "both"
    if "canvas" in sheet_lower and "teams" not in sheet_lower:
        default_plat = "canvas"
    elif "teams" in sheet_lower and "canvas" not in sheet_lower:
        default_plat = "teams"

    async def create_course_row(r_idx):
        nombre = str(ws.cell(row=r_idx, column=col_nombre).value or "").strip()
        sis_id = str(ws.cell(row=r_idx, column=col_sis).value or "").strip() if col_sis else ""
        periodo = str(ws.cell(row=r_idx, column=col_periodo).value or "").strip() if col_periodo else ""
        subcuenta_raw = str(ws.cell(row=r_idx, column=col_subcuenta).value or "").strip() if col_subcuenta else ""
        docente_email = str(ws.cell(row=r_idx, column=col_docente).value or "").strip() if col_docente else ""
        if docente_email and "@" not in docente_email:
            docente_email = ""

        coordinadores_raw = str(ws.cell(row=r_idx, column=col_coordinadores).value or "").strip() if col_coordinadores else ""
        coordinador_emails = [
            e.strip() for e in re.split(r'[,;/\n]+', coordinadores_raw) if e.strip() and "@" in e.strip()
        ]

        if not nombre:
            return

        existing_estado = str(ws.cell(row=r_idx, column=col_estado).value or "").strip()
        if "✅" in existing_estado:
            return

        error_canvas = None
        error_teams = None
        error_docente_canvas = None
        error_docente_teams = None
        canvas_id = None
        teams_id = None
        docente_canvas_ok = False
        docente_teams_ok = False
        coordinadores_canvas_ok: list[str] = []
        coordinadores_canvas_failed: list[str] = []
        coordinadores_teams_ok: list[str] = []
        coordinadores_teams_failed: list[str] = []

        # course_code_str combina nombre+período — se usa SOLO para derivar el
        # mailNickname del equipo de Teams (necesita distinguir períodos, ver
        # search_group_by_name_and_nickname_prefix). El código del curso en
        # Canvas (course_code) debe ser igual al nombre del curso, sin el
        # período pegado atrás — eso es lo que se pedía corregir.
        course_code_str = f"{nombre} {periodo}".strip()

        # 1. Canvas Creation
        if default_plat in ("canvas", "both"):
            # Si se indicó una subcuenta, el curso se crea DIRECTAMENTE ahí
            # (ej. Administración/Marketing/Negocios) en vez de en la cuenta
            # raíz + período — reemplaza al período, no se combinan. Si el
            # nombre no matchea ninguna subcuenta real, se corta acá con un
            # error claro en vez de crear el curso en el lugar equivocado.
            target_account_id = _ACCOUNT_LOCAL
            subcuenta_error = None
            if subcuenta_raw:
                sub_id = subaccount_map.get(_norm(subcuenta_raw))
                if sub_id:
                    target_account_id = sub_id
                else:
                    subcuenta_error = f"No se encontró la subcuenta '{subcuenta_raw}' en Canvas."

            # Si ya existe un curso con este SIS ID exacto (el SIS ID ya es
            # específico de nombre+período por diseño), lo reutilizamos en
            # vez de intentar crear uno duplicado.
            if sis_id and sis_id != "None":
                try:
                    existing = await canvas.get(f"/courses/sis_course_id:{sis_id}")
                    if existing:
                        canvas_id = existing.get("id")
                except Exception:
                    pass

            try:
                if subcuenta_error:
                    raise ValueError(subcuenta_error)
                if not canvas_id:
                    payload = {
                        "course": {
                            "name": nombre,
                            "course_code": nombre,
                        }
                    }
                    if sis_id and sis_id != "None":
                        payload["course"]["sis_course_id"] = sis_id
                    if periodo and not subcuenta_raw:
                        term_val = f"sis_term_id:{periodo}"
                        if periodo.isdigit():
                            term_val = periodo
                        else:
                            p_lower = periodo.strip().lower()
                            for t in terms_data:
                                t_name = str(t.get("name") or "").strip().lower()
                                t_sis = str(t.get("sis_term_id") or "").strip().lower()
                                if t_name == p_lower or str(t.get("id")) == periodo or t_sis == p_lower:
                                    term_val = t.get("id")
                                    break
                        payload["course"]["term_id"] = term_val

                    data = await canvas.post(f"/accounts/{target_account_id}/courses", payload)
                    canvas_id = data.get("id")
            except Exception as e:
                err_text = getattr(e, "detail", None) or str(e)
                # Diagnóstico: si el rechazo es por SIS ID duplicado, casi
                # siempre es un curso eliminado (soft-delete) que retiene el
                # SIS ID sin aparecer en los listados normales. Se identifica
                # el curso real que lo está bloqueando para no tener que
                # buscarlo a mano — NO se reutiliza automáticamente.
                if sis_id and sis_id != "None" and "sis" in err_text.lower() and "already" in err_text.lower():
                    try:
                        blocking = await canvas.get(f"/courses/sis_course_id:{sis_id}")
                        if blocking:
                            err_text += (
                                f" — El SIS ID '{sis_id}' ya está en uso por el curso ID {blocking.get('id')} "
                                f"'{blocking.get('name')}' (estado: {blocking.get('workflow_state')}). "
                                f"Si ese curso está eliminado y no lo necesitás, hay que liberar el SIS ID "
                                f"desde Canvas (o asignarle un SIS ID distinto a esta fila)."
                            )
                    except Exception:
                        pass
                error_canvas = err_text

            # 1b. Inscribir al docente indicado como Teacher del curso recién creado
            if canvas_id and docente_email:
                try:
                    match = await canvas.find_user_exact(_ACCOUNT_LOCAL, docente_email, fields=("email", "login_id"))
                    docente_canvas_id = match["id"] if match else None
                    if not docente_canvas_id:
                        error_docente_canvas = f"No se encontró el usuario '{docente_email}' en Canvas"
                    else:
                        await canvas.post(f"/courses/{canvas_id}/enrollments", {
                            "enrollment": {
                                "user_id": docente_canvas_id,
                                "type": "TeacherEnrollment",
                                "enrollment_state": "active",
                                "notify": False,
                            }
                        })
                        docente_canvas_ok = True
                except Exception as e:
                    error_docente_canvas = str(e)

            # 1c. Inscribir a cada coordinador/diseñador como DesignerEnrollment del curso
            if canvas_id and coordinador_emails:
                async def _enroll_coordinador_canvas(email: str):
                    try:
                        match = await canvas.find_user_exact(_ACCOUNT_LOCAL, email, fields=("email", "login_id"))
                        coord_canvas_id = match["id"] if match else None
                        if not coord_canvas_id:
                            coordinadores_canvas_failed.append(f"{email} (no encontrado en Canvas)")
                            return
                        enrollment_payload = {
                            "user_id": coord_canvas_id,
                            "enrollment_state": "active",
                            "notify": False,
                        }
                        if coordinador_role_id:
                            enrollment_payload["role_id"] = coordinador_role_id
                        else:
                            enrollment_payload["type"] = "DesignerEnrollment"
                        await canvas.post(f"/courses/{canvas_id}/enrollments", {"enrollment": enrollment_payload})
                        coordinadores_canvas_ok.append(email)
                    except Exception as e:
                        coordinadores_canvas_failed.append(f"{email} ({e})")

                await asyncio.gather(*(_enroll_coordinador_canvas(e) for e in coordinador_emails))

        # 2. Teams Creation
        if default_plat in ("teams", "both"):
            docente_azure_id = None
            if docente_email:
                try:
                    found = await graph.get_user_by_upn_exact(docente_email)
                    docente_azure_id = found["id"] if found else None
                    if not docente_azure_id:
                        error_docente_teams = f"No se encontró el usuario '{docente_email}' en Microsoft 365"
                except Exception as e:
                    error_docente_teams = str(e)

            coordinador_azure_ids: list[str] = []
            if coordinador_emails:
                async def _resolve_coordinador_teams(email: str):
                    try:
                        found = await graph.get_user_by_upn_exact(email)
                        azure_id = found["id"] if found else None
                        if azure_id:
                            coordinador_azure_ids.append(azure_id)
                            coordinadores_teams_ok.append(email)
                        else:
                            coordinadores_teams_failed.append(f"{email} (no encontrado en Microsoft 365)")
                    except Exception as e:
                        coordinadores_teams_failed.append(f"{email} ({e})")

                await asyncio.gather(*(_resolve_coordinador_teams(e) for e in coordinador_emails))

            try:
                # base_nickname identifica nombre+período (sin el sufijo random
                # de unicidad). Si ya existe un equipo con el mismo displayName
                # Y cuyo mailNickname arranca con esta base, es el MISMO curso
                # del MISMO período — se reutiliza. Si el nombre coincide pero
                # el mailNickname no (otro período con la misma materia), NO
                # se reutiliza: se crea un equipo nuevo para este período.
                base_nickname = _safe_mail_nickname(course_code_str)
                teams_id = await graph.search_group_by_name_and_nickname_prefix(nombre, base_nickname)

                owner_ids = []
                try:
                    admin_user = await graph.get(f"/users/resteche@usil.edu.py", params={"$select": "id"})
                    if admin_user and admin_user.get("id"):
                        owner_ids.append(admin_user["id"])
                except: pass
                if docente_azure_id and docente_azure_id not in owner_ids:
                    owner_ids.append(docente_azure_id)
                for coord_id in coordinador_azure_ids:
                    if coord_id not in owner_ids:
                        owner_ids.append(coord_id)

                if not teams_id:
                    nickname = f"{base_nickname}{int(time.time() * 1000) % 100000}"[:64]
                    new_team = await create_team_via_group(
                        display_name=nombre,
                        mail_nickname=nickname,
                        description=f"Grupo para {nombre}",
                        visibility="Private",
                        owner_ids=owner_ids
                    )
                    teams_id = new_team.get("id")
                else:
                    # Equipo reutilizado: asegurar que docente/coordinadores
                    # queden como owners aunque ya existiera (ignorando "ya es owner").
                    for oid in owner_ids:
                        try:
                            await graph.post(f"/groups/{teams_id}/owners/$ref", {
                                "@odata.id": f"https://graph.microsoft.com/v1.0/directoryObjects/{oid}"
                            })
                        except Exception:
                            pass

                if docente_azure_id and teams_id:
                    docente_teams_ok = True
            except Exception as e:
                error_teams = str(e)

            if not teams_id and coordinadores_teams_ok:
                # El equipo no se creó — ningún coordinador quedó como owner realmente,
                # aunque se hayan resuelto sus cuentas en Microsoft 365.
                coordinadores_teams_failed.extend(f"{e} (equipo no creado)" for e in coordinadores_teams_ok)
                coordinadores_teams_ok = []

        # Update Excel
        ws.cell(row=r_idx, column=col_fecha, value=datetime.now().strftime("%d/%m/%Y"))
        
        if canvas_id:
            ws.cell(row=r_idx, column=col_canvas_id, value=canvas_id)
        if teams_id:
            ws.cell(row=r_idx, column=col_teams_id, value=teams_id)
            
        final_errors = []
        if error_canvas: final_errors.append(f"Canvas: {error_canvas}")
        if error_teams: final_errors.append(f"Teams: {error_teams}")

        docente_warnings = []
        if docente_email:
            if error_docente_canvas: docente_warnings.append(f"Docente/Canvas: {error_docente_canvas}")
            if error_docente_teams: docente_warnings.append(f"Docente/Teams: {error_docente_teams}")
        if coordinadores_canvas_failed:
            docente_warnings.append(f"Coordinadores/Canvas: {', '.join(coordinadores_canvas_failed)}")
        if coordinadores_teams_failed:
            docente_warnings.append(f"Coordinadores/Teams: {', '.join(coordinadores_teams_failed)}")

        succeeded_entry = {
            "nombre": nombre,
            "canvas_id": canvas_id,
            "teams_id": teams_id,
            "sis_course_id": sis_id,
        }
        if docente_email:
            succeeded_entry["docente_email"] = docente_email
            succeeded_entry["docente_canvas_enrolled"] = docente_canvas_ok
            succeeded_entry["docente_teams_enrolled"] = docente_teams_ok
        if coordinador_emails:
            succeeded_entry["coordinadores_canvas_enrolled"] = coordinadores_canvas_ok
            succeeded_entry["coordinadores_teams_enrolled"] = coordinadores_teams_ok

        if not final_errors:
            if docente_warnings:
                ws.cell(row=r_idx, column=col_estado, value=f"✅ OK ⚠️ {' | '.join(docente_warnings)}")
                ws.cell(row=r_idx, column=col_estado).font = Font(color="D97706", bold=True)
            else:
                ws.cell(row=r_idx, column=col_estado, value="✅ OK")
                ws.cell(row=r_idx, column=col_estado).font = Font(color="00B050", bold=True)
            result.succeeded.append(succeeded_entry)
        else:
            error_msg = " | ".join(final_errors + docente_warnings)
            ws.cell(row=r_idx, column=col_estado, value=f"⚠️ {error_msg}")
            ws.cell(row=r_idx, column=col_estado).font = Font(color="FF0000", bold=True)
            result.failed.append({"input": {"nombre": nombre}, "error": error_msg})

    row_idxs = list(range(header_row_idx + 1, ws.max_row + 1))
    total_rows = len(row_idxs)

    batch_size = 3
    for i in range(0, total_rows, batch_size):
        chunk = row_idxs[i:i + batch_size]
        await asyncio.gather(*(create_course_row(r) for r in chunk))
        await jobs.update_job_progress(
            job_id,
            len(result.succeeded),
            len(result.failed),
            data_json=f'{{"total_to_process": {total_rows}, "processed": {len(result.succeeded) + len(result.failed)}}}'
        )

    out_io = io.BytesIO()
    wb.save(out_io)
    out_io.seek(0)

    try:
        await graph.put_raw(f"/shares/{encoded_url}/driveItem/content", out_io.read())
        save_note = ""
    except Exception as e:
        save_note = f" (⚠️ no se pudo guardar el archivo actualizado en OneDrive: {e})"

    if req.report_url and (result.succeeded or result.failed):
        await append_report_onedrive(req.report_url, result.succeeded, result.failed)

    await jobs.complete_job(
        job_id,
        len(result.succeeded),
        len(result.failed),
        f"{len(result.succeeded)} cursos creados, {len(result.failed)} fallaron.{save_note}"
    )




@router.post("/excel/egreso/preview", summary="Pre-visualizar planilla de Egreso/Eliminación")
async def preview_egreso_onedrive(req: DiplomadosUrlRequest) -> PreviewResponse:
    """Previsualiza los usuarios que se darán de baja leyendo la planilla."""
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")
    
    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo de OneDrive. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail="El archivo descargado no es un Excel válido.")

    if req.sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"La pestaña '{req.sheet_name}' no existe.")

    ws = wb[req.sheet_name]
    
    sample_rows = []
    
    header_row_idx, headers_dict, headers_raw = _find_header_row_and_headers(ws)
    
    if not header_row_idx:
        raise HTTPException(status_code=400, detail="No se encontró la fila de encabezados.")

    headers = [h for h in headers_raw if h]

    def get_col_idx(*keys):
        return _match_col_idx(headers_dict, *keys)

    col_nombre = get_col_idx("nombre", "alumno", "estudiante")
    col_correo = get_col_idx("correo", "email")
    col_cedula = get_col_idx("cedula", "cédula", "ci", "documento", "dni")
    # Solo "desvinculado"/"baja" — NO genérico "estado"/"enviado": si la misma
    # hoja se usó antes para CREAR a estos usuarios, ya trae su propia columna
    # "Enviado/Estado" (de la creación, con "✅ OK"). Buscar genéricamente
    # "estado"/"enviado" la confundiría con el estado de egreso y marcaría a
    # todos como ya dados de baja sin haberlo hecho nunca.
    col_enviado = get_col_idx("desvinculado", "dado de baja")

    if not col_cedula and not col_correo:
        raise HTTPException(status_code=400, detail="Falta la columna de Cédula o Correo para identificar a los usuarios.")

    students_to_process = 0
    students_already_processed = 0
    student_details = []

    for row_idx in range(header_row_idx + 1, ws.max_row + 1):
        row_vals = [str(ws.cell(row=row_idx, column=c).value or "").strip() for c in range(1, len(headers_raw) + 1)]

        if not any(row_vals):
            continue

        nombre = row_vals[col_nombre - 1] if col_nombre else ""
        correo = row_vals[col_correo - 1] if col_correo else ""
        cedula = _clean_cedula(row_vals[col_cedula - 1]) if col_cedula else ""
        enviado = (row_vals[col_enviado - 1] if col_enviado else "").lower()

        if not cedula and not correo:
            continue

        if "ok" in enviado or "eliminado" in enviado or "baja" in enviado or "enviado" in enviado:
            students_already_processed += 1
        else:
            students_to_process += 1
            # Sin límite de muestra: esta acción da de baja/elimina cuentas de
            # forma permanente, así que conviene poder revisar la lista
            # completa antes de confirmar, no solo los primeros 3.
            student_details.append({"nombre": nombre or cedula or correo, "correo": correo or cedula})
            sample_rows.append({h: v for h, v in zip(headers_raw, row_vals) if h})

    return PreviewResponse(
        sheet_name=req.sheet_name,
        students_to_process=students_to_process,
        students_already_processed=students_already_processed,
        headers=headers[:min(5, len(headers))],
        sample_rows=sample_rows,
        student_details=student_details,
    )

@router.post("/excel/egreso", summary="Procesar planilla de Egreso/Eliminación")
async def import_egreso_onedrive(req: DiplomadosUrlRequest) -> BulkResult:
    try:
        return await _import_egreso_onedrive_inner(req)
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        tb = traceback.format_exc()
        print(tb)
        raise HTTPException(status_code=500, detail=f"Error interno: {str(e)} | Trace: {tb}")

async def _import_egreso_onedrive_inner(req: DiplomadosUrlRequest) -> BulkResult:
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")
    
    encoded_url = _encode_share_url(req.url)
    
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo. {e}")

    try:
        import io
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(contents))
    except Exception as e:
        raise HTTPException(status_code=400, detail="El archivo no es un Excel válido.")

    result = BulkResult()

    if req.sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"La pestaña '{req.sheet_name}' no existe.")

    ws = wb[req.sheet_name]
    
    header_row_idx, headers, _ = _find_header_row_and_headers(ws)

    
    if not header_row_idx:
        raise HTTPException(status_code=400, detail="No se encontraron cabeceras.")

    def get_col_idx(*keys):
        return _match_col_idx(headers, *keys)

    col_nombre = get_col_idx("nombre", "alumno", "estudiante")
    col_correo = get_col_idx("correo", "email")
    col_cedula = get_col_idx("cedula", "cédula", "ci", "documento", "dni")
    # Solo "desvinculado"/"baja" — ver comentario equivalente en el preview:
    # evita reutilizar la columna "Enviado/Estado" que la misma hoja pudo
    # traer de un proceso de creación anterior.
    col_enviado = get_col_idx("desvinculado", "dado de baja")

    col_cc = get_col_idx("cc", "copia")
    sheet_cc_list = []
    if col_cc:
        for r_idx in range(header_row_idx + 1, ws.max_row + 1):
            cc_val = str(ws.cell(row=r_idx, column=col_cc).value or "").strip()
            if cc_val:
                for email in cc_val.replace(";", ",").replace("\n", ",").split(","):
                    email = email.strip()
                    if "@" in email and email not in sheet_cc_list:
                        sheet_cc_list.append(email)
    col_usuario = get_col_idx("usuario")

    if not col_cedula and not col_correo:
        raise HTTPException(status_code=400, detail="Falta la columna de Cédula o Correo para identificar a los usuarios.")

    if not col_enviado:
        from openpyxl.styles import Font
        col_enviado = ws.max_column + 1
        ws.cell(row=header_row_idx, column=col_enviado, value="Desvinculado").font = Font(bold=True)

    users_to_process = []
    for row_idx in range(header_row_idx + 1, ws.max_row + 1):
        nombre = str(ws.cell(row=row_idx, column=col_nombre).value or "").strip() if col_nombre else ""
        correo = str(ws.cell(row=row_idx, column=col_correo).value or "").strip() if col_correo else ""
        cedula = _clean_cedula(str(ws.cell(row=row_idx, column=col_cedula).value or "").strip()) if col_cedula else ""
        enviado = str(ws.cell(row=row_idx, column=col_enviado).value or "").strip().lower()
        usuario_val = str(ws.cell(row=row_idx, column=col_usuario).value or "").strip() if col_usuario else ""

        if not cedula and not correo:
            continue

        if "ok" in enviado or "eliminado" in enviado or "baja" in enviado or "enviado" in enviado:
            continue

        users_to_process.append({
            "r_idx": row_idx,
            "correo": correo,
            "usuario": usuario_val,
            "nombre": nombre or cedula or correo,
            "cedula": cedula,
        })

    if len(users_to_process) > 50:
        raise HTTPException(status_code=400, detail=f"Demasiados registros nuevos pendientes ({len(users_to_process)}). Máximo 50 por ejecución.")

    if len(users_to_process) > 0:
        try:
            # Verificar si el archivo está bloqueado
            await graph.put_raw(f"/shares/{encoded_url}/driveItem/content", contents)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"El archivo Excel está abierto o el enlace es de Solo Lectura. Detalle real: {e}")

    for user_data in users_to_process:
        correo = user_data["correo"]
        usuario_upn = user_data["usuario"]
        cedula = user_data["cedula"]
        nombre = user_data["nombre"]
        r_idx = user_data["r_idx"]

        canvas_email = ""
        error = ""

        # 1. Canvas Delete — por cédula (SIS ID) primero, igual que el egreso
        # individual: es más confiable que buscar por nombre/correo, que
        # puede no coincidir exactamente con lo cargado en Canvas.
        try:
            if cedula:
                c_user = await canvas.get(f"/accounts/{settings.canvas_account_id}/users/sis_user_id:{cedula}")
                canvas_email = c_user.get("email") or ""
                await canvas.delete(f"/accounts/{settings.canvas_account_id}/users/{c_user['id']}")
            else:
                raise ValueError("sin_cedula")
        except Exception as e:
            if cedula and ("404" not in str(e)):
                error = f"Error Canvas: {str(e)}"
            else:
                # Sin cédula, o no encontrado por SIS ID: fallback por búsqueda
                # de texto. Es una acción DESTRUCTIVA (borra la cuenta), así
                # que nunca se toma "el primero que aparezca" a ciegas — sólo
                # se borra si hay una coincidencia de nombre EXACTA (sin
                # tildes/mayúsculas) entre los candidatos, para no eliminar
                # la cuenta de otra persona con nombre/correo parecido.
                search_term = usuario_upn if (usuario_upn and "@" in usuario_upn) else (correo or nombre)
                try:
                    candidatos = await canvas.get(f"/accounts/{settings.canvas_account_id}/users", params={"search_term": search_term})
                    match = next((u for u in (candidatos or []) if _norm(u.get("name") or "") == _norm(nombre)), None) if nombre else None
                    if match:
                        canvas_email = match.get("email") or ""
                        await canvas.delete(f"/accounts/{settings.canvas_account_id}/users/{match['id']}")
                    elif candidatos:
                        error = f"Usuario no encontrado en Canvas: hay {len(candidatos)} candidato(s) por búsqueda de texto pero ninguno coincide exacto con el nombre — revisar manualmente"
                    else:
                        error = "Usuario no encontrado en Canvas"
                except Exception as e2:
                    error = "Usuario no encontrado en Canvas" if "404" in str(e2) else f"Error Canvas: {str(e2)}"

        # 2. Azure AD Disable or Delete
        search_term = usuario_upn if (usuario_upn and "@" in usuario_upn) else (correo or canvas_email or nombre)
        try:
            if "@" in search_term:
                found = await graph.get_user_by_upn_exact(search_term)
                target_id = found["id"] if found else None
            else:
                # Mismo criterio que arriba: sin correo/UPN conocido, sólo se
                # actúa sobre un candidato de búsqueda de texto si su
                # displayName coincide EXACTO con el nombre esperado.
                candidatos = await graph.search_users(search_term)
                match = next((u for u in (candidatos or []) if _norm(u.get("displayName") or "") == _norm(nombre)), None) if nombre else None
                target_id = match["id"] if match else None
                if not target_id and candidatos:
                    error = (error + " | " if error else "") + f"Azure AD: hay {len(candidatos)} candidato(s) por búsqueda de texto pero ninguno coincide exacto con el nombre — revisar manualmente"

            if target_id:
                if req.delete_account:
                    await graph.delete(f"/users/{target_id}")
                else:
                    await graph.patch(f"/users/{target_id}", {"accountEnabled": False})
                    try:
                        await graph.remove_all_licenses(target_id)
                    except Exception:
                        pass
            elif not error:
                error = "No encontrado en Azure AD"
        except Exception as e:
            error = error + f" | Error Azure: {str(e)}" if error else f"Error Azure: {str(e)}"
        
        if error:
            ws.cell(row=r_idx, column=col_enviado, value=f"Error: {error}")
            result.failed.append({"correo": correo, "error": error})
        else:
            status_text = "OK (Eliminado permanentemente)" if req.delete_account else "OK (Deshabilitado)"
            ws.cell(row=r_idx, column=col_enviado, value=status_text)
            result.succeeded.append({"correo": correo})

    # Guardar y subir
    if len(users_to_process) > 0:
        out_io = io.BytesIO()
        wb.save(out_io)
        out_io.seek(0)
        try:
            await graph.put_raw(f"/shares/{encoded_url}/driveItem/content", out_io.read())
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"No se pudo guardar el archivo actualizado en OneDrive. {e}")

    return result


@router.post("/excel/docentes-onedrive/preview", summary="Previsualizar Docentes desde OneDrive")
async def preview_docentes_onedrive(req: DiplomadosUrlRequest) -> DocentesPreviewResponse:
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")

    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar de OneDrive. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents))
    except Exception as e:
        raise HTTPException(status_code=400, detail="El archivo no es un Excel válido.")

    if req.sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"La pestaña '{req.sheet_name}' no existe.")

    ws = wb[req.sheet_name]

    header_row_idx, headers, _ = _find_header_row_and_headers(ws)


    if not header_row_idx:
        raise HTTPException(status_code=400, detail="No se encontraron las columnas de 'Nombre' y 'Cédula'.")

    def get_col_idx(*keys):
        return _match_col_idx(headers, *keys)

    col_nombre = get_col_idx("nombre", "alumno", "estudiante")
    col_cedula = get_col_idx("cedula", "cdula", "ci")
    col_correo = get_col_idx("correo", "email")
    col_plat = get_col_idx("plataforma")
    col_curso = get_col_idx("curso", "id curso", "canvas")
    col_equipo = get_col_idx("equipo", "id equipo", "teams")
    col_curso_nombre = get_col_idx("nombre del curso", "curso", "diplomado")
    col_enviado = get_col_idx("estado", "enviado")
    col_usuario = get_col_idx("usuario")

    col_cc = get_col_idx("cc", "copia")
    sheet_cc_list = []
    if col_cc:
        for r_idx in range(header_row_idx + 1, ws.max_row + 1):
            cc_val = str(ws.cell(row=r_idx, column=col_cc).value or "").strip()
            if cc_val:
                for email in cc_val.replace(";", ",").replace("\n", ",").split(","):
                    email = email.strip()
                    if "@" in email and email not in sheet_cc_list:
                        sheet_cc_list.append(email)

    rows = []
    already_processed = []
    for r_idx in range(header_row_idx + 1, ws.max_row + 1):
        nombre = str(ws.cell(row=r_idx, column=col_nombre).value or "").strip()
        cedula = _clean_cedula(str(ws.cell(row=r_idx, column=col_cedula).value or "").strip())

        if not nombre or not cedula or cedula == "None":
            continue

        correo = str(ws.cell(row=r_idx, column=col_correo).value or "").strip() if col_correo else ""
        plat = str(ws.cell(row=r_idx, column=col_plat).value or "both").strip().lower() if col_plat else "both"
        id_curso = str(ws.cell(row=r_idx, column=col_curso).value or "").strip() if col_curso else ""
        id_equipo = str(ws.cell(row=r_idx, column=col_equipo).value or "").strip() if col_equipo else ""

        usuario_preview = str(ws.cell(row=r_idx, column=col_usuario).value or "").strip() if col_usuario else ""
        enviado = str(ws.cell(row=r_idx, column=col_enviado).value or "").strip() if col_enviado else ""
        enviado_lower = enviado.lower()
        if col_enviado and ("✅" in enviado or enviado_lower in ["si", "yes", "true", "enviado", "ok"]
                or "creado ok" in enviado_lower or "ya exist" in enviado_lower
                or (usuario_preview and "@" in usuario_preview)):
            already_processed.append({
                "nombre": nombre, "cedula": cedula, "usuario": usuario_preview, "estado": enviado,
            })
            continue

        rows.append({
            "nombre": nombre,
            "cedula": cedula,
            "correo": correo,
            "plataforma": plat,
            "curso": id_curso,
            "equipo": id_equipo
        })

    return DocentesPreviewResponse(
        total_rows=ws.max_row - header_row_idx,
        valid_rows=len(rows),
        headers=list(headers.keys()),
        sample_rows=rows,
        already_processed_details=already_processed,
    )

@router.post("/excel/docentes-onedrive", summary="Alta Docentes OneDrive")
async def import_docentes_onedrive(req: DiplomadosUrlRequest) -> BulkResult:
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")

    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo de OneDrive. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents))
    except Exception as e:
        raise HTTPException(status_code=400, detail="El archivo descargado no es un Excel válido.")

    _ACCOUNT_LOCAL = settings.canvas_account_id
    result = BulkResult()

    if req.sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"La pestaña '{req.sheet_name}' no existe.")

    ws = wb[req.sheet_name]

    header_row_idx, headers, _ = _find_header_row_and_headers(ws)


    if not header_row_idx:
        raise HTTPException(status_code=400, detail="Columnas de Nombre y Cédula no encontradas.")

    def get_col_idx(*keys):
        return _match_col_idx(headers, *keys)

    col_nombre = get_col_idx("nombre", "alumno", "estudiante")
    col_cedula = get_col_idx("cedula", "cdula", "ci")
    col_correo = get_col_idx("correo", "email")
    col_plat = get_col_idx("plataforma")
    col_curso = get_col_idx("curso", "id curso", "canvas")
    col_equipo = get_col_idx("equipo", "id equipo", "teams")
    # Nota: "curso" NO va en las claves de este fallback — si sólo existe una
    # columna "Curso" (sin una columna aparte de "Nombre del Curso"/"Diplomado"),
    # col_curso_nombre debe quedar vacío en vez de apuntar a la misma columna
    # que col_curso; si no, la validación de más abajo compara el ID contra sí
    # mismo como si fuera un nombre y siempre falla.
    col_curso_nombre = get_col_idx("nombre del curso", "diplomado")

    col_usuario = get_col_idx("usuario")
    col_contra = get_col_idx("contrasena", "contrasea", "clave")
    col_enviado = get_col_idx("estado", "enviado")

    col_cc = get_col_idx("cc", "copia")
    sheet_cc_list = []
    if col_cc:
        for r_idx in range(header_row_idx + 1, ws.max_row + 1):
            cc_val = str(ws.cell(row=r_idx, column=col_cc).value or "").strip()
            if cc_val:
                for email in cc_val.replace(";", ",").replace("\n", ",").split(","):
                    email = email.strip()
                    if "@" in email and email not in sheet_cc_list:
                        sheet_cc_list.append(email)

    next_col = ws.max_column + 1
    if not col_usuario:
        col_usuario = next_col; ws.cell(row=header_row_idx, column=col_usuario, value="Usuario").font = Font(bold=True); next_col += 1
    if not col_contra:
        col_contra = next_col; ws.cell(row=header_row_idx, column=col_contra, value="Contrasea").font = Font(bold=True); next_col += 1
    if not col_enviado:
        col_enviado = next_col; ws.cell(row=header_row_idx, column=col_enviado, value="Estado").font = Font(bold=True); next_col += 1

    sheet_lower = req.sheet_name.lower()
    default_plat = "both"
    if "canvas" in sheet_lower and "teams" not in sheet_lower:
        default_plat = "canvas"
    elif "teams" in sheet_lower and "canvas" not in sheet_lower:
        default_plat = "teams"

    users_to_process = []
    for r_idx in range(header_row_idx + 1, ws.max_row + 1):
        nombre = str(ws.cell(row=r_idx, column=col_nombre).value or "").strip()
        cedula = _clean_cedula(str(ws.cell(row=r_idx, column=col_cedula).value or "").strip())
        
        if not nombre or not cedula or cedula == "None":
            continue
            
        enviado = str(ws.cell(row=r_idx, column=col_enviado).value or "").strip()
        enviado_lower = enviado.lower()
        usuario_val = str(ws.cell(row=r_idx, column=col_usuario).value or "").strip() if col_usuario else ""
        if ("✅" in enviado or enviado_lower in ["si", "yes", "true", "enviado", "ok"]
                or "creado ok" in enviado_lower or "ya exist" in enviado_lower
                or (usuario_val and "@" in usuario_val)):
            continue

        users_to_process.append(r_idx)

    for r_idx in users_to_process:
        nombre = str(ws.cell(row=r_idx, column=col_nombre).value or "").strip()
        cedula = _clean_cedula(str(ws.cell(row=r_idx, column=col_cedula).value or "").strip())
        correo = str(ws.cell(row=r_idx, column=col_correo).value or "").strip() if col_correo else ""
        plat = str(ws.cell(row=r_idx, column=col_plat).value or default_plat).strip().lower() if col_plat else default_plat
        id_curso = str(ws.cell(row=r_idx, column=col_curso).value or "").strip() if col_curso else ""
        id_equipo = str(ws.cell(row=r_idx, column=col_equipo).value or "").strip() if col_equipo else ""
        curso_nombre = str(ws.cell(row=r_idx, column=col_curso_nombre).value or "").strip() if col_curso_nombre else ""

        # Planillas solo-Teams (ej. "Usuarios Docentes (Teams)") suelen tener
        # una única columna "Curso" y ninguna columna "Equipo" dedicada. En
        # ese caso el ID que se cargó en "Curso" es el ID del equipo de Teams
        # al que hay que matricular directamente al docente.
        if not id_equipo and not col_equipo and id_curso and id_curso != "None" and plat == "teams":
            id_equipo = id_curso

        creds, status = await user_service.generate_unique_credentials(nombre, cedula, plat)
        login_id = creds["email"]
        pwd = creds["password"]
        
        entry = {"cedula": cedula, "nombre": creds["full_name"], "login_id": login_id}
        error = ""
        parts = creds["full_name"].strip().split()

        # Azure AD
        azure_id = None
        if plat in ("teams", "both"):
            try:
                au_payload = {
                    "displayName": creds["full_name"],
                    "givenName": parts[0],
                    "surname": " ".join(parts[1:]) if len(parts) > 1 else "",
                    "userPrincipalName": login_id,
                    "mailNickname": login_id.replace(".", "_").replace("@", "_"),
                    "usageLocation": settings.usage_location,
                    "accountEnabled": True,
                    "jobTitle": "Docente",
                    "country": "Paraguay",
                    "passwordProfile": {
                        "forceChangePasswordNextSignIn": False,
                        "password": pwd,
                    },
                }
                if cedula:
                    au_payload["postalCode"] = cedula
                au = await graph.post("/users", au_payload)
                azure_id = au["id"]
                await graph.assign_license(azure_id, settings.azure_sku_teachers)
                entry["teams"] = "creado"
            except Exception as e:
                if "already exists" in str(e).lower() or "Request_BadRequest" in str(e):
                    # Recuperar la cuenta existente por UPN EXACTO (no
                    # search_users, que es un $search difuso y puede traer a
                    # otra persona con nombre/correo parecido en vez del
                    # login_id que efectivamente ya existe).
                    existing = await graph.get_user_by_upn_exact(login_id)
                    if existing:
                        azure_id = existing["id"]
                        entry["teams"] = "exista"
                else:
                    error += f"Teams: {str(e)} | "

        # Validation: Avoid Copy-Paste Errors
        if id_curso and id_curso != "None" and curso_nombre:
            cn = await canvas.get_course_name_by_id(id_curso)
            if cn and cn.strip().lower() != curso_nombre.strip().lower():
                error += f"Error de Validación: El ID Curso {id_curso} pertenece a '{cn}' y no coincide con '{curso_nombre}' | "
                
        if id_equipo and id_equipo != "None" and curso_nombre:
            gn = await graph.get_group_name_by_id(id_equipo)
            if gn and gn.strip().lower() != curso_nombre.strip().lower():
                error += f"Error de Validación: El ID Equipo {id_equipo} pertenece a '{gn}' y no coincide con '{curso_nombre}' | "

        if error:
            # Skip enrollment and creation if validation fails
            ws.cell(row=r_idx, column=col_enviado, value=f"⚠️ Error: {error}")
            ws.cell(row=r_idx, column=col_enviado).font = Font(color="D97706", bold=True)
            result.failed.append({"correo": str(ws.cell(row=r_idx, column=col_correo).value), "error": error})
            continue

        # Bidirectional Canvas Logic (Name <-> ID)
        if plat in ("canvas", "both"):
            if not id_curso and curso_nombre:
                try:
                    existing_cid = await canvas.search_course_by_name(_ACCOUNT_LOCAL, curso_nombre)
                    if existing_cid:
                        id_curso = existing_cid
                    else:
                        id_curso = await canvas.create_course(_ACCOUNT_LOCAL, curso_nombre)
                    if id_curso and col_curso:
                        ws.cell(row=r_idx, column=col_curso, value=id_curso)
                except Exception as e:
                    error += f"CanvasCourseLogic: {str(e)} | "
            elif id_curso and id_curso != "None" and not curso_nombre:
                try:
                    cn = await canvas.get_course_name_by_id(id_curso)
                    if cn:
                        curso_nombre = cn
                        if col_curso_nombre:
                            ws.cell(row=r_idx, column=col_curso_nombre, value=cn)
                except Exception as e:
                    pass

        # Bidirectional Teams Logic (Name <-> ID)
        if plat in ("teams", "both"):
            if not id_equipo and curso_nombre:
                try:
                    existing_tid = await graph.search_group_by_name(curso_nombre)
                    if existing_tid:
                        id_equipo = existing_tid
                    else:
                        nickname = _safe_mail_nickname(curso_nombre)
                        owner_ids = [azure_id] if azure_id else []
                        new_team = await graph.create_team_via_group(
                            display_name=curso_nombre,
                            mail_nickname=nickname,
                            description=f"Grupo para {curso_nombre}",
                            visibility="Private",
                            owner_ids=owner_ids
                        )
                        id_equipo = new_team.get("id")
                    if id_equipo and col_equipo:
                        ws.cell(row=r_idx, column=col_equipo, value=id_equipo)
                except Exception as e:
                    error += f"TeamsGroupLogic: {str(e)} | "
            elif id_equipo and id_equipo != "None" and not curso_nombre:
                try:
                    gn = await graph.get_group_name_by_id(id_equipo)
                    if gn:
                        curso_nombre = gn
                        if col_curso_nombre:
                            ws.cell(row=r_idx, column=col_curso_nombre, value=gn)
                except Exception as e:
                    pass

        # Azure Teams Enrollment
        if id_equipo and azure_id and id_equipo != "None":
            try:
                await graph.post(f"/groups/{id_equipo}/owners/$ref", {
                    "@odata.id": f"https://graph.microsoft.com/v1.0/directoryObjects/{azure_id}"
                })
                entry["teams_enroll"] = "owner"
            except Exception as e:
                if "already" not in str(e).lower():
                    error += f"TeamsEnroll: {str(e)} | "

        # Canvas
        canvas_id = None
        if plat in ("canvas", "both"):
            try:
                cu = await canvas.post(f"/accounts/{_ACCOUNT_LOCAL}/users", {
                    "user": {
                        "name": creds["full_name"],
                        "sortable_name": creds["full_name"],
                        "short_name": parts[0] + " " + parts[-1] if len(parts)>1 else creds["full_name"]
                    },
                    "pseudonym": {
                        "unique_id": login_id,
                        "sis_user_id": cedula,
                        "password": pwd,
                        "send_confirmation": False
                    },
                    "communication_channel": {
                        "type": "email", "address": login_id,
                        "skip_confirmation": True,
                    },
                })
                canvas_id = cu["id"]
                entry["canvas"] = "creado"
            except Exception as e:
                try:
                    match = await canvas.find_user_exact(_ACCOUNT_LOCAL, login_id)
                    if match:
                        canvas_id = match["id"]
                        entry["canvas"] = "exista"
                except:
                    pass
                if not canvas_id:
                    error += f"Canvas: {str(e)} | "

        # Canvas Enrollment
        if id_curso and canvas_id and id_curso != "None":
            try:
                await canvas.post(f"/courses/{id_curso}/enrollments", {
                    "enrollment": {
                        "user_id": canvas_id,
                        "type": "TeacherEnrollment",
                        "enrollment_state": "active",
                        "notify": False
                    }
                })
                entry["canvas_enroll"] = "teacher"
            except Exception as e:
                error += f"CanvasEnroll: {str(e)} | "
                
        if error:
            ws.cell(row=r_idx, column=col_enviado, value=f"⚠️ Error: {error}")
            ws.cell(row=r_idx, column=col_enviado).font = Font(color="D97706", bold=True)
            result.failed.append({"correo": login_id, "error": error})
        else:
            ws.cell(row=r_idx, column=col_usuario, value=login_id)
            ws.cell(row=r_idx, column=col_contra, value=pwd)
            collision_note = _collision_note(creds)
            if collision_note:
                ws.cell(row=r_idx, column=col_enviado, value=f"✅ OK{collision_note}")
                ws.cell(row=r_idx, column=col_enviado).font = Font(color="D97706", bold=True)
            else:
                ws.cell(row=r_idx, column=col_enviado, value="✅ OK")
                ws.cell(row=r_idx, column=col_enviado).font = Font(color="00B050", bold=True)
            result.succeeded.append(entry)

    if len(users_to_process) > 0:
        out_io = io.BytesIO()
        wb.save(out_io)
        out_io.seek(0)
        try:
            await graph.put_raw(f"/shares/{encoded_url}/driveItem/content", out_io.read())
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"No se pudo guardar el archivo actualizado en OneDrive. {e}")

    return result



@router.post("/excel/docentes-onedrive/sheets", response_model=list[str])
async def get_docentes_sheets(req: UrlOnlyRequest) -> list[str]:
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL invalida.")
    
    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo. {e}")

    try:
        import io
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
        sheets = wb.sheetnames
        wb.close()
        return sheets
    except Exception as e:
        raise HTTPException(status_code=400, detail="El archivo no es un Excel válido.")



@router.post("/excel/matriculaciones-onedrive/sheets")
async def get_matriculaciones_sheets(req: UrlOnlyRequest):
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")
    
    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo de OneDrive: {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
        return {"sheets": wb.sheetnames}
    except Exception as e:
        raise HTTPException(status_code=400, detail="El archivo descargado no es un Excel válido.")


@router.post("/excel/matriculaciones-onedrive/preview", summary="Previsualizar Matriculaciones desde OneDrive")
async def preview_matriculaciones_onedrive(req: DiplomadosUrlRequest) -> PreviewResponse:
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")
    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo: {e}")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail="El archivo descargado no es un Excel válido.")
    if req.sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"La pestaña '{req.sheet_name}' no existe.")
    ws = wb[req.sheet_name]
    
    sample_rows = []

    
    header_row_idx, headers_dict, headers_raw = _find_header_row_and_headers(ws)

    
    if not header_row_idx:
        raise HTTPException(status_code=400, detail="No se encontraron encabezados.")

    headers = [h for h in headers_raw if h]


    col_enviado = -1
    for i, h in enumerate(headers_raw):
        if "estado" in h.lower() or "enviado" in h.lower():
            col_enviado = i

    students_to_process = 0
    students_already_processed = 0

    for row_idx in range(header_row_idx + 1, ws.max_row + 1):
        row_vals = [str(ws.cell(row=row_idx, column=c).value or "").strip() for c in range(1, len(headers_raw) + 1)]
        if not any(row_vals):
            continue
            
        estado_val = row_vals[col_enviado] if col_enviado >= 0 else ""
        if estado_val.lower() == "ok" or "matriculado" in estado_val.lower():
            students_already_processed += 1
        else:
            students_to_process += 1
            
        if len(sample_rows) < 500:
            sample_rows.append({h: v for h, v in zip(headers_raw, row_vals) if h})

    wb.close()
    return PreviewResponse(
        sheet_name=req.sheet_name,
        students_to_process=students_to_process,
        students_already_processed=students_already_processed,
        headers=headers,
        sample_rows=sample_rows,
        student_details=[]
    )

async def _process_matriculaciones_bg(job_id: int, req: DiplomadosUrlRequest, contents: bytes, encoded_url: str):
    await jobs.start_job(job_id)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents))
    except Exception as e:
        await jobs.fail_job(job_id, f"El archivo no es un Excel válido: {e}")
        return

    if req.sheet_name not in wb.sheetnames:
        await jobs.fail_job(job_id, f"La pestaña '{req.sheet_name}' no existe.")
        return

    ws = wb[req.sheet_name]

    # Find header row (mismo helper robusto que usa el preview, para que
    # ambos detecten siempre la misma fila/columnas)
    header_row_idx, headers_norm, _headers_raw = _find_header_row_and_headers(ws)
    if not header_row_idx:
        await jobs.fail_job(job_id, "No se encontraron encabezados.")
        return
    headers = {n: c for n, c in headers_norm.items()}

    # Identify columns. La exclusión de columnas "(Autogenerado)" es SOLO
    # para usuario/correo: ahí sí hay un decoy real conocido ("Correo
    # (Autogenerado)", que registra el envío de credenciales, no un
    # identificador) — si se toma por error, la fila queda sin identificador
    # y se salta en silencio, dando "no hay filas nuevas" aunque el preview
    # sí muestre filas pendientes. "ID CANVAS (Autogenerado)"/"ID TEAMS
    # (Autogenerado)" en cambio SON el identificador real en planillas que
    # usan esa convención de nombres (ej. la misma que "Cursos y Equipos
    # (Plataformas)") — excluirlas ahí dejaba la matriculación sin columna
    # de Canvas/Teams detectada en absoluto.
    # "canvas"/"curso"/"equipo"/"teams" se evalúan ANTES que los
    # identificadores genéricos de usuario: un encabezado como "SIS ID
    # (Canvas)" contiene tanto "sis" como "canvas" — si "sis" ganara primero
    # (como pasaba antes, con un if/elif en ese orden), esa columna se
    # tomaba como identificador de usuario y la de curso quedaba sin
    # detectar. "canvas"/"curso"/"equipo"/"teams" son señales inequívocas de
    # a qué se refiere la columna, mientras que "sis" por sí solo es
    # ambiguo (puede ser el SIS ID del alumno o el SIS Course ID).
    user_col, canvas_col, teams_col, rol_col, env_col = None, None, None, None, None
    env_col_is_estado = False
    for n, col_idx in headers.items():
        is_autogen = "autogenerado" in n
        if "curso" in n or "canvas" in n:
            canvas_col = col_idx
        elif "equipo" in n or "teams" in n:
            teams_col = col_idx
        elif "rol" in n:
            rol_col = col_idx
        elif "estado" in n and not env_col_is_estado:
            env_col = col_idx
            env_col_is_estado = True
        elif "enviado" in n and not env_col_is_estado:
            env_col = col_idx
        elif not is_autogen and ("usuario" in n or "correo" in n or "email" in n or "cedula" in n or "sis" in n or "alumno" in n):
            user_col = col_idx

    if not user_col or not (canvas_col or teams_col):
        detected = [h for h in _headers_raw if h] or ["(ninguna)"]
        faltan = []
        if not user_col:
            faltan.append("usuario/correo/cédula/SIS")
        if not (canvas_col or teams_col):
            faltan.append("curso/canvas o equipo/teams")
        await jobs.fail_job(
            job_id,
            f"Falta columna de {' y '.join(faltan)}. Columnas detectadas en la planilla: {', '.join(detected)}.",
        )
        return
        
    if not env_col:
        env_col = ws.max_column + 1
        ws.cell(row=header_row_idx, column=env_col, value="Enviado")

    from app.routers.sync import _enroll_single, UnifiedEnrollment

    tasks = []
    success_count = 0
    error_count = 0
    total_to_process = ws.max_row - header_row_idx
    
    # Pre-calcular tareas válidas (filas)
    valid_rows = []
    for r_idx in range(header_row_idx + 1, ws.max_row + 1):
        user_val = str(ws.cell(row=r_idx, column=user_col).value or "").strip()
        if not user_val:
            continue
        estado_val = str(ws.cell(row=r_idx, column=env_col).value or "").strip().lower()
        if estado_val == "ok" or "matriculado" in estado_val:
            continue
        valid_rows.append(r_idx)
    
    if not valid_rows:
        await jobs.complete_job(job_id, 0, 0, "No hay filas nuevas para procesar.")
        return

    async def process_row(r_idx):
        nonlocal success_count, error_count
        user_val = str(ws.cell(row=r_idx, column=user_col).value or "").strip()
        canvas_val = str(ws.cell(row=r_idx, column=canvas_col).value or "").strip() if canvas_col else ""
        teams_val = str(ws.cell(row=r_idx, column=teams_col).value or "").strip() if teams_col else ""
        rol_val = str(ws.cell(row=r_idx, column=rol_col).value or "estudiante").strip().lower() if rol_col else "estudiante"
        
        # Mapear rol
        if "asistente" in rol_val or "ta" in rol_val:
            mapped_role = "ta"
        elif "diseñador" in rol_val or "designer" in rol_val:
            mapped_role = "designer"
        elif "observador" in rol_val or "observer" in rol_val:
            mapped_role = "observer"
        elif "prof" in rol_val or "own" in rol_val or "propiet" in rol_val:
            mapped_role = "teacher"
        else:
            mapped_role = "student"

        enroll_item = UnifiedEnrollment(
            user_identifier=user_val,
            canvas_course_id=canvas_val,
            teams_team_id=teams_val,
            role=mapped_role
        )
        
        try:
            enroll_res = await _enroll_single(enroll_item)
            if enroll_res.get("status") == "success":
                success_count += 1
                ws.cell(row=r_idx, column=env_col, value="OK").fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            else:
                msg = enroll_res.get("message", "Error desconocido")
                error_count += 1
                ws.cell(row=r_idx, column=env_col, value=f"Error: {msg}").fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        except Exception as ex:
            msg = str(ex)
            error_count += 1
            ws.cell(row=r_idx, column=env_col, value=f"Error: {msg}").fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    # Process in chunks and update DB progress periodically
    batch_size = 5
    for i in range(0, len(valid_rows), batch_size):
        chunk = valid_rows[i:i+batch_size]
        await asyncio.gather(*(process_row(r) for r in chunk))
        
        # Report progress to DB every batch
        await jobs.update_job_progress(
            job_id, 
            success_count, 
            error_count, 
            data_json=f'{{"total_to_process": {len(valid_rows)}, "processed": {success_count + error_count}}}'
        )

    # Save to OneDrive
    out_io = io.BytesIO()
    wb.save(out_io)
    out_io.seek(0)
    
    try:
        await graph.put_raw(f"/shares/{encoded_url}/driveItem/content", out_io.read())
        await jobs.complete_job(job_id, success_count, error_count, "Guardado en OneDrive correctamente.")
    except Exception as e:
        await jobs.complete_job(job_id, success_count, error_count, f"Procesado, pero no se pudo guardar en OneDrive: {e}")

@router.post("/excel/matriculaciones-onedrive")
async def import_matriculaciones_onedrive(req: DiplomadosUrlRequest, bg_tasks: BackgroundTasks) -> dict:
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")
    
    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error descargando archivo de OneDrive: {e}")

    job_id = await jobs.create_job(
        job_type="matriculacion_masiva_onedrive",
        operation="import",
        username="admin" # Ideally from auth token
    )
    
    bg_tasks.add_task(_process_matriculaciones_bg, job_id, req, contents, encoded_url)
    
    return {"status": "success", "job_id": job_id, "message": "Proceso de matriculación iniciado en segundo plano."}

@router.get("/excel/jobs/{job_id}")
async def get_job_status(job_id: int):
    job = await jobs.get_job(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Job no encontrado.")
    return job

async def _process_rollback_bg(job_id: int, req: DiplomadosUrlRequest, contents: bytes, encoded_url: str):
    _ACCOUNT_LOCAL = settings.canvas_account_id
    await jobs.start_job(job_id)
    try:
        import io, openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(contents))
    except Exception as e:
        await jobs.fail_job(job_id, f"El archivo no es un Excel válido: {e}")
        return

    if req.sheet_name not in wb.sheetnames:
        await jobs.fail_job(job_id, f"La pestaña '{req.sheet_name}' no existe en el archivo.")
        return

    ws = wb[req.sheet_name]
    
    def _norm(s):
        import unicodedata
        return unicodedata.normalize('NFKD', s).encode('ASCII', 'ignore').decode('utf-8').lower()
        
    header_row_idx = 1
    headers = {}
    for r in range(1, min(10, ws.max_row + 1)):
        row_vals = [str(ws.cell(row=r, column=c).value or "").strip() for c in range(1, ws.max_column + 1)]
        if any("nombre" in _norm(v) for v in row_vals):
            header_row_idx = r
            headers = {_norm(v): c for c, v in enumerate(row_vals, start=1) if v}
            break

    def get_col_idx(*keys):
        return _match_col_idx(headers, *keys)

    col_usuario = get_col_idx("usuario")
    col_enviado = get_col_idx("estado", "enviado")
    col_curso = get_col_idx("curso", "id curso", "canvas")
    col_equipo = get_col_idx("equipo", "id equipo", "teams")
    col_contra = get_col_idx("contrasena", "contraseña", "clave")

    if not col_usuario or not col_enviado:
        await jobs.fail_job(job_id, "No se encontraron las columnas necesarias (Usuario, Enviado/Estado).")
        return

    # Filter rows to rollback
    valid_rows = []
    for r_idx in range(header_row_idx + 1, ws.max_row + 1):
        enviado = str(ws.cell(row=r_idx, column=col_enviado).value or "").strip().lower()
        login_id = str(ws.cell(row=r_idx, column=col_usuario).value or "").strip()
        if login_id and login_id != "None" and ("ok" in enviado or "completado" in enviado or "creado" in enviado or "✅" in enviado):
            valid_rows.append(r_idx)

    success_count = 0
    error_count = 0

    async def process_rollback_row(r_idx):
        nonlocal success_count, error_count
        login_id = str(ws.cell(row=r_idx, column=col_usuario).value or "").strip()
        id_curso = str(ws.cell(row=r_idx, column=col_curso).value or "").strip() if col_curso else ""
        id_equipo = str(ws.cell(row=r_idx, column=col_equipo).value or "").strip() if col_equipo else ""
        
        error = ""
        uid = None
        try:
            user_data = await graph.get(f"/users/{login_id}", params={"$select": "id"})
            if user_data and user_data.get("id"):
                uid = user_data["id"]
        except: pass
        
        # Canvas Rollback — coincidencia EXACTA de login_id/email/SIS ID, no
        # "el primero que aparezca" de una búsqueda difusa: eso podría
        # desmatricular a otro alumno con correo/nombre parecido en vez del
        # que realmente corresponde a esta fila.
        if id_curso and id_curso != "None":
            try:
                match = await canvas.find_user_exact(_ACCOUNT_LOCAL, login_id)
                if match:
                    await canvas.remove_user_from_course(id_curso, str(match["id"]))
            except Exception as e:
                error += f"CanvasUnenroll: {e} | "
        
        # Teams Rollback
        if id_equipo and id_equipo != "None" and uid:
            try:
                await graph.remove_member_from_group(id_equipo, uid)
                await graph.remove_owner_from_group(id_equipo, uid)
            except Exception as e:
                error += f"TeamsUnenroll: {e} | "
                
        if error:
            error_count += 1
        else:
            ws.cell(row=r_idx, column=col_enviado, value="")
            if col_contra:
                ws.cell(row=r_idx, column=col_contra, value="")
            ws.cell(row=r_idx, column=col_usuario, value="")
            success_count += 1

    import asyncio
    batch_size = 5
    for i in range(0, len(valid_rows), batch_size):
        chunk = valid_rows[i:i+batch_size]
        await asyncio.gather(*(process_rollback_row(r) for r in chunk))
        
        await jobs.update_job_progress(
            job_id, 
            success_count, 
            error_count, 
            data_json=f'{{"total_to_process": {len(valid_rows)}, "processed": {success_count + error_count}}}'
        )

    out_io = io.BytesIO()
    wb.save(out_io)
    out_io.seek(0)
    
    try:
        await graph.put_raw(f"/shares/{encoded_url}/driveItem/content", out_io.read())
        await jobs.complete_job(job_id, success_count, error_count, "Reversión y guardado en OneDrive completados.")
    except Exception as e:
        await jobs.complete_job(job_id, success_count, error_count, f"Procesado, pero no se pudo guardar en OneDrive: {e}")


@router.post("/excel/rollback")
async def rollback_onedrive(req: DiplomadosUrlRequest, bg_tasks: BackgroundTasks) -> dict:
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")

    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo. {e}")

    job_id = await jobs.create_job(
        job_type="rollback_masivo_onedrive",
        operation="rollback",
        username="admin"
    )
    
    bg_tasks.add_task(_process_rollback_bg, job_id, req, contents, encoded_url)
    
    return {"status": "success", "job_id": job_id, "message": "Proceso de reversión iniciado en segundo plano."}
async def _process_matriculaciones_json_bg(job_id: int, data: list[dict]):
    from app.routers.sync import _enroll_single, UnifiedEnrollment
    import asyncio
    import json
    
    await jobs.start_job(job_id)
    
    success_count = 0
    error_count = 0
    results = []
    
    async def process_row(row):
        nonlocal success_count, error_count
        user_val = str(row.get("usuario") or "").strip()
        canvas_val = str(row.get("curso") or "").strip()
        teams_val = str(row.get("equipo") or "").strip()
        rol_val = str(row.get("rol") or "estudiante").strip().lower()
        
        if not user_val:
            return
            
        if "asistente" in rol_val or "ta" in rol_val:
            mapped_role = "ta"
        elif "diseñador" in rol_val or "designer" in rol_val:
            mapped_role = "designer"
        elif "observador" in rol_val or "observer" in rol_val:
            mapped_role = "observer"
        elif "prof" in rol_val or "own" in rol_val or "propiet" in rol_val:
            mapped_role = "teacher"
        else:
            mapped_role = "student"

        enroll_item = UnifiedEnrollment(
            user_identifier=user_val,
            canvas_course_id=canvas_val,
            teams_team_id=teams_val,
            role=mapped_role
        )
        
        try:
            enroll_res = await _enroll_single(enroll_item)
            if enroll_res.get("status") == "success":
                success_count += 1
                results.append({"usuario": user_val, "curso": canvas_val, "equipo": teams_val, "status": "✅ OK"})
            else:
                error_count += 1
                msg = enroll_res.get("message", "Error desconocido")
                results.append({"usuario": user_val, "curso": canvas_val, "equipo": teams_val, "status": f"❌ {msg}"})
        except Exception as e:
            error_count += 1
            results.append({"usuario": user_val, "curso": canvas_val, "equipo": teams_val, "status": f"❌ Error: {e}"})

    batch_size = 5
    for i in range(0, len(data), batch_size):
        chunk = data[i:i+batch_size]
        await asyncio.gather(*(process_row(r) for r in chunk))
        
        await jobs.update_job_progress(
            job_id, 
            success_count, 
            error_count, 
            data_json=json.dumps({"total_to_process": len(data), "processed": success_count + error_count, "results": results})
        )

    await jobs.complete_job(job_id, success_count, error_count, "Matriculación finalizada.")

@router.post("/excel/matriculaciones/json")
async def import_matriculaciones_json(req: JsonDataRequest, bg_tasks: BackgroundTasks) -> dict:
    if not req.data:
        raise HTTPException(status_code=400, detail="No hay datos para procesar.")
        
    job_id = await jobs.create_job(
        job_type="matriculacion_masiva_json",
        operation="import",
        username="admin"
    )
    
    bg_tasks.add_task(_process_matriculaciones_json_bg, job_id, req.data)
    return {"status": "success", "job_id": job_id, "message": "Proceso iniciado."}

async def _process_rollback_json_bg(job_id: int, data: list[dict]):
    import asyncio
    import json
    _ACCOUNT_LOCAL = settings.canvas_account_id
    await jobs.start_job(job_id)
    
    success_count = 0
    error_count = 0
    results = []
    
    async def process_rollback_row(row):
        nonlocal success_count, error_count
        login_id = str(row.get("usuario") or "").strip()
        id_curso = str(row.get("curso") or "").strip()
        id_equipo = str(row.get("equipo") or "").strip()
        
        if not login_id:
            return
            
        error = ""
        uid = None
        try:
            user_data = await graph.get(f"/users/{login_id}", params={"$select": "id"})
            if user_data and user_data.get("id"):
                uid = user_data["id"]
        except: pass
        
        # Canvas Rollback — coincidencia EXACTA, no "el primero que
        # aparezca" de la búsqueda difusa (ver mismo fix más arriba).
        if id_curso:
            try:
                match = await canvas.find_user_exact(_ACCOUNT_LOCAL, login_id)
                if match:
                    await canvas.remove_user_from_course(id_curso, str(match["id"]))
            except Exception as e:
                error += f"CanvasUnenroll: {e} | "
        
        # Teams Rollback
        if id_equipo and uid:
            try:
                await graph.remove_member_from_group(id_equipo, uid)
                await graph.remove_owner_from_group(id_equipo, uid)
            except Exception as e:
                error += f"TeamsUnenroll: {e} | "
                
        if error:
            error_count += 1
            results.append({"usuario": login_id, "status": f"❌ {error}"})
        else:
            success_count += 1
            results.append({"usuario": login_id, "status": "✅ Revertido"})

    batch_size = 5
    for i in range(0, len(data), batch_size):
        chunk = data[i:i+batch_size]
        await asyncio.gather(*(process_rollback_row(r) for r in chunk))
        
        await jobs.update_job_progress(
            job_id, 
            success_count, 
            error_count, 
            data_json=json.dumps({"total_to_process": len(data), "processed": success_count + error_count, "results": results})
        )

    await jobs.complete_job(job_id, success_count, error_count, "Reversión completada.")

@router.post("/excel/rollback/json")
async def rollback_json(req: JsonDataRequest, bg_tasks: BackgroundTasks) -> dict:
    if not req.data:
        raise HTTPException(status_code=400, detail="No hay datos para procesar.")
        
    job_id = await jobs.create_job(
        job_type="rollback_masivo_json",
        operation="rollback",
        username="admin"
    )
    
    bg_tasks.add_task(_process_rollback_json_bg, job_id, req.data)
    return {"status": "success", "job_id": job_id, "message": "Proceso de reversión iniciado."}

@router.post("/excel/diplomados/json")
async def import_diplomados_json(req: JsonDataRequest):
    if not req.data:
        raise HTTPException(status_code=400, detail="No hay datos para procesar.")
    if len(req.data) > 50:
        raise HTTPException(status_code=400, detail="Límite de 50 alumnos por lote para Diplomados excedido.")

    result = BulkResult(succeeded=[], failed=[])

    import re
    import time
    
    for row in req.data:
        nombre = str(row.get("nombre") or "").strip()
        cedula = _clean_cedula(str(row.get("cedula") or "").strip())
        curso_nombre = str(row.get("curso_nombre") or "").strip()
        row_team_owner_ids = list(_DIPLOMADO_DEFAULT_OWNERS) + (
            [_MBA_EXTRA_OWNER] if "mba" in curso_nombre.lower() else []
        )

        if not nombre or not cedula:
            continue
            
        try:
            # platform="both" (no solo "teams"): si la persona ya tiene cuenta
            # en Canvas (de este mismo flujo de MBA, de Alta Docentes, de
            # Ingreso, etc.) se detecta por CÉDULA -el chequeo más confiable,
            # porque es el SIS ID exacto- y se reutiliza su correo real, en
            # vez de intentar adivinarlo de nuevo y terminar creando una
            # cuenta duplicada cuando el correo real tiene otro formato
            # (ej. por una colisión de nombre resuelta distinto la vez
            # anterior).
            creds, status = await user_service.generate_unique_credentials(nombre, cedula, "both")
            login_id = creds["email"]
            pwd = creds["password"]
            
            error = None
            au_id = None
            
            # Azure AD Creation
            parts = creds["full_name"].strip().split()
            try:
                au_payload = {
                    "displayName": creds["full_name"],
                    "givenName": parts[0],
                    "surname": " ".join(parts[1:]) if len(parts) > 1 else "",
                    "userPrincipalName": login_id,
                    "mailNickname": login_id.replace(".", "_").replace("@", "_"),
                    "usageLocation": settings.usage_location,
                    "accountEnabled": True,
                    "jobTitle": "Alumno",
                    "country": "Paraguay",
                    "passwordProfile": {
                        "forceChangePasswordNextSignIn": False,
                        "password": pwd,
                    },
                }
                if cedula:
                    au_payload["postalCode"] = cedula
                au = await graph.post("/users", au_payload)
                au_id = au.get("id")
                await graph.assign_license(au_id, settings.azure_sku_students)
            except Exception as e:
                if "already exists" not in str(e).lower() and "Request_BadRequest" not in str(e):
                    error = str(e)
                else:
                    error = "Ya existía en Azure AD" 
                    
            if not error or "Ya existía" in error:
                # Group logic
                target_equipo = None
                if row.get("id_equipo"):
                    target_equipo = row.get("id_equipo")
                elif curso_nombre:
                    existing_tid = await graph.search_group_by_name(curso_nombre)
                    if existing_tid:
                        target_equipo = existing_tid
                    else:
                        nickname = _safe_mail_nickname(curso_nombre)
                        new_team = await graph.create_team_via_group(
                            display_name=curso_nombre,
                            mail_nickname=nickname,
                            description=f"Grupo para {curso_nombre}",
                            visibility="Private",
                            owner_ids=row_team_owner_ids
                        )
                        target_equipo = new_team.get("id")
                
                if target_equipo:
                    uid = au_id
                    if not uid:
                        try:
                            user_data = await graph.get(f"/users/{login_id}", params={"$select": "id"})
                            if user_data and user_data.get("id"):
                                uid = user_data["id"]
                        except: pass
                    
                    if uid:
                        try:
                            await graph.post(f"/groups/{target_equipo}/members/$ref", {"@odata.id": f"https://graph.microsoft.com/v1.0/directoryObjects/{uid}"})
                        except Exception as e:
                            if "already exist" not in str(e).lower():
                                error = f"{error or ''} | TeamsEnroll: {e}"
                                
            if not error or "Ya existía" in error:
                result.succeeded.append({"cedula": cedula, "nombre": creds["full_name"], "login_id": login_id, "password": pwd})
            else:
                result.failed.append({"input": {"cedula": cedula}, "error": error, "nombre": creds["full_name"]})

        except Exception as general_e:
            result.failed.append({"input": {"cedula": cedula}, "error": str(general_e), "nombre": nombre})
            
    return result

# ═══════════════════════════════════════════════════════════════════════════════
# Carga Masiva Genérica (OneDrive)
# ═══════════════════════════════════════════════════════════════════════════════
@router.post("/excel/masivo/sheets", response_model=list[str])
async def get_masivo_sheets(req: UrlOnlyRequest) -> list[str]:
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")
    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo. {e}")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
        sheets = wb.sheetnames
        wb.close()
        return sheets
    except Exception as e:
        raise HTTPException(status_code=400, detail="El archivo no es un Excel válido.")

@router.post("/excel/masivo/preview", summary="Previsualizar Carga Masiva de Usuarios")
async def preview_masivo_onedrive(req: DiplomadosUrlRequest) -> PreviewResponse:
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")
    
    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo de OneDrive. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail="El archivo descargado no es un Excel válido.")

    if req.sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"La pestaña '{req.sheet_name}' no existe.")

    ws = wb[req.sheet_name]
    
    sample_rows = []

    
    header_row_idx, headers_dict, headers_raw = _find_header_row_and_headers(ws)

    
    if not header_row_idx:
        raise HTTPException(status_code=400, detail="No se encontró la fila de encabezados.")

    headers = [h for h in headers_raw if h]


    col_estado = -1
    col_nombre = -1
    col_cedula = -1
    col_usuario = -1
    for i, h in enumerate(headers_raw):
        h_lower = h.lower()
        if "estado" in h_lower or "enviado" in h_lower:
            col_estado = i
        if "nombre" in h_lower:
            if col_nombre == -1: col_nombre = i
        if "cedula" in h_lower or "cédula" in h_lower or "ci" in h_lower:
            if col_cedula == -1: col_cedula = i
        if "usuario" in h_lower:
            if col_usuario == -1: col_usuario = i

    students_to_process = 0
    students_already_processed = 0
    student_details = []
    already_processed_details = []

    for row_idx in range(header_row_idx + 1, ws.max_row + 1):
        row_vals = [str(ws.cell(row=row_idx, column=c).value or "").strip() for c in range(1, len(headers_raw) + 1)]

        if not any(row_vals):
            continue

        nombre_val = row_vals[col_nombre] if col_nombre >= 0 else ""
        cedula_val = row_vals[col_cedula] if col_cedula >= 0 else ""
        if not nombre_val or not cedula_val:
            continue

        estado_val = row_vals[col_estado] if col_estado >= 0 else ""
        estado_lower = estado_val.lower()
        usuario_val = row_vals[col_usuario] if col_usuario >= 0 else ""

        if ("✅" in estado_val or estado_lower in ["si", "yes", "true", "enviado", "ok"]
                or "creado ok" in estado_lower or "ya exist" in estado_lower
                or (usuario_val and "@" in usuario_val)):
            students_already_processed += 1
            already_processed_details.append({
                "nombre": nombre_val, "cedula": cedula_val, "usuario": usuario_val, "estado": estado_val,
            })
        else:
            students_to_process += 1
            student_details.append({"nombre": nombre_val, "cedula": cedula_val})

        if len(sample_rows) < 10:
            sample_rows.append({h: v for h, v in zip(headers_raw, row_vals) if h})

    wb.close()
    return PreviewResponse(
        sheet_name=req.sheet_name,
        students_to_process=students_to_process,
        students_already_processed=students_already_processed,
        headers=headers,
        sample_rows=sample_rows,
        student_details=student_details,
        already_processed_details=already_processed_details,
    )



@router.post("/excel/masivo", summary="Importar Masivamente desde OneDrive (Sin Matriculación)", response_model=BulkResult)
async def import_masivo_onedrive(req: DiplomadosUrlRequest) -> BulkResult:
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")
    
    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo de OneDrive. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents))
    except Exception as e:
        raise HTTPException(status_code=400, detail="El archivo descargado no es un Excel válido.")

    _ACCOUNT_LOCAL = settings.canvas_account_id
    result = BulkResult()

    if req.sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"La pestaña '{req.sheet_name}' no existe.")

    ws = wb[req.sheet_name]
    
    header_row_idx, headers, _ = _find_header_row_and_headers(ws)

    
    if not header_row_idx:
        raise HTTPException(status_code=400, detail="No se encontraron las columnas requeridas (Nombre, Cedula).")

    def get_col_idx(*keys):
        return _match_col_idx(headers, *keys)

    col_nombre = get_col_idx("nombre", "alumno", "estudiante")
    col_cedula = get_col_idx("cedula", "cédula", "ci", "documento", "dni")
    col_correo = get_col_idx("correo", "email")
    col_plataforma = get_col_idx("plataforma", "platform")
    
    col_usuario = get_col_idx("usuario")
    col_contra = get_col_idx("contrasena", "contraseña", "clave")
    col_enviado = get_col_idx("estado", "enviado")

    next_col = ws.max_column + 1
    if not col_usuario:
        col_usuario = next_col
        ws.cell(row=header_row_idx, column=col_usuario, value="Usuario").font = Font(bold=True)
        next_col += 1
    if not col_contra:
        col_contra = next_col
        ws.cell(row=header_row_idx, column=col_contra, value="Contraseña").font = Font(bold=True)
        next_col += 1
    if not col_enviado:
        col_enviado = next_col
        ws.cell(row=header_row_idx, column=col_enviado, value="Enviado").font = Font(bold=True)
    
    sheet_lower = req.sheet_name.lower()
    default_plat = "both"
    if "canvas" in sheet_lower and "teams" not in sheet_lower:
        default_plat = "canvas"
    elif "teams" in sheet_lower and "canvas" not in sheet_lower:
        default_plat = "teams"

    async def process_row(r_idx):
        nombre = str(ws.cell(row=r_idx, column=col_nombre).value or "").strip()
        cedula = _clean_cedula(str(ws.cell(row=r_idx, column=col_cedula).value or "").strip())
        correo = str(ws.cell(row=r_idx, column=col_correo).value or "").strip() if col_correo else ""
        plat_val = str(ws.cell(row=r_idx, column=col_plataforma).value or default_plat).strip().lower() if col_plataforma else default_plat
        
        enviado = str(ws.cell(row=r_idx, column=col_enviado).value or "").strip()
        
        if not nombre or not cedula or cedula == "None":
            return
            
        usuario_val = str(ws.cell(row=r_idx, column=col_usuario).value or "").strip() if col_usuario else ""
        enviado_lower = enviado.lower()
        if "✅" in enviado or enviado_lower in ["si", "yes", "true", "enviado", "ok"] or "creado ok" in enviado_lower or "ya exist" in enviado_lower or (usuario_val and "@" in usuario_val):
            return

        if "canvas" in plat_val and "teams" in plat_val:
            plat = "both"
        elif "canvas" in plat_val:
            plat = "canvas"
        elif "teams" in plat_val:
            plat = "teams"
        else:
            plat = "both"

        creds, status = await user_service.generate_unique_credentials(nombre, cedula, plat)
        login_id = creds["email"]
        pwd = creds["password"]
        error_teams = ""
        error_canvas = ""
        
        # 1. Teams Creation
        if plat in ("teams", "both"):
            parts = creds["full_name"].strip().split()
            try:
                au_payload = {
                    "displayName": creds["full_name"],
                    "givenName": parts[0],
                    "surname": " ".join(parts[1:]) if len(parts) > 1 else "",
                    "userPrincipalName": login_id,
                    "mailNickname": login_id.replace(".", "_").replace("@", "_"),
                    "usageLocation": settings.usage_location,
                    "accountEnabled": True,
                    "jobTitle": "Alumno",
                    "country": "Paraguay",
                    "passwordProfile": {
                        "forceChangePasswordNextSignIn": False,
                        "password": pwd,
                    },
                }
                if cedula:
                    au_payload["postalCode"] = cedula
                au = await graph.post("/users", au_payload)
                await graph.assign_license(au["id"], settings.azure_sku_students)
            except Exception as e:
                if "already exists" not in str(e).lower() and "Request_BadRequest" not in str(e):
                    error_teams = f"Teams Error: {str(e)}"
                else:
                    error_teams = "Ya existía en Azure AD"
        
        # 2. Canvas Creation
        if plat in ("canvas", "both"):
            payload = {
                "user": {"name": creds["full_name"], "short_name": creds["full_name"]},
                "pseudonym": {"unique_id": login_id, "password": pwd, "sis_user_id": cedula, "send_confirmation": False},
                "communication_channel": {"type": "email", "address": login_id, "skip_confirmation": True},
            }
            try:
                await canvas.post(f"/accounts/{_ACCOUNT_LOCAL}/users", payload)
            except Exception as e:
                err_str = str(e).lower()
                if "already in use" not in err_str and "taken" not in err_str:
                    error_canvas = f"Canvas Error: {str(e)}"
                else:
                    error_canvas = "Ya existía en Canvas"

        final_error = []
        if error_teams: final_error.append(error_teams)
        if error_canvas: final_error.append(error_canvas)

        error_msg = " | ".join(final_error)

        # El correo de credenciales NO se envía aquí: es un paso aparte y
        # posterior (primero se crea, después se matricula, recién al final
        # se envían las credenciales) — ver /excel/envio-credenciales.
        if not error_msg and (not correo or correo == "None"):
            error_msg = "Creado OK (Sin correo personal)"

        if not error_msg or "Creado OK" in error_msg or "Ya existía" in error_msg:
            ws.cell(row=r_idx, column=col_usuario, value=login_id)
            ws.cell(row=r_idx, column=col_contra, value=pwd)

            if error_msg and "Ya existía" in error_msg:
                ws.cell(row=r_idx, column=col_enviado, value=f"⚠️ {error_msg}")
                ws.cell(row=r_idx, column=col_enviado).font = Font(color="D97706", bold=True)
                result.succeeded.append({"nombre": nombre, "correo": login_id, "mensaje": "Ya existía"})
            elif error_msg and "Sin correo personal" in error_msg:
                ws.cell(row=r_idx, column=col_enviado, value="✅ OK (Sin correo)")
                ws.cell(row=r_idx, column=col_enviado).font = Font(color="00B050", bold=True)
                result.succeeded.append({"nombre": nombre, "correo": login_id, "mensaje": "OK"})
            else:
                collision_note = _collision_note(creds)
                if collision_note:
                    ws.cell(row=r_idx, column=col_enviado, value=f"✅ OK{collision_note}")
                    ws.cell(row=r_idx, column=col_enviado).font = Font(color="D97706", bold=True)
                else:
                    ws.cell(row=r_idx, column=col_enviado, value="✅ OK")
                    ws.cell(row=r_idx, column=col_enviado).font = Font(color="00B050", bold=True)
                result.succeeded.append({"nombre": nombre, "correo": login_id, "mensaje": "OK"})
        else:
            ws.cell(row=r_idx, column=col_enviado, value=f"❌ Error: {error_msg}")
            ws.cell(row=r_idx, column=col_enviado).font = Font(color="FF0000", bold=True)
            result.failed.append({"nombre": nombre, "correo": login_id, "error": error_msg})

    tasks = []
    for r_idx in range(header_row_idx + 1, ws.max_row + 1):
        tasks.append(process_row(r_idx))
    
    batch_size = 5
    for i in range(0, len(tasks), batch_size):
        await asyncio.gather(*tasks[i:i+batch_size])

    out_io = io.BytesIO()
    wb.save(out_io)
    out_io.seek(0)
    
    try:
        await graph.put_raw(f"/shares/{encoded_url}/driveItem/content", out_io.read())
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"No se pudo guardar el archivo en OneDrive: {e}")

    return result


@router.post("/excel/courses/delete/preview")
async def preview_delete_courses_onedrive(req: DiplomadosUrlRequest) -> PreviewResponse:
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")
    
    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail="El archivo no es un Excel válido.")

    if req.sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"La pestaña '{req.sheet_name}' no existe.")

    ws = wb[req.sheet_name]
    header_row_idx, headers, _ = _find_header_row_and_headers(ws)
    if not header_row_idx:
        raise HTTPException(status_code=400, detail="No se encontraron cabeceras.")

    def get_col_idx(*keys):
        return _match_col_idx(headers, *keys)

    col_sis = get_col_idx("sisid", "sis", "id")
    col_nombre = get_col_idx("nombre", "curso")
    col_estado = get_col_idx("estado", "enviado", "eliminado", "baja")

    if not col_sis:
        raise HTTPException(status_code=400, detail="Se requiere una columna 'SIS ID' o 'ID' para borrar cursos.")

    sample_rows = []
    courses_to_process = 0
    courses_skipped = 0

    for r_idx in range(header_row_idx + 1, ws.max_row + 1):
        sis = str(ws.cell(row=r_idx, column=col_sis).value or "").strip()
        if not sis:
            continue
            
        estado = ""
        if col_estado:
            estado = str(ws.cell(row=r_idx, column=col_estado).value or "").strip().lower()

        if "ok" in estado or "eliminado" in estado or "baja" in estado:
            courses_skipped += 1
        else:
            courses_to_process += 1
            if len(sample_rows) < 10:
                row_vals = []
                for _, idx in headers.items():
                    val = ws.cell(row=r_idx, column=idx).value
                    row_vals.append(str(val).strip() if val is not None else "")
                sample_rows.append(dict(zip(headers.keys(), row_vals)))

    wb.close()
    return PreviewResponse(
        sheet_name=req.sheet_name,
        students_to_process=courses_to_process,
        students_already_processed=courses_skipped,
        headers=list(headers.keys()),
        sample_rows=sample_rows,
    )

@router.post("/excel/courses/delete/import")
async def import_delete_courses_onedrive(req: DiplomadosUrlRequest) -> BulkResult:
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")
    
    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents))
    except Exception as e:
        raise HTTPException(status_code=400, detail="El archivo no es un Excel válido.")

    result = BulkResult()
    if req.sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"La pestaña '{req.sheet_name}' no existe.")

    ws = wb[req.sheet_name]
    header_row_idx, headers, _ = _find_header_row_and_headers(ws)
    
    if not header_row_idx:
        raise HTTPException(status_code=400, detail="No se encontraron cabeceras.")

    def get_col_idx(*keys):
        return _match_col_idx(headers, *keys)

    # Sin fallback genérico "id": en un endpoint que ELIMINA cursos, si la
    # planilla no trae una columna de SIS ID reconocible, es más seguro
    # cortar con un error claro que adivinar con la primera columna que
    # contenga "id" en el nombre (ej. "ID Equipo" de Teams, un GUID que no
    # es un SIS ID de Canvas).
    col_sis = get_col_idx("sis id", "sisid", "sis")
    col_nombre = get_col_idx("nombre", "curso")
    col_estado = get_col_idx("estado", "enviado", "eliminado", "baja")

    if not col_sis:
        raise HTTPException(status_code=400, detail="Falta columna SIS ID.")

    if not col_estado:
        from openpyxl.styles import Font
        col_estado = ws.max_column + 1
        ws.cell(row=header_row_idx, column=col_estado, value="Estado/Eliminado").font = Font(bold=True)

    courses_to_process = []
    for r_idx in range(header_row_idx + 1, ws.max_row + 1):
        sis = str(ws.cell(row=r_idx, column=col_sis).value or "").strip()
        nombre = str(ws.cell(row=r_idx, column=col_nombre).value or "").strip() if col_nombre else ""
        estado = str(ws.cell(row=r_idx, column=col_estado).value or "").strip().lower()

        if not sis:
            continue
            
        if "ok" in estado or "eliminado" in estado or "baja" in estado:
            continue

        courses_to_process.append({"r_idx": r_idx, "sis": sis, "nombre": nombre})

    if len(courses_to_process) > 100:
        raise HTTPException(status_code=400, detail=f"Demasiados cursos a eliminar ({len(courses_to_process)}). Máximo 100 por ejecución.")

    if len(courses_to_process) > 0:
        try:
            await graph.put_raw(f"/shares/{encoded_url}/driveItem/content", contents)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"El archivo está bloqueado (cierra el Excel). Detalle: {e}")

    for course_data in courses_to_process:
        sis = course_data["sis"]
        r_idx = course_data["r_idx"]
        error = ""
        
        try:
            # 1. Search canvas by SIS ID
            canvas_courses = await canvas.paginate_limited(f"/accounts/{settings.canvas_account_id}/courses", {"search_term": sis, "per_page": 5}, max_records=5)
            target_course = None
            for c in canvas_courses:
                if c.get("sis_course_id") == sis or str(c.get("id")) == sis:
                    target_course = c
                    break
            
            if target_course:
                course_id = target_course["id"]
                course_name = target_course.get("name", sis)
                # 2. Delete Canvas
                await canvas.delete(f"/courses/{course_id}", {"event": "delete"})
                # 3. Delete Teams
                team_id = await graph.search_group_by_name(course_name)
                if team_id:
                    await graph.delete(f"/groups/{team_id}")
                else:
                    error += " | Teams: Equipo no encontrado"
            else:
                error = "No encontrado en Canvas"
                
        except Exception as e:
            error = f"Error: {str(e)}"
            
        if error and not error.startswith(" | Teams: Equipo no"):
            ws.cell(row=r_idx, column=col_estado, value=error)
            result.failed.append({"sis": sis, "error": error})
        else:
            ws.cell(row=r_idx, column=col_estado, value="OK (Eliminado)")
            result.succeeded.append({"sis": sis})

    if len(courses_to_process) > 0:
        out_io = io.BytesIO()
        wb.save(out_io)
        out_io.seek(0)
        try:
            await graph.put_raw(f"/shares/{encoded_url}/driveItem/content", out_io.read())
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"No se pudo guardar: {e}")

    return result


# ═══════════════════════════════════════════════════════════════════════════════
# Envío de Credenciales para Altas Nuevas (planilla "Usuarios (Nuevos)": Fecha,
# Nombre, Cédula, Correo (personal), Usuario Institucional (Autogenerado),
# Contraseña (Autogenerado), Estado (de creación), Correo (Estado) — última
# columna, la de envío). A diferencia del reenvío masivo, acá se usan el
# usuario/contraseña YA escritos en la planilla por el proceso de alta (no se
# recalculan), y solo se procesan filas cuya creación ya salió bien.
#
# Sin CC por defecto todavía (a pedido explícito): se pasa cc_override=[] para
# no aplicar el CC fijo genérico hasta que se defina a quién copiar.
# ═══════════════════════════════════════════════════════════════════════════════

@router.post("/excel/envio-credenciales/sheets", response_model=list[str])
async def get_envio_credenciales_sheets(req: UrlOnlyRequest) -> list[str]:
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")

    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
        sheets = wb.sheetnames
        wb.close()
        return sheets
    except Exception:
        raise HTTPException(status_code=400, detail="El archivo no es un Excel válido.")


def _envio_credenciales_cols(headers_dict: dict):
    def get_col_idx(*keys):
        return _match_col_idx(headers_dict, *keys)

    return {
        "nombre": get_col_idx("nombre"),
        "cedula": get_col_idx("cedula", "cédula", "ci"),
        "correo": get_col_idx("correo"),
        "usuario": get_col_idx("usuario institucional", "usuario"),
        "contra": get_col_idx("contrasena autogenerado", "contraseña autogenerado", "contrasena", "contraseña", "clave"),
        "estado_creacion": get_col_idx("estado"),
        "correo_estado": get_col_idx("correo estado", "correo enviado", "email enviado"),
    }


_ENVIO_CREDENCIALES_CC_KEY = "envio_credenciales_cc"


class CcSettingsResponse(BaseModel):
    cc: list[str] = []


class CcSettingsRequest(BaseModel):
    cc: list[str] = []


@router.get("/excel/envio-credenciales/cc-settings", response_model=CcSettingsResponse,
            summary="Obtener la lista de copia (CC) configurada para Envío de Credenciales")
async def get_envio_credenciales_cc() -> CcSettingsResponse:
    raw = await database.get_setting(_ENVIO_CREDENCIALES_CC_KEY)
    cc = [e.strip() for e in (raw or "").split(",") if e.strip()]
    return CcSettingsResponse(cc=cc)


@router.post("/excel/envio-credenciales/cc-settings", response_model=CcSettingsResponse,
             summary="Configurar la lista de copia (CC) para Envío de Credenciales")
async def set_envio_credenciales_cc(req: CcSettingsRequest) -> CcSettingsResponse:
    cc_clean = list(dict.fromkeys(e.strip() for e in req.cc if e.strip()))
    invalid = [e for e in cc_clean if "@" not in e]
    if invalid:
        raise HTTPException(status_code=400, detail=f"Correo(s) inválido(s): {', '.join(invalid)}")
    await database.set_setting(_ENVIO_CREDENCIALES_CC_KEY, ",".join(cc_clean))
    return CcSettingsResponse(cc=cc_clean)


class EnvioCredencialesPreviewResponse(BaseModel):
    sheet_name: str
    to_process: int
    already_processed: int
    not_created_yet: int
    headers: list[str] = []
    sample_rows: list[dict] = []


@router.post("/excel/envio-credenciales/preview", summary="Pre-visualizar envío de credenciales de altas nuevas")
async def preview_envio_credenciales(req: DiplomadosUrlRequest) -> EnvioCredencialesPreviewResponse:
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")

    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo de OneDrive. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="El archivo descargado no es un Excel válido.")

    if req.sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"La pestaña '{req.sheet_name}' no existe.")

    ws = wb[req.sheet_name]
    header_row_idx, headers_dict, headers_raw = _find_header_row_and_headers(ws)
    if not header_row_idx:
        raise HTTPException(status_code=400, detail="No se encontró la fila de encabezados.")

    cols = _envio_credenciales_cols(headers_dict)
    if not cols["nombre"] or not cols["correo"] or not cols["usuario"] or not cols["contra"]:
        raise HTTPException(status_code=400, detail="Faltan columnas de Nombre, Correo, Usuario Institucional o Contraseña.")

    to_process = 0
    already_processed = 0
    not_created_yet = 0
    sample_rows = []

    for r_idx in range(header_row_idx + 1, ws.max_row + 1):
        row_vals = [str(ws.cell(row=r_idx, column=c).value or "").strip() for c in range(1, len(headers_raw) + 1)]
        if not any(row_vals):
            continue
        nombre = row_vals[cols["nombre"] - 1]
        correo = row_vals[cols["correo"] - 1]
        usuario = row_vals[cols["usuario"] - 1]
        contra = row_vals[cols["contra"] - 1]
        estado_creacion = row_vals[cols["estado_creacion"] - 1].lower() if cols["estado_creacion"] else ""
        correo_estado = row_vals[cols["correo_estado"] - 1].lower() if cols["correo_estado"] else ""

        if not nombre or not correo or not usuario or not contra:
            continue

        creado_ok = "✅" in row_vals[cols["estado_creacion"] - 1] or "ok" in estado_creacion if cols["estado_creacion"] else True
        if not creado_ok:
            not_created_yet += 1
            continue

        if "ok" in correo_estado or "enviado" in correo_estado or "✅" in correo_estado:
            already_processed += 1
        else:
            to_process += 1
            if len(sample_rows) < 10:
                sample_rows.append({h: v for h, v in zip(headers_raw, row_vals) if h})

    return EnvioCredencialesPreviewResponse(
        sheet_name=req.sheet_name,
        to_process=to_process,
        already_processed=already_processed,
        not_created_yet=not_created_yet,
        headers=[h for h in headers_raw if h][:min(8, len(headers_raw))],
        sample_rows=sample_rows,
    )


@router.post("/excel/envio-credenciales", summary="Enviar credenciales masivamente a altas nuevas desde OneDrive")
async def envio_credenciales_onedrive(req: DiplomadosUrlRequest, bg_tasks: BackgroundTasks) -> dict:
    if not req.url or "http" not in req.url:
        raise HTTPException(status_code=400, detail="URL inválida.")

    encoded_url = _encode_share_url(req.url)
    try:
        contents = await graph.get_raw(f"/shares/{encoded_url}/driveItem/content")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo descargar el archivo de OneDrive. {e}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents))
    except Exception:
        raise HTTPException(status_code=400, detail="El archivo descargado no es un Excel válido.")

    if req.sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"La pestaña '{req.sheet_name}' no existe.")

    ws = wb[req.sheet_name]
    header_row_idx, headers_dict, headers_raw = _find_header_row_and_headers(ws)
    if not header_row_idx:
        raise HTTPException(status_code=400, detail="No se encontró la fila de encabezados.")

    cols = _envio_credenciales_cols(headers_dict)
    if not cols["nombre"] or not cols["correo"] or not cols["usuario"] or not cols["contra"]:
        raise HTTPException(status_code=400, detail="Faltan columnas de Nombre, Correo, Usuario Institucional o Contraseña.")

    col_correo_estado = cols["correo_estado"]
    if not col_correo_estado:
        col_correo_estado = ws.max_column + 1
        ws.cell(row=header_row_idx, column=col_correo_estado, value="Correo (Estado)").font = Font(bold=True)

    is_diplomado = "diplomado" in req.sheet_name.lower()

    rows_to_process = []
    for r_idx in range(header_row_idx + 1, ws.max_row + 1):
        nombre = str(ws.cell(row=r_idx, column=cols["nombre"]).value or "").strip()
        correo = str(ws.cell(row=r_idx, column=cols["correo"]).value or "").strip()
        usuario = str(ws.cell(row=r_idx, column=cols["usuario"]).value or "").strip()
        contra = str(ws.cell(row=r_idx, column=cols["contra"]).value or "").strip()
        estado_creacion_raw = str(ws.cell(row=r_idx, column=cols["estado_creacion"]).value or "").strip() if cols["estado_creacion"] else ""
        correo_estado = str(ws.cell(row=r_idx, column=col_correo_estado).value or "").strip().lower()

        if not nombre or not correo or not usuario or not contra:
            continue

        creado_ok = ("✅" in estado_creacion_raw or "ok" in estado_creacion_raw.lower()) if cols["estado_creacion"] else True
        if not creado_ok:
            continue
        if "ok" in correo_estado or "enviado" in correo_estado or "✅" in correo_estado:
            continue

        rows_to_process.append({"r_idx": r_idx, "nombre": nombre, "correo": correo, "usuario": usuario, "contra": contra})

    if len(rows_to_process) > 300:
        raise HTTPException(status_code=400, detail=f"Demasiados registros pendientes ({len(rows_to_process)}). Máximo 300 por ejecución.")

    if len(rows_to_process) > 0:
        try:
            await graph.put_raw(f"/shares/{encoded_url}/driveItem/content", contents)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"El archivo Excel está abierto o el enlace es de Solo Lectura. Detalle real: {e}")

    job_id = await jobs.create_job(job_type="envio_credenciales_altas_onedrive", operation="import", username="admin")
    bg_tasks.add_task(_process_envio_credenciales_bg, job_id, wb, ws.title, encoded_url, col_correo_estado, rows_to_process, is_diplomado)
    return {"status": "success", "job_id": job_id, "message": "Envío de credenciales iniciado en segundo plano.", "to_process": len(rows_to_process)}


async def _process_envio_credenciales_bg(
    job_id: int, wb, sheet_name: str, encoded_url: str, col_correo_estado: int,
    rows_to_process: list[dict], is_diplomado: bool,
):
    import json
    await jobs.start_job(job_id)
    ws = wb[sheet_name]

    # CC configurado desde la web (/ui/envio-credenciales) — vacío por
    # defecto hasta que se cargue explícitamente ahí, nunca el CC fijo
    # genérico de otros programas.
    cc_raw = await database.get_setting(_ENVIO_CREDENCIALES_CC_KEY)
    cc_list = [e.strip() for e in (cc_raw or "").split(",") if e.strip()]

    success_count = 0
    error_count = 0
    results: list[dict] = []

    async def process_row(row: dict):
        nonlocal success_count, error_count
        try:
            await email_service.send_credentials_email(
                to_email=row["correo"],
                full_name=row["nombre"],
                login_id=row["usuario"],
                password=row["contra"],
                program_type="diplomado" if is_diplomado else "grado",
                cc_override=cc_list,
            )
            ws.cell(row=row["r_idx"], column=col_correo_estado, value="✅ Enviado")
            success_count += 1
            results.append({"nombre": row["nombre"], "correo": row["correo"], "status": "✅ Enviado"})
        except Exception as e:
            msg = str(getattr(e, "detail", e))
            ws.cell(row=row["r_idx"], column=col_correo_estado, value=f"❌ Error: {msg}")
            error_count += 1
            results.append({"nombre": row["nombre"], "correo": row["correo"], "status": f"❌ {msg}"})

    # Concurrencia baja a propósito: cada correo de "grado" lleva un ZIP de
    # ~7MB adjunto, y con varios en simultáneo la app pisa el límite de
    # bytes entrantes de Graph ("ApplicationThrottled").
    batch_size = 2
    for i in range(0, len(rows_to_process), batch_size):
        chunk = rows_to_process[i:i + batch_size]
        await asyncio.gather(*(process_row(r) for r in chunk))
        await jobs.update_job_progress(
            job_id, success_count, error_count,
            data_json=json.dumps({"total_to_process": len(rows_to_process), "processed": success_count + error_count, "results": results}),
        )

    out_io = io.BytesIO()
    wb.save(out_io)
    out_io.seek(0)
    try:
        await graph.put_raw(f"/shares/{encoded_url}/driveItem/content", out_io.read())
        if error_count and success_count == 0:
            await jobs.fail_job(job_id, f"Todos los envíos fallaron ({error_count}).")
        else:
            await jobs.complete_job(job_id, success_count, error_count, "Envío y guardado en OneDrive completados.")
    except Exception as e:
        await jobs.complete_job(job_id, success_count, error_count, f"Procesado, pero no se pudo guardar en OneDrive: {e}")
